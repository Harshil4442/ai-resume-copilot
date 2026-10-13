"""Build the source-pinned document image and check genuine engines locally.

This is local native scanner/parser evidence, never Cloud Run sandbox proof.
Run with --platform linux/amd64 (CI) or linux/arm64 (a native ARM machine).
"""
from __future__ import annotations

import argparse
import base64
import hashlib
import importlib.util
import io
import json
import math
import os
import re
import selectors
import signal
import subprocess
import sys
import tempfile
import time
import uuid
from contextvars import ContextVar
from pathlib import Path
from typing import NamedTuple
from zipfile import ZIP_DEFLATED, ZipFile

ROOT = Path(__file__).resolve().parents[2]
SOURCES = (
    "infra/document-processing/Dockerfile", "infra/document-processing/processor.py",
    "infra/document-processing/service.py", "backend/app/services/parsing.py",
    "backend/app/services/market/skill_extractor.py", "backend/app/services/market/skill_taxonomy.py",
    "backend/app/services/market/security.py",
    "infra/document-processing/docx_renderer.py", "backend/app/services/resume_layout.py",
    "backend/app/services/pdf_streams.py", "backend/app/services/docx_layout_validation.py",
)
IMAGE_ID = re.compile(r"sha256:[0-9a-f]{64}\Z")
LABEL = "hirewiz.document-test.source."
MEMORY = 2 * 1024**3
TMPFS = "rw,noexec,nosuid,nodev,size=64m,mode=1777"
ERROR = "document_engine_check_refused"
_FAILURE_STAGES = frozenset({
    "setup", "daemon", "source_hashes", "build", "image_metadata", "runtime_sources",
    "corpus", "case_request", "case_execution", "case_validation", "layout_corpus",
    "source_recheck", "report_write", "clamd_corpus",
})
_FAILURE_CASES = frozenset({
    "clean_pdf", "clean_docx", "page_counter", "active_pdf", "includetext",
    "eicar_exact_file", "eicar_zip_member", "missing_definitions",
    "styled_multipage_fit", "rendered_reflow_refusal",
    "clamd_two_clean", "clamd_source_eicar", "clamd_candidate_eicar", "clamd_zip_member_eicar",
    "clamd_encrypted_pdf", "clamd_encrypted_zip", "clamd_max_filesize", "clamd_max_scansize",
    "clamd_max_files", "clamd_max_recursion", "clamd_resident_parser",
})
CLAMD_FAILURE_LINES = {'_command': (100, 103, 109, 112, 115), '_native_file': (202, 206, 214, 217), '_installed_clamd': (227, 231, 235), '_clamd_definition': (257, 261, 263), '_zip_check': (283,), '_zip_admit': (495, 497), '_read_clamd_maps': (850,), '_remaining': (528, 531), '_diagnostics': (565, 568, 572, 574, 581, 594, 596), '_drain': (609,), '_socket': (617, 625, 631), '_exchange': (639, 645, 652, 656), '_start': (667, 670, 677, 679, 682, 684, 686, 691, 695, 718, 722, 732), 'scan': (737, 742, 749, 751, 761, 772, 777, 778, 780, 782), 'close': (792, 806, 812, 815)}
CLAMD_FAILURE_PHASES = frozenset(('setup', 'request', 'source_scan', 'resident_parser', 'candidate_scan', 'owner_cleanup', 'ownership', 'memory', 'writer_version', 'report'))
CLAMD_FAILURE_CLASSES = frozenset(('InspectionDenied', 'InspectionUnavailable', 'MemoryError', 'TimeoutExpired', 'FileNotFoundError', 'PermissionError', 'TimeoutError', 'UnicodeError', 'OSError', 'ValueError', 'TypeError', 'KeyError', 'StopIteration', 'RuntimeError', 'SystemExit', 'KeyboardInterrupt', 'Unexpected'))
CLAMD_STARTUP_INDEX_SHA = "0b3ae37b66fd51626cef76d633051bc66348a52c0ce23f7653160e30647bf077"
CLAMD_STARTUP_STATUSES = frozenset(("unavailable", "observed_prefix", "no_unknown_complete_record_in_snapshot"))
CLAMD_STARTUP_SHAPES = frozenset(("empty", "non_printable_ascii", "leading_whitespace", "printable_ascii"))


_FAILURE: ContextVar[dict | None] = ContextVar("document_image_failure", default=None)
# Finite processor refusal protocol, synchronized with the source-pinned core.
# Never load app packages or copy a native response into diagnostic storage.
_DOCX_REASONS = frozenset({
    "document_safety", "renderer_identity", "document_size", "document_package", "package_changed", "edit_binding",
    "hyperlink_changed", "no_changed_paragraphs", "undeclared_native_change", "pdf_active_content", "pdf_link",
    "unsupported_graphics", "pdf_complexity", "unsupported_paint", "pdf_size", "pdf_pages", "pdf_geometry",
    "font_mapping", "pdf_text", "outside_region_pixels", "page_geometry_changed", "nontext_changed",
    "ambiguous_rendered_span", "text_or_reading_order_changed", "paragraph_page_changed", "line_wrap_changed",
    "font_changed", "retained_style_changed", "paragraph_overflow", "unsupported_graphics_overlap",
    "protected_text_changed", "protected_geometry_changed",
})
_DOCX_BINDING_KEYS = frozenset({"version", "operation", "sha256", "size_bytes", "source_format", "candidate_sha256",
                                "candidate_size_bytes", "edits_sha256"})
_DOCX_FIT_CHECKS = ("validator", "pages", "unchanged_links", "outside_changed_regions_pixel_delta",
                    "scan_engine", "candidate_scan")


class _DocxRefusal(NamedTuple):
    case: str
    error: str
    reason: str


class _DocxFitMismatch(NamedTuple):
    case: str
    failed_checks: tuple[str, ...]
    pages: int | None
    links: int | None


_LINK_FIELDS = ("native_hyperlinks", "imported_hyperlinks", "pdf_link_annotations",
                "native_uri_match", "imported_uri_match", "pdf_uri_match", "imported_label_match",
                "pdf_label_text_present", "pdf_label_glyph_occurrences", "pdf_label_in_page",
                "styled_pdf_link_annotations", "styled_pdf_uri_match", "styled_fixed_label_present")
_LINK_COUNT_INDEXES = (0, 1, 2, 8, 10)
_LINK_BOOL_INDEXES = (3, 4, 5, 6, 7, 9, 11, 12)


class _DocxLinkObservation(NamedTuple):
    container_calls: int
    elapsed: float
    values: tuple | None


EICAR = (b"X5O!P%@AP[4" + bytes([92]) + b"PZX54(P^)7CC)7}$"
         + b"EICAR-STANDARD-ANTIVIRUS-TEST-FILE!$H+H*")
# Runtime hashes attest every copied Python source; the Dockerfile remains
# bound to build input by its source label, like the native TeX helper.
RUNTIME = {
    path: "/opt/document/" + (path.rsplit("/", 1)[1] if path.startswith("infra/")
                            else path.removeprefix("backend/"))
    for path in SOURCES if not path.endswith("Dockerfile")
}
HASH_PROGRAM = ("import hashlib,json,pathlib; print(json.dumps({k:hashlib.sha256("
                "pathlib.Path(v).read_bytes()).hexdigest() for k,v in " + repr(RUNTIME) + ".items()}))")
DETECTION_PROGRAM = '''import base64,json,pathlib,subprocess,sys
sys.path.insert(0,"/opt/document")
import processor
payload=pathlib.Path("/tmp/request.json").read_bytes()
kind,content,sha256=processor.decode_request(payload)
target=pathlib.Path("/tmp/antivirus-test."+kind);target.write_bytes(content)
r=subprocess.run(["/usr/bin/clamscan","--stdout","--database=/var/lib/clamav",
 "--alert-exceeds-max=yes","--max-filesize=6M","--max-scansize=20M",
 "--max-files=500","--max-recursion=16","--max-scantime=20000","--",str(target)],
 capture_output=True,timeout=30)
found=r.returncode==1 and len(r.stdout)<=65536 and not r.stderr and b"EICAR" in r.stdout.upper() and b" FOUND" in r.stdout
p=subprocess.run([sys.executable,"-I","/opt/document/processor.py","/tmp/request.json"],
 capture_output=True,timeout=40)
expected={"version":1,"sha256":sha256,"size_bytes":len(content),"source_format":kind,"error":"document_inspection_refused"}
refused=p.returncode==1 and len(p.stdout)<=4096 and json.loads(p.stdout)==expected and not p.stderr
print(json.dumps({"eicar_detected":found,"processor_refused":refused}))
'''
CLAMD_NATIVE_PROGRAM = '"""Future Linux-image proof program; preparation must never execute this file."""\nimport base64\nimport hashlib\nimport json\nimport os\nfrom pathlib import Path\nimport resource\nimport sys\nimport tempfile\nimport time\n\nsys.path.insert(0, "/opt/document")\nimport processor\n\n\nCLAMD_FAILURE_LINES = {\'_command\': (100, 103, 109, 112, 115), \'_native_file\': (202, 206, 214, 217), \'_installed_clamd\': (227, 231, 235), \'_clamd_definition\': (257, 261, 263), \'_zip_check\': (283,), \'_zip_admit\': (495, 497), \'_read_clamd_maps\': (850,), \'_remaining\': (528, 531), \'_diagnostics\': (565, 568, 572, 574, 581, 594, 596), \'_drain\': (609,), \'_socket\': (617, 625, 631), \'_exchange\': (639, 645, 652, 656), \'_start\': (667, 670, 677, 679, 682, 684, 686, 691, 695, 718, 722, 732), \'scan\': (737, 742, 749, 751, 761, 772, 777, 778, 780, 782), \'close\': (792, 806, 812, 815)}\nCLAMD_FAILURE_PHASES = frozenset((\'setup\', \'request\', \'source_scan\', \'resident_parser\', \'candidate_scan\', \'owner_cleanup\', \'ownership\', \'memory\', \'writer_version\', \'report\'))\nCLAMD_FAILURE_CLASSES = frozenset((\'InspectionDenied\', \'InspectionUnavailable\', \'MemoryError\', \'TimeoutExpired\', \'FileNotFoundError\', \'PermissionError\', \'TimeoutError\', \'UnicodeError\', \'OSError\', \'ValueError\', \'TypeError\', \'KeyError\', \'StopIteration\', \'RuntimeError\', \'SystemExit\', \'KeyboardInterrupt\', \'Unexpected\'))\nCLAMD_PHASE = "setup"\n\n\ndef failure_receipt(error):\n    """Fixed classes and literal owned raise sites only; never format errors."""\n    classes = ((processor.InspectionDenied, "InspectionDenied"),\n        (processor.InspectionUnavailable, "InspectionUnavailable"),\n        (MemoryError, "MemoryError"), (processor.subprocess.TimeoutExpired, "TimeoutExpired"),\n        (FileNotFoundError, "FileNotFoundError"), (PermissionError, "PermissionError"),\n        (TimeoutError, "TimeoutError"), (UnicodeError, "UnicodeError"), (OSError, "OSError"),\n        (ValueError, "ValueError"), (TypeError, "TypeError"), (KeyError, "KeyError"),\n        (StopIteration, "StopIteration"), (RuntimeError, "RuntimeError"),\n        (SystemExit, "SystemExit"), (KeyboardInterrupt, "KeyboardInterrupt"))\n    chain, seen, truncated = [], set(), False\n    while error is not None:\n        if id(error) in seen or len(chain) == 8:\n            truncated = True\n            break\n        seen.add(id(error))\n        category = next((label for kind, label in classes if isinstance(error, kind)), "Unexpected")\n        frames, trace_seen, walked = [], set(), 0\n        trace = error.__traceback__\n        while trace is not None:\n            if id(trace) in trace_seen or walked == 64:\n                truncated = True\n                break\n            trace_seen.add(id(trace))\n            walked += 1\n            code, line = trace.tb_frame.f_code, trace.tb_lineno\n            if (code.co_filename == "/opt/document/processor.py"\n                    and line in CLAMD_FAILURE_LINES.get(code.co_name, ())):\n                if len(frames) == 8:\n                    truncated = True\n                    break\n                frames.append({"function": code.co_name, "line": line})\n            trace = trace.tb_next\n        chain.append({"class": category, "frames": frames})\n        error = error.__context__\n    return {"version": 1, "status": "native_clamd_proof_unavailable",\n            "phase": CLAMD_PHASE if CLAMD_PHASE in CLAMD_FAILURE_PHASES else "setup",\n            "chain": chain, "truncated": truncated}\n\n\ndef require(value):\n    if not value:\n        raise ValueError("native_clamd_proof_unavailable")\n\n\ndef memory_value(name):\n    with open("/sys/fs/cgroup/" + name, "rb") as stream:\n        raw = stream.read(33)\n    require(0 < len(raw) <= 32 and raw.strip().isdigit())\n    return int(raw)\n\n\ndef main():\n    global CLAMD_PHASE\n    os.umask(0o077)\n    resource.setrlimit(resource.RLIMIT_CORE, (0, 0))\n    resource.setrlimit(resource.RLIMIT_CPU, (40, 40))\n    resource.setrlimit(resource.RLIMIT_AS, (2 * 1024**3, 2 * 1024**3))\n    resource.setrlimit(resource.RLIMIT_FSIZE, (32 * 1024**2, 32 * 1024**2))\n    resource.setrlimit(resource.RLIMIT_NOFILE, (64, 64))\n    deadline = time.monotonic() + 50\n    CLAMD_PHASE = "request"\n    with open("/tmp/request.json", "rb") as stream:\n        raw = stream.read(processor.MAX_REQUEST + 1)\n    require(0 < len(raw) <= processor.MAX_REQUEST)\n    request = json.loads(raw, object_pairs_hook=processor._pairs)\n    require(set(request) == {"case", "source", "candidate", "expected", "resident_parser",\n                            "candidate_sha256", "candidate_size_bytes", "edits"})\n    require(request["expected"] in {"clean", "source_refused", "candidate_refused", "candidate_nonclean"})\n    require(type(request["resident_parser"]) is bool)\n    source = base64.b64decode(request["source"], validate=True)\n    candidate = base64.b64decode(request["candidate"], validate=True)\n    require(0 < len(source) <= processor.MAX_FILE and 0 < len(candidate) <= processor.MAX_FILE)\n    require(type(request["candidate_size_bytes"]) is int and len(candidate) == request["candidate_size_bytes"]\n            and hashlib.sha256(candidate).hexdigest() == request["candidate_sha256"])\n    outcome = "clean"\n    generated = False\n    with tempfile.TemporaryDirectory(prefix="clamd-proof-", dir="/tmp") as temporary:\n        directory = Path(temporary)\n        directory.chmod(0o700)\n        original = directory / "original.docx"\n        target = directory / "candidate.docx"\n        original.write_bytes(source)\n        original.chmod(0o600)\n        owner = processor.OwnedClamd(directory, deadline)\n        source_metadata = candidate_metadata = None\n        resident_current = None\n        with owner:\n            try:\n                CLAMD_PHASE = "source_scan"\n                source_metadata = owner.scan(original, source)\n            except processor.InspectionDenied:\n                outcome = "source_refused"\n            if outcome == "clean":\n                if request["resident_parser"]:\n                    # Actual source parsing and candidate regeneration while the\n                    # genuine engine remains resident, never synthetic telemetry.\n                    CLAMD_PHASE = "resident_parser"\n                    processor.validate_docx(source)\n                    from app.services.resume_layout import _apply_docx\n                    require(type(request["edits"]) is list and len(request["edits"]) == 1)\n                    candidate = _apply_docx(source, request["edits"])\n                    require(len(candidate) == request["candidate_size_bytes"]\n                            and hashlib.sha256(candidate).hexdigest() == request["candidate_sha256"])\n                    processor.validate_docx(candidate)\n                    generated = True\n                    resident_current = memory_value("memory.current")\n                target.write_bytes(candidate)\n                target.chmod(0o600)\n                try:\n                    CLAMD_PHASE = "candidate_scan"\n                    candidate_metadata = owner.scan(target, candidate)\n                except processor.InspectionDenied:\n                    outcome = "candidate_refused"\n                except processor.InspectionUnavailable:\n                    require(request["expected"] == "candidate_nonclean")\n                    outcome = "candidate_nonclean"\n            require(original.read_bytes() == source)\n            if target.exists():\n                require(target.read_bytes() == candidate)\n            CLAMD_PHASE = "owner_cleanup"\n        CLAMD_PHASE = "ownership"\n        require(owner.stopped and not owner.forced and owner.child is not None and owner.child.returncode == 0)\n        require(owner.peer is not None and owner.peer["pid"] == owner.child.pid)\n        require(owner.peer["uid"] == os.getuid() and owner.peer["gid"] == os.getgid())\n        require(len(owner.signatures) == 1 and owner.signatures[0] > 0)\n        require(not owner.log_path.exists())\n        require(time.monotonic() < deadline)\n        actual_expected = "candidate_refused" if request["expected"] == "candidate_nonclean" and outcome == "candidate_refused" else request["expected"]\n        require(outcome == actual_expected)\n        if request["expected"] == "clean":\n            require(source_metadata == candidate_metadata and len(owner.scan_bindings) == 2)\n        else:\n            require(outcome != "clean")\n        CLAMD_PHASE = "memory"\n        maximum, peak = memory_value("memory.max"), memory_value("memory.peak")\n        require(maximum == 2 * 1024**3 and 0 < peak <= maximum)\n        if request["resident_parser"]:\n            require(generated and type(resident_current) is int and 0 < resident_current <= maximum)\n        # A native Writer version operation starts only after the verified stop.\n        CLAMD_PHASE = "writer_version"\n        import docx_renderer\n        renderer_version = docx_renderer._version(deadline=deadline)\n        response = {"case": request["case"], "outcome": outcome, "native_installation": owner.manifest,\n                    "owned_peer": owner.peer, "loaded_signatures": owner.signatures[0],\n                    "scans": owner.scan_bindings, "daemon_returncode": owner.child.returncode,\n                    "protocol_scans": owner.protocol_bindings, "instream_attempts": owner.instream_attempts,\n                    "zip_admission_failure": owner.zip_admission_failure,\n                    "daemon_stopped_before_writer": owner.stopped, "native_log_removed": not owner.log_path.exists(),\n                    "renderer_version_after_stop": renderer_version,\n                    "resident_parser_exercised": generated, "resident_parser_cgroup_current_bytes": resident_current,\n                    "cgroup_memory_max_bytes": maximum, "cgroup_memory_peak_bytes": peak,\n                    "performance_proven": False, "cloud_run_isolation_proven": False}\n        CLAMD_PHASE = "report"\n        encoded = json.dumps(response, separators=(",", ":"), allow_nan=False).encode()\n        require(len(encoded) <= 16384)\n        sys.stdout.buffer.write(encoded + b"\\n")\n\n\nNATIVE_REQUIREMENT_SITES = {144: \'shutdown\', 145: \'pid\', 146: \'uid\', 147: \'signatures\', 148: \'log_cleanup\', 149: \'deadline\', 151: \'case_outcome\', 153: \'clean_metadata\', 155: \'negative_nonclean\'}\nNATIVE_REQUIREMENTS = frozenset((\'shutdown\', \'pid\', \'uid\', \'signatures\', \'log_cleanup\', \'deadline\', \'case_outcome\', \'clean_metadata\', \'negative_nonclean\'))\n\n\ndef native_requirement_receipt(error):\n    """Only exact own code frames at fixed require sites; no locals or values."""\n    if type(error) is not ValueError:\n        return None\n    trace, previous, seen = error.__traceback__, None, set()\n    for _ in range(64):\n        if trace is None or id(trace) in seen:\n            return None\n        seen.add(id(trace))\n        code = trace.tb_frame.f_code\n        if (code is require.__code__ and previous is not None\n                and previous[0] is main.__code__ and trace.tb_next is None):\n            label = NATIVE_REQUIREMENT_SITES.get(previous[1])\n            return label if label in NATIVE_REQUIREMENTS else None\n        previous = code, trace.tb_lineno\n        trace = trace.tb_next\n    return None\n\n\nif __name__ == "__main__":\n    try:\n        main()\n    except BaseException as error:\n        diagnostic = failure_receipt(error)\n        requirement = native_requirement_receipt(error)\n        if requirement is not None:\n            diagnostic["native_requirement"] = requirement\n        print(json.dumps(diagnostic, separators=(",", ":"), allow_nan=False))\n        raise SystemExit(1) from None\n'
CLAMD_STARTUP_DIAGNOSTIC_PROGRAM = '"""Future Linux-image proof program; preparation must never execute this file."""\nimport base64\nimport hashlib\nimport json\nimport os\nfrom pathlib import Path\nimport resource\nimport sys\nimport tempfile\nimport time\n\nsys.path.insert(0, "/opt/document")\nimport processor\n\n"""PRIVATE diagnostic derivative only; no native execution or product change on import."""\nimport hashlib\nimport os\nimport re\nimport stat\nimport time\n\nPROCESSOR_SHA = "803a150fe615abea0a8d4bdc367bcf5177f8e7f9abdadf93494b372d3ffece35"\nPROCESSOR_PATH = "/opt/document/processor.py"\nCAP = 65536\nLINE_CAP = 2048\nRECORD_CAP = 4\nPUBLIC_TEMPLATE_INDEX_SHA = "0b3ae37b66fd51626cef76d633051bc66348a52c0ce23f7653160e30647bf077"\nPUBLIC_PATTERNS = ((0, \'\\\\+\\\\+\\\\+\\\\ Started\\\\ at\\\\ [ -~]{0,256}\'), (1, \'Received\\\\ -?[0-9]{1,20}\\\\ file\\\\ descriptor\\\\(s\\\\)\\\\ from\\\\ systemd\\\\.\'), (2, \'clamd\\\\ daemon\\\\ [ -~]{0,256}\\\\ \\\\(OS:\\\\ [ -~]{0,256},\\\\ ARCH:\\\\ [ -~]{0,256},\\\\ CPU:\\\\ [ -~]{0,256}\\\\)\'), (3, \'Running\\\\ as\\\\ user\\\\ [ -~]{0,256}\\\\ \\\\(UID\\\\ [0-9]{1,20},\\\\ GID\\\\ [0-9]{1,20}\\\\)\'), (4, \'Running\\\\ on\\\\ 32\\\\-bit\\\\ system,\\\\ and\\\\ RLIMIT_DATA\\\\ >\\\\ 2GB,\\\\ lowering\\\\ to\\\\ 2GB!\'), (5, \'Log\\\\ file\\\\ size\\\\ limited\\\\ to\\\\ -?[0-9]{1,20}\\\\ bytes\\\\.\'), (6, \'Log\\\\ file\\\\ size\\\\ limit\\\\ disabled\\\\.\'), (7, \'Reading\\\\ databases\\\\ from\\\\ [ -~]{0,256}\'), (8, \'Excluded\\\\ PUA\\\\ categories:\'), (9, \'\\\\ [ -~]{0,256}\'), (10, \'\'), (11, \'Included\\\\ PUA\\\\ categories:\'), (12, \'\\\\ [ -~]{0,256}\'), (13, \'\'), (14, \'Not\\\\ loading\\\\ PUA\\\\ signatures\\\\.\'), (15, \'Only\\\\ loading\\\\ official\\\\ signatures\\\\.\'), (16, \'Not\\\\ loading\\\\ phishing\\\\ signatures\\\\.\'), (17, \'Bytecode:\\\\ Security\\\\ mode\\\\ set\\\\ to\\\\ "TrustSigned"\\\\.\'), (18, \'Bytecode:\\\\ Security\\\\ mode\\\\ set\\\\ to\\\\ "Paranoid"\\\\.\'), (19, \'Bytecode:\\\\ Enabled\\\\ support\\\\ for\\\\ unsigned\\\\ bytecode\\\\.\'), (20, \'Bytecode\\\\ support\\\\ disabled\\\\.\'), (21, \'Disabling\\\\ URL\\\\ based\\\\ phishing\\\\ detection\\\\.\'), (22, \'Only\\\\ using\\\\ the\\\\ A\\\\-C\\\\ matcher\\\\.\'), (23, \'Max\\\\ A\\\\-C\\\\ depth\\\\ set\\\\ to\\\\ [0-9]{1,20}\'), (24, \'Loaded\\\\ [0-9]{1,20}\\\\ signatures\\\\.\'), (25, \'Socket\\\\ file\\\\ removed\\\\.\'), (26, \'No\\\\ stats\\\\ for\\\\ Database\\\\ check\\\\ \\\\-\\\\ forcing\\\\ reload\'), (27, \'SelfCheck:\\\\ Database\\\\ modification\\\\ detected\\\\.\\\\ Forcing\\\\ reload\\\\.\'), (28, \'SelfCheck:\\\\ Database\\\\ status\\\\ OK\\\\.\'), (29, \'Reading\\\\ databases\\\\ from\\\\ [ -~]{0,256}\'), (30, \'Database\\\\ correctly\\\\ reloaded\\\\ \\\\([0-9]{1,20}\\\\ signatures\\\\)\'), (31, \'Database\\\\ reload\\\\ completed\\\\.\'), (32, \'Limits:\\\\ Global\\\\ time\\\\ limit\\\\ set\\\\ to\\\\ [0-9]{1,20}\\\\ milliseconds\\\\.\'), (33, \'Limits:\\\\ Global\\\\ size\\\\ limit\\\\ set\\\\ to\\\\ [0-9]{1,20}\\\\ bytes\\\\.\'), (34, \'Limits:\\\\ File\\\\ size\\\\ limit\\\\ set\\\\ to\\\\ [0-9]{1,20}\\\\ bytes\\\\.\'), (35, \'Limits:\\\\ Recursion\\\\ level\\\\ limit\\\\ set\\\\ to\\\\ [0-9]{1,20}\\\\.\'), (36, \'Limits:\\\\ Files\\\\ limit\\\\ set\\\\ to\\\\ [0-9]{1,20}\\\\.\'), (37, \'Limits:\\\\ MaxEmbeddedPE\\\\ limit\\\\ set\\\\ to\\\\ [0-9]{1,20}\\\\ bytes\\\\.\'), (38, \'Limits:\\\\ MaxHTMLNormalize\\\\ limit\\\\ set\\\\ to\\\\ [0-9]{1,20}\\\\ bytes\\\\.\'), (39, \'Limits:\\\\ MaxHTMLNoTags\\\\ limit\\\\ set\\\\ to\\\\ [0-9]{1,20}\\\\ bytes\\\\.\'), (40, \'Limits:\\\\ MaxScriptNormalize\\\\ limit\\\\ set\\\\ to\\\\ [0-9]{1,20}\\\\ bytes\\\\.\'), (41, \'Limits:\\\\ MaxZipTypeRcg\\\\ limit\\\\ set\\\\ to\\\\ [0-9]{1,20}\\\\ bytes\\\\.\'), (42, \'Limits:\\\\ MaxPartitions\\\\ limit\\\\ set\\\\ to\\\\ [0-9]{1,20}\\\\.\'), (43, \'Limits:\\\\ MaxIconsPE\\\\ limit\\\\ set\\\\ to\\\\ [0-9]{1,20}\\\\.\'), (44, \'Limits:\\\\ MaxRecHWP3\\\\ limit\\\\ set\\\\ to\\\\ [0-9]{1,20}\\\\.\'), (45, \'Limits:\\\\ PCREMatchLimit\\\\ limit\\\\ set\\\\ to\\\\ [0-9]{1,20}\\\\.\'), (46, \'Limits:\\\\ PCRERecMatchLimit\\\\ limit\\\\ set\\\\ to\\\\ [0-9]{1,20}\\\\.\'), (47, \'Limits:\\\\ PCREMaxFileSize\\\\ limit\\\\ set\\\\ to\\\\ [0-9]{1,20}\\\\.\'), (48, \'Archive\\\\ support\\\\ enabled\\\\.\'), (49, \'Archive\\\\ support\\\\ disabled\\\\.\'), (50, \'Image\\\\ \\\\(graphics\\\\)\\\\ scanning\\\\ support\\\\ enabled\\\\.\'), (51, \'Image\\\\ \\\\(graphics\\\\)\\\\ scanning\\\\ support\\\\ disabled\\\\.\'), (52, \'Detection\\\\ using\\\\ image\\\\ fuzzy\\\\ hash\\\\ enabled\\\\.\'), (53, \'Detection\\\\ using\\\\ image\\\\ fuzzy\\\\ hash\\\\ disabled\\\\.\'), (54, \'Alerting\\\\ of\\\\ encrypted\\\\ archives\\\\ _and_\\\\ documents\\\\ enabled\\\\.\'), (55, \'Alerting\\\\ of\\\\ encrypted\\\\ documents\\\\ enabled\\\\.\'), (56, \'Alerting\\\\ of\\\\ encrypted\\\\ archives\\\\ enabled\\\\.\'), (57, \'Alerting\\\\ of\\\\ encrypted\\\\ documents\\\\ enabled\\\\.\'), (58, \'Heuristic\\\\ alerting\\\\ enabled\\\\ for\\\\ scans\\\\ that\\\\ exceed\\\\ set\\\\ maximums\\\\.\'), (59, \'AlertExceedsMax\\\\ heuristic\\\\ detection\\\\ disabled\\\\.\'), (60, \'Heuristic\\\\ alerts\\\\ disabled\\\\.\'), (61, \'Heuristic\\\\ alerts\\\\ enabled\\\\.\'), (62, \'Portable\\\\ Executable\\\\ support\\\\ enabled\\\\.\'), (63, \'Portable\\\\ Executable\\\\ support\\\\ disabled\\\\.\'), (64, \'ELF\\\\ support\\\\ enabled\\\\.\'), (65, \'ELF\\\\ support\\\\ disabled\\\\.\'), (66, \'Alerting\\\\ on\\\\ broken\\\\ executables\\\\ enabled\\\\.\'), (67, \'Media\\\\ \\\\(Graphics\\\\)\\\\ Format\\\\ Validation\\\\ enabled\'), (68, \'Mail\\\\ files\\\\ support\\\\ enabled\\\\.\'), (69, \'Mail:\\\\ RFC1341\\\\ handling\\\\ enabled\\\\.\'), (70, \'Mail\\\\ files\\\\ support\\\\ disabled\\\\.\'), (71, \'OLE2\\\\ support\\\\ enabled\\\\.\'), (72, \'OLE2:\\\\ Alerting\\\\ on\\\\ all\\\\ VBA\\\\ macros\\\\.\'), (73, \'OLE2\\\\ support\\\\ disabled\\\\.\'), (74, \'PDF\\\\ support\\\\ enabled\\\\.\'), (75, \'PDF\\\\ support\\\\ disabled\\\\.\'), (76, \'SWF\\\\ support\\\\ enabled\\\\.\'), (77, \'SWF\\\\ support\\\\ disabled\\\\.\'), (78, \'HTML\\\\ support\\\\ enabled\\\\.\'), (79, \'HTML\\\\ support\\\\ disabled\\\\.\'), (80, \'XMLDOCS\\\\ support\\\\ enabled\\\\.\'), (81, \'XMLDOCS\\\\ support\\\\ disabled\\\\.\'), (82, \'HWP3\\\\ support\\\\ enabled\\\\.\'), (83, \'HWP3\\\\ support\\\\ disabled\\\\.\'), (84, \'OneNote\\\\ support\\\\ enabled\\\\.\'), (85, \'OneNote\\\\ support\\\\ disabled\\\\.\'), (86, \'Phishing:\\\\ Always\\\\ checking\\\\ for\\\\ cloaked\\\\ urls\'), (87, \'Phishing:\\\\ Always\\\\ checking\\\\ for\\\\ ssl\\\\ mismatches\'), (88, \'Raw\\\\ DMG:\\\\ Alert\\\\ on\\\\ partitions\\\\ intersections\'), (89, \'Heuristic:\\\\ precedence\\\\ enabled\'), (90, \'Structured:\\\\ Minimum\\\\ Credit\\\\ Card\\\\ Number\\\\ Count\\\\ set\\\\ to\\\\ [0-9]{1,20}\'), (91, \'Structured:\\\\ Minimum\\\\ Social\\\\ Security\\\\ Number\\\\ Count\\\\ set\\\\ to\\\\ [0-9]{1,20}\'), (92, \'Self\\\\ checking\\\\ disabled\\\\.\'), (93, \'Self\\\\ checking\\\\ every\\\\ [0-9]{1,20}\\\\ seconds\\\\.\'), (94, \'SIGHUP\\\\ caught:\\\\ re\\\\-opening\\\\ log\\\\ file\\\\.\'), (95, \'Activating\\\\ the\\\\ newly\\\\ loaded\\\\ database\\\\.\\\\.\\\\.\'), (96, \'Pid\\\\ file\\\\ removed\\\\.\'), (97, \'\\\\-\\\\-\\\\-\\\\ Stopped\\\\ at\\\\ [ -~]{0,256}\'), (98, \'LOCAL:\\\\ No\\\\ local\\\\ AF_UNIX\\\\ SOCK_STREAM\\\\ socket\\\\ received\\\\ from\\\\ systemd\\\\.\'), (99, \'LOCAL:\\\\ Received\\\\ AF_UNIX\\\\ SOCK_STREAM\\\\ socket\\\\ from\\\\ systemd\\\\.\'), (100, \'Localserver:\\\\ Creating\\\\ socket\\\\ directory:\\\\ [ -~]{0,256}\'), (101, \'LOCAL:\\\\ Removing\\\\ stale\\\\ socket\\\\ file\\\\ [ -~]{0,256}\'), (102, \'LOCAL:\\\\ Unix\\\\ socket\\\\ file\\\\ [ -~]{0,256}\'), (103, \'LOCAL:\\\\ Setting\\\\ connection\\\\ queue\\\\ length\\\\ to\\\\ -?[0-9]{1,20}\'))\nKNOWN_PATTERNS = (\'Loaded ([1-9][0-9]{0,8}) signatures\\\\.\', \'clamd daemon 1\\\\.4\\\\.3 \\\\(OS: [A-Za-z0-9_.-]+, ARCH: [A-Za-z0-9_.-]+, CPU: [A-Za-z0-9_.-]+\\\\)\', \'Limits: (?:MaxEmbeddedPE|MaxHTMLNormalize|MaxHTMLNoTags|MaxScriptNormalize|MaxZipTypeRcg) limit set to [0-9]+ bytes\\\\.\', \'Limits: (?:MaxPartitions|MaxIconsPE|MaxRecHWP3|PCREMatchLimit|PCRERecMatchLimit|PCREMaxFileSize) limit set to [0-9]+\\\\.\', \'(?:\\\\+\\\\+\\\\+ Started|--- Stopped) at [A-Za-z0-9 :]{24}\', \'instream\\\\(local\\\\): [A-Za-z0-9][A-Za-z0-9._-]{0,199} FOUND\')\n_PUBLIC = tuple((index, re.compile(pattern)) for index, pattern in PUBLIC_PATTERNS)\n_KNOWN = tuple(re.compile(pattern) for pattern in KNOWN_PATTERNS)\n_WITNESS = None\n\n\ndef _unknown_raise(error):\n    seen = set()\n    for _ in range(8):\n        if error is None or id(error) in seen:\n            return False\n        seen.add(id(error))\n        trace = error.__traceback__\n        walked = set()\n        for _ in range(64):\n            if trace is None or id(trace) in walked:\n                break\n            walked.add(id(trace))\n            code = trace.tb_frame.f_code\n            if (code.co_filename == PROCESSOR_PATH and code.co_name == "_diagnostics"\n                    and trace.tb_lineno == 594):\n                return True\n            trace = trace.tb_next\n        error = error.__context__\n    return False\n\n\ndef _known(line, self, processor):\n    try:\n        value = line.decode("ascii", errors="strict")\n    except UnicodeError:\n        return False\n    DATABASE = processor.DATABASE\n    CLAMD_SETTINGS = processor.CLAMD_SETTINGS\n    allowed = {\'Received 0 file descriptor(s) from systemd.\', \'Not loading PUA signatures.\', \'Only loading official signatures.\', \'Bytecode: Security mode set to "TrustSigned".\', \'Reading databases from \' + DATABASE, \'Self checking disabled.\', \'Socket file removed.\', \'Log file size limit disabled.\', \'Log file size limited to 9223372036854775807 bytes.\', \'LOCAL: Unix socket file \' + str(self.socket_path), \'LOCAL: Setting connection queue length to \' + CLAMD_SETTINGS[\'MaxConnectionQueueLength\'], \'Alerting of encrypted archives _and_ documents enabled.\', \'Archive support enabled.\', \'Image (graphics) scanning support enabled.\', \'Detection using image fuzzy hash enabled.\', \'Alerting of encrypted archives enabled.\', \'Alerting of encrypted documents enabled.\', \'Heuristic alerts enabled.\', \'Heuristic alerting enabled for scans that exceed set maximums.\', \'Portable Executable support enabled.\', \'ELF support enabled.\', \'Mail files support enabled.\', \'OLE2 support enabled.\', \'PDF support enabled.\', \'SWF support enabled.\', \'HTML support enabled.\', \'XMLDOCS support enabled.\', \'HWP3 support enabled.\', \'OneNote support enabled.\', \'Limits: Global time limit set to 20000 milliseconds.\', \'Limits: Global size limit set to 20971520 bytes.\', \'Limits: File size limit set to 6291456 bytes.\', \'Limits: Recursion level limit set to 16.\', \'Limits: Files limit set to 500.\'}\n    return value in allowed or any(pattern.fullmatch(value) for pattern in _KNOWN)\n\n\ndef _shape(line):\n    if not line:\n        return "empty"\n    if any(value < 32 or value > 126 for value in line):\n        return "non_printable_ascii"\n    if line[:1] in (b" ", b"\\t"):\n        return "leading_whitespace"\n    return "printable_ascii"\n\n\ndef _snapshot(self):\n    # Only the original owned byte buffers and original same-inode descriptor.\n    # No reads from arbitrary paths, no original descriptor offset change.\n    child = self.child\n    if (child is None or child.returncode is not None or type(child.pid) is not int\n            or child.pid <= 0 or os.getpgid(child.pid) != child.pid):\n        raise ValueError("witness_unavailable")\n    directory = self.directory.lstat()\n    if (not stat.S_ISDIR(directory.st_mode) or stat.S_IMODE(directory.st_mode) != 0o700\n            or directory.st_uid != os.getuid() or directory.st_gid != os.getgid()):\n        raise ValueError("witness_unavailable")\n    stdout = self.pipes.get("stdout")\n    if (type(stdout) is not dict or stdout.get("stream") is not child.stdout\n            or type(stdout.get("raw")) is not bytearray or len(stdout["raw"]) > CAP):\n        raise ValueError("witness_unavailable")\n    observed = [("stdout", bytes(stdout["raw"]))]\n    if self.log_fd is None or self.log_path.parent != self.directory or self.log_path.name != "clamd.log":\n        raise ValueError("witness_unavailable")\n    info, path = os.fstat(self.log_fd), self.log_path.lstat()\n    if (not stat.S_ISREG(info.st_mode) or not stat.S_ISREG(path.st_mode)\n            or (info.st_dev, info.st_ino) != self.log_identity\n            or (path.st_dev, path.st_ino) != self.log_identity\n            or stat.S_IMODE(info.st_mode) != 0o600 or stat.S_IMODE(path.st_mode) != 0o600\n            or info.st_uid != os.getuid() or info.st_gid != os.getgid()\n            or path.st_uid != os.getuid() or path.st_gid != os.getgid()\n            or not 0 <= info.st_size <= CAP):\n        raise ValueError("witness_unavailable")\n    prefix = os.pread(self.log_fd, info.st_size, 0)\n    after = os.fstat(self.log_fd)\n    if (len(prefix) != info.st_size or (after.st_dev, after.st_ino) != self.log_identity\n            or not info.st_size <= after.st_size <= CAP):\n        raise ValueError("witness_unavailable")\n    observed.append(("log", prefix))\n    return observed\n\n\ndef _capture(self, processor):\n    result = {"version": 1, "status": "unavailable", "snapshot_final": False,\n        "records": [], "truncated": False, "public_template_index_sha256": PUBLIC_TEMPLATE_INDEX_SHA}\n    try:\n        if time.monotonic() >= self.deadline:\n            return result\n        observed = _snapshot(self)\n        for stream, raw in observed:\n            # Only complete newline-terminated prefix records; incomplete tails\n            # cannot support a record or any last/final/causal claim.\n            complete = raw[:raw.rfind(b"\\n") + 1]\n            lines = complete.splitlines()\n            if raw and not raw.endswith(b"\\n"):\n                result["truncated"] = True\n            for ordinal, line in enumerate(lines):\n                if time.monotonic() >= self.deadline or ordinal >= 1024:\n                    result["truncated"] = True\n                    break\n                if _known(line, self, processor):\n                    continue\n                if len(result["records"]) == RECORD_CAP or len(line) > LINE_CAP:\n                    result["truncated"] = True\n                    continue\n                matches = []\n                try:\n                    value = line.decode("ascii", errors="strict")\n                except UnicodeError:\n                    value = None\n                if value is not None:\n                    for index, pattern in _PUBLIC:\n                        if time.monotonic() >= self.deadline:\n                            result["truncated"] = True\n                            break\n                        if pattern.fullmatch(value):\n                            if len(matches) == 8:\n                                result["truncated"] = True\n                                break\n                            matches.append(index)\n                result["records"].append({"stream": stream, "ordinal": ordinal,\n                    "line_bytes": len(line), "line_sha256": hashlib.sha256(line).hexdigest(),\n                    "shape": _shape(line), "public_template_indexes": matches})\n        result["status"] = "observed_prefix" if result["records"] else "no_unknown_complete_record_in_snapshot"\n    except BaseException:\n        # No exception text, path or error object escapes this diagnostic.\n        pass\n    return result\n\n\ndef _wrap(original, processor):\n    def diagnostic(self):\n        global _WITNESS\n        try:\n            return original(self)\n        except BaseException as error:\n            if _WITNESS is None and _unknown_raise(error):\n                _WITNESS = _capture(self, processor)\n            raise\n    return diagnostic\n\n\ndef install(processor):\n    """Future reviewed synthetic child only. Never call during preparation."""\n    if processor.__file__ != PROCESSOR_PATH:\n        raise ValueError("witness_pin_refused")\n    descriptor = os.open(PROCESSOR_PATH, os.O_RDONLY | os.O_NOFOLLOW)\n    try:\n        info = os.fstat(descriptor)\n        if (not stat.S_ISREG(info.st_mode) or info.st_uid != 0 or info.st_mode & 0o022\n                or not 0 < info.st_size <= 131072):\n            raise ValueError("witness_pin_refused")\n        raw = os.read(descriptor, 131073)\n        if len(raw) != info.st_size or hashlib.sha256(raw).hexdigest() != PROCESSOR_SHA:\n            raise ValueError("witness_pin_refused")\n    finally:\n        os.close(descriptor)\n    original = processor.OwnedClamd._diagnostics\n    if original.__code__.co_filename != PROCESSOR_PATH or original.__name__ != "_diagnostics":\n        raise ValueError("witness_pin_refused")\n    processor.OwnedClamd._diagnostics = _wrap(original, processor)\n    return original\n\n\ndef witness():\n    # This projection is error-only metadata, never successful acceptance.\n    return _WITNESS\n\n\n\nCLAMD_FAILURE_LINES = {\'_command\': (100, 103, 109, 112, 115), \'_native_file\': (202, 206, 214, 217), \'_installed_clamd\': (227, 231, 235), \'_clamd_definition\': (257, 261, 263), \'_zip_check\': (283,), \'_zip_admit\': (495, 497), \'_read_clamd_maps\': (850,), \'_remaining\': (528, 531), \'_diagnostics\': (565, 568, 572, 574, 581, 594, 596), \'_drain\': (609,), \'_socket\': (617, 625, 631), \'_exchange\': (639, 645, 652, 656), \'_start\': (667, 670, 677, 679, 682, 684, 686, 691, 695, 718, 722, 732), \'scan\': (737, 742, 749, 751, 761, 772, 777, 778, 780, 782), \'close\': (792, 806, 812, 815)}\nCLAMD_FAILURE_PHASES = frozenset((\'setup\', \'request\', \'source_scan\', \'resident_parser\', \'candidate_scan\', \'owner_cleanup\', \'ownership\', \'memory\', \'writer_version\', \'report\'))\nCLAMD_FAILURE_CLASSES = frozenset((\'InspectionDenied\', \'InspectionUnavailable\', \'MemoryError\', \'TimeoutExpired\', \'FileNotFoundError\', \'PermissionError\', \'TimeoutError\', \'UnicodeError\', \'OSError\', \'ValueError\', \'TypeError\', \'KeyError\', \'StopIteration\', \'RuntimeError\', \'SystemExit\', \'KeyboardInterrupt\', \'Unexpected\'))\nCLAMD_PHASE = "setup"\n\n\ndef failure_receipt(error):\n    """Fixed classes and literal owned raise sites only; never format errors."""\n    classes = ((processor.InspectionDenied, "InspectionDenied"),\n        (processor.InspectionUnavailable, "InspectionUnavailable"),\n        (MemoryError, "MemoryError"), (processor.subprocess.TimeoutExpired, "TimeoutExpired"),\n        (FileNotFoundError, "FileNotFoundError"), (PermissionError, "PermissionError"),\n        (TimeoutError, "TimeoutError"), (UnicodeError, "UnicodeError"), (OSError, "OSError"),\n        (ValueError, "ValueError"), (TypeError, "TypeError"), (KeyError, "KeyError"),\n        (StopIteration, "StopIteration"), (RuntimeError, "RuntimeError"),\n        (SystemExit, "SystemExit"), (KeyboardInterrupt, "KeyboardInterrupt"))\n    chain, seen, truncated = [], set(), False\n    while error is not None:\n        if id(error) in seen or len(chain) == 8:\n            truncated = True\n            break\n        seen.add(id(error))\n        category = next((label for kind, label in classes if isinstance(error, kind)), "Unexpected")\n        frames, trace_seen, walked = [], set(), 0\n        trace = error.__traceback__\n        while trace is not None:\n            if id(trace) in trace_seen or walked == 64:\n                truncated = True\n                break\n            trace_seen.add(id(trace))\n            walked += 1\n            code, line = trace.tb_frame.f_code, trace.tb_lineno\n            if (code.co_filename == "/opt/document/processor.py"\n                    and line in CLAMD_FAILURE_LINES.get(code.co_name, ())):\n                if len(frames) == 8:\n                    truncated = True\n                    break\n                frames.append({"function": code.co_name, "line": line})\n            trace = trace.tb_next\n        chain.append({"class": category, "frames": frames})\n        error = error.__context__\n    return {"version": 1, "status": "native_clamd_proof_unavailable",\n            "phase": CLAMD_PHASE if CLAMD_PHASE in CLAMD_FAILURE_PHASES else "setup",\n            "chain": chain, "truncated": truncated}\n\n\ndef require(value):\n    if not value:\n        raise ValueError("native_clamd_proof_unavailable")\n\n\ndef memory_value(name):\n    with open("/sys/fs/cgroup/" + name, "rb") as stream:\n        raw = stream.read(33)\n    require(0 < len(raw) <= 32 and raw.strip().isdigit())\n    return int(raw)\n\n\ndef main():\n    global CLAMD_PHASE\n    os.umask(0o077)\n    resource.setrlimit(resource.RLIMIT_CORE, (0, 0))\n    resource.setrlimit(resource.RLIMIT_CPU, (40, 40))\n    resource.setrlimit(resource.RLIMIT_AS, (2 * 1024**3, 2 * 1024**3))\n    resource.setrlimit(resource.RLIMIT_FSIZE, (32 * 1024**2, 32 * 1024**2))\n    resource.setrlimit(resource.RLIMIT_NOFILE, (64, 64))\n    deadline = time.monotonic() + 50\n    install(processor)  # Private witness admission is inside the unchanged work budget.\n    CLAMD_PHASE = "request"\n    with open("/tmp/request.json", "rb") as stream:\n        raw = stream.read(processor.MAX_REQUEST + 1)\n    require(0 < len(raw) <= processor.MAX_REQUEST)\n    request = json.loads(raw, object_pairs_hook=processor._pairs)\n    require(set(request) == {"case", "source", "candidate", "expected", "resident_parser",\n                            "candidate_sha256", "candidate_size_bytes", "edits"})\n    require(request["expected"] in {"clean", "source_refused", "candidate_refused", "candidate_nonclean"})\n    require(type(request["resident_parser"]) is bool)\n    source = base64.b64decode(request["source"], validate=True)\n    candidate = base64.b64decode(request["candidate"], validate=True)\n    require(0 < len(source) <= processor.MAX_FILE and 0 < len(candidate) <= processor.MAX_FILE)\n    require(type(request["candidate_size_bytes"]) is int and len(candidate) == request["candidate_size_bytes"]\n            and hashlib.sha256(candidate).hexdigest() == request["candidate_sha256"])\n    outcome = "clean"\n    generated = False\n    with tempfile.TemporaryDirectory(prefix="clamd-proof-", dir="/tmp") as temporary:\n        directory = Path(temporary)\n        directory.chmod(0o700)\n        original = directory / "original.docx"\n        target = directory / "candidate.docx"\n        original.write_bytes(source)\n        original.chmod(0o600)\n        owner = processor.OwnedClamd(directory, deadline)\n        source_metadata = candidate_metadata = None\n        resident_current = None\n        with owner:\n            try:\n                CLAMD_PHASE = "source_scan"\n                source_metadata = owner.scan(original, source)\n            except processor.InspectionDenied:\n                outcome = "source_refused"\n            if outcome == "clean":\n                if request["resident_parser"]:\n                    # Actual source parsing and candidate regeneration while the\n                    # genuine engine remains resident, never synthetic telemetry.\n                    CLAMD_PHASE = "resident_parser"\n                    processor.validate_docx(source)\n                    from app.services.resume_layout import _apply_docx\n                    require(type(request["edits"]) is list and len(request["edits"]) == 1)\n                    candidate = _apply_docx(source, request["edits"])\n                    require(len(candidate) == request["candidate_size_bytes"]\n                            and hashlib.sha256(candidate).hexdigest() == request["candidate_sha256"])\n                    processor.validate_docx(candidate)\n                    generated = True\n                    resident_current = memory_value("memory.current")\n                target.write_bytes(candidate)\n                target.chmod(0o600)\n                try:\n                    CLAMD_PHASE = "candidate_scan"\n                    candidate_metadata = owner.scan(target, candidate)\n                except processor.InspectionDenied:\n                    outcome = "candidate_refused"\n                except processor.InspectionUnavailable:\n                    require(request["expected"] == "candidate_nonclean")\n                    outcome = "candidate_nonclean"\n            require(original.read_bytes() == source)\n            if target.exists():\n                require(target.read_bytes() == candidate)\n            CLAMD_PHASE = "owner_cleanup"\n        CLAMD_PHASE = "ownership"\n        require(owner.stopped and not owner.forced and owner.child is not None and owner.child.returncode == 0)\n        require(owner.peer is not None and owner.peer["pid"] == owner.child.pid)\n        require(owner.peer["uid"] == os.getuid() and owner.peer["gid"] == os.getgid())\n        require(len(owner.signatures) == 1 and owner.signatures[0] > 0)\n        require(not owner.log_path.exists())\n        require(time.monotonic() < deadline)\n        actual_expected = "candidate_refused" if request["expected"] == "candidate_nonclean" and outcome == "candidate_refused" else request["expected"]\n        require(outcome == actual_expected)\n        if request["expected"] == "clean":\n            require(source_metadata == candidate_metadata and len(owner.scan_bindings) == 2)\n        else:\n            require(outcome != "clean")\n        CLAMD_PHASE = "memory"\n        maximum, peak = memory_value("memory.max"), memory_value("memory.peak")\n        require(maximum == 2 * 1024**3 and 0 < peak <= maximum)\n        if request["resident_parser"]:\n            require(generated and type(resident_current) is int and 0 < resident_current <= maximum)\n        # A native Writer version operation starts only after the verified stop.\n        CLAMD_PHASE = "writer_version"\n        import docx_renderer\n        renderer_version = docx_renderer._version(deadline=deadline)\n        response = {"case": request["case"], "outcome": outcome, "native_installation": owner.manifest,\n                    "owned_peer": owner.peer, "loaded_signatures": owner.signatures[0],\n                    "scans": owner.scan_bindings, "daemon_returncode": owner.child.returncode,\n                    "protocol_scans": owner.protocol_bindings, "instream_attempts": owner.instream_attempts,\n                    "zip_admission_failure": owner.zip_admission_failure,\n                    "daemon_stopped_before_writer": owner.stopped, "native_log_removed": not owner.log_path.exists(),\n                    "renderer_version_after_stop": renderer_version,\n                    "resident_parser_exercised": generated, "resident_parser_cgroup_current_bytes": resident_current,\n                    "cgroup_memory_max_bytes": maximum, "cgroup_memory_peak_bytes": peak,\n                    "performance_proven": False, "cloud_run_isolation_proven": False}\n        CLAMD_PHASE = "report"\n        encoded = json.dumps(response, separators=(",", ":"), allow_nan=False).encode()\n        require(len(encoded) <= 16384)\n        sys.stdout.buffer.write(encoded + b"\\n")\n\n\nNATIVE_REQUIREMENT_SITES = {330: \'shutdown\', 331: \'pid\', 332: \'uid\', 333: \'signatures\', 334: \'log_cleanup\', 335: \'deadline\', 337: \'case_outcome\', 339: \'clean_metadata\', 341: \'negative_nonclean\'}\nNATIVE_REQUIREMENTS = frozenset((\'shutdown\', \'pid\', \'uid\', \'signatures\', \'log_cleanup\', \'deadline\', \'case_outcome\', \'clean_metadata\', \'negative_nonclean\'))\n\n\ndef native_requirement_receipt(error):\n    """Only exact own code frames at fixed require sites; no locals or values."""\n    if type(error) is not ValueError:\n        return None\n    trace, previous, seen = error.__traceback__, None, set()\n    for _ in range(64):\n        if trace is None or id(trace) in seen:\n            return None\n        seen.add(id(trace))\n        code = trace.tb_frame.f_code\n        if (code is require.__code__ and previous is not None\n                and previous[0] is main.__code__ and trace.tb_next is None):\n            label = NATIVE_REQUIREMENT_SITES.get(previous[1])\n            return label if label in NATIVE_REQUIREMENTS else None\n        previous = code, trace.tb_lineno\n        trace = trace.tb_next\n    return None\n\n\nif __name__ == "__main__":\n    try:\n        main()\n    except BaseException as error:\n        diagnostic = failure_receipt(error)\n        requirement = native_requirement_receipt(error)\n        if requirement is not None:\n            diagnostic["native_requirement"] = requirement\n        startup = witness()\n        if startup is not None:\n            diagnostic["startup_log_witness"] = startup\n        print(json.dumps(diagnostic, separators=(",", ":"), allow_nan=False))\n        raise SystemExit(1) from None\n'
DOCX_LAYOUT_PROGRAM = '''import json,pathlib,subprocess,sys
payload=pathlib.Path("/tmp/request.json").read_bytes()
p=subprocess.run([sys.executable,"-I","/opt/document/processor.py","--docx-layout-stdin"],
 input=payload,capture_output=True,timeout=55)
if len(p.stdout)>32768 or p.stderr: raise SystemExit(3)
sys.stdout.buffer.write(p.stdout);raise SystemExit(p.returncode)
'''
DOCX_LINK_PROGRAM = '''import io,json,math,pathlib,re,sys,tempfile,time
from zipfile import ZIP_DEFLATED,ZipFile
sys.path.insert(0,"/opt/document")
import processor,docx_renderer
from lxml import etree
from pypdf import PdfReader
import pdfplumber
URI="https://example.test/portfolio"
W="{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
R="{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
T="{urn:oasis:names:tc:opendocument:xmlns:text:1.0}"
X="{http://www.w3.org/1999/xlink}href"
def xml(data):
 if len(data)>10*1024**2: raise ValueError()
 tree=etree.fromstring(data,etree.XMLParser(resolve_entities=False,no_network=True,load_dtd=False,huge_tree=False))
 if tree.getroottree().docinfo.doctype: raise ValueError()
 nodes=list(tree.iter())
 if len(nodes)>100000 or any(len(list(n.iterancestors()))>64 for n in nodes): raise ValueError()
 return tree
def pdf_observe(pdf):
 processor.validate_pdf(pdf)
 reader=PdfReader(io.BytesIO(pdf),strict=True)
 annotations=[]
 text_present=False
 for page in reader.pages:
  text_present=text_present or "Portfolio" in (page.extract_text() or "")
  for item in page.get("/Annots",[]):
   item=item.get_object()
   if item.get("/Subtype")=="/Link": annotations.append(item)
 targets=[]
 for item in annotations:
  action=item.get("/A")
  action=action.get_object() if action is not None else {}
  targets.append(action.get("/S")=="/URI" and action.get("/URI")==URI)
 occurrences=0;in_page=True;total=0
 with pdfplumber.open(io.BytesIO(pdf)) as document:
  for page in document.pages:
   chars=page.chars;total+=len(chars)
   if total>200000: raise ValueError()
   glyphs=[(letter,char) for char in chars for letter in char["text"]]
   if len(glyphs)>200000: raise ValueError()
   for match in re.finditer("Portfolio","".join(letter for letter,char in glyphs)):
    occurrences+=1
    for letter,char in glyphs[match.start():match.end()]:
     box=tuple(float(char[k]) for k in ("x0","top","x1","bottom"))
     in_page=in_page and all(math.isfinite(n) for n in box) and 0<=box[0]<=box[2]<=page.width and 0<=box[1]<=box[3]<=page.height
 if not 0<=len(annotations)<=100 or not 0<=occurrences<=100: raise ValueError()
 return len(annotations),bool(targets) and all(targets),text_present,occurrences,occurrences>0 and in_page

def styled_copy(source):
 with ZipFile(io.BytesIO(source)) as archive:
  infos=archive.infolist();parts={entry.filename:archive.read(entry) for entry in infos}
 expected_parts={"[Content_Types].xml","_rels/.rels","docProps/core.xml","docProps/app.xml","docProps/thumbnail.jpeg",
  "word/document.xml","word/_rels/document.xml.rels","word/styles.xml","word/stylesWithEffects.xml","word/settings.xml",
  "word/webSettings.xml","word/fontTable.xml","word/theme/theme1.xml","word/numbering.xml","customXml/item1.xml",
  "customXml/_rels/item1.xml.rels","customXml/itemProps1.xml"}
 if set(parts)!=expected_parts: raise ValueError()
 document=xml(parts["word/document.xml"]);styles=xml(parts["word/styles.xml"])
 body=document.find(W+"body")
 if body is None or [n.tag for n in body]!=[W+"p"]*6+[W+"sectPr"]: raise ValueError()
 paragraphs=list(body)[:6]
 fixed=["Taylor Example","Built scalable systems with Python and SQL.","Reduced API latency by 40% through focused profiling.",
  "Portfolio","","Original education and second-page running content."]
 if ["".join(n.text or "" for n in p.iter(W+"t")) for p in paragraphs]!=fixed: raise ValueError()
 if ["".join(n.text or "" for n in r.iter(W+"t")) for r in paragraphs[1]]!=["Built scalable systems with ","Python"," and ","SQL","."]: raise ValueError()
 if paragraphs[1][1].find(W+"rPr/"+W+"b") is None or paragraphs[1][3].find(W+"rPr/"+W+"i") is None: raise ValueError()
 link=paragraphs[3].find(W+"hyperlink")
 if link is None or len(paragraphs[3])!=1 or set(link.attrib)!={R+"id"} or len(link)!=1 or link[0].tag!=W+"r" or list(link[0].attrib) or [n.tag for n in link[0]]!=[W+"t"] or link[0][0].text!="Portfolio": raise ValueError()
 if len(list(document.iter(W+"hyperlink")))!=1 or paragraphs[4].find(W+"r/"+W+"br").get(W+"type")!="page": raise ValueError()
 if styles.tag!=W+"styles" or any(n.get(W+"styleId")=="Hyperlink" for n in styles): raise ValueError()
 default=[n for n in styles if n.get(W+"styleId")=="DefaultParagraphFont" and n.get(W+"type")=="character"]
 normal=[n for n in styles if n.get(W+"styleId")=="Normal"]
 if len(default)!=1 or len(normal)!=1 or normal[0].find(W+"rPr/"+W+"rFonts").get(W+"ascii")!="Cambria" or normal[0].find(W+"rPr/"+W+"sz").get(W+"val")!="22": raise ValueError()
 style=etree.SubElement(styles,W+"style",{W+"type":"character",W+"styleId":"Hyperlink"})
 etree.SubElement(style,W+"name",{W+"val":"Hyperlink"});etree.SubElement(style,W+"basedOn",{W+"val":"DefaultParagraphFont"})
 properties=etree.Element(W+"rPr");etree.SubElement(properties,W+"rStyle",{W+"val":"Hyperlink"});link[0].insert(0,properties)
 parts["word/document.xml"]=etree.tostring(document,xml_declaration=True,encoding="UTF-8",standalone=True)
 parts["word/styles.xml"]=etree.tostring(styles,xml_declaration=True,encoding="UTF-8",standalone=True)
 buffer=io.BytesIO()
 with ZipFile(buffer,"w",compression=ZIP_DEFLATED) as archive:
  for entry in infos:archive.writestr(entry,parts[entry.filename])
 result=buffer.getvalue()
 if not 0<len(result)<=5*1024**2: raise ValueError()
 return result

def main():
 deadline=time.monotonic()+55
 with pathlib.Path("/tmp/request.json").open("rb") as stream: payload=stream.read(processor.MAX_REQUEST+1)
 source,decoded=processor.decode_docx_request(payload)
 with pathlib.Path("/tmp/diagnostic-binding.json").open("rb") as stream: expected_raw=stream.read(2049)
 if len(expected_raw)>2048: raise ValueError()
 expected=json.loads(expected_raw)
 binding=decoded["binding"]
 if type(expected) is not dict or set(expected)!=set(binding) or any(type(expected[k]) is not type(v) or expected[k]!=v for k,v in binding.items()): raise ValueError()
 with tempfile.TemporaryDirectory(prefix="docx-link-observation-",dir="/tmp") as temporary:
  directory=pathlib.Path(temporary);directory.chmod(0o700)
  original=directory/"input.docx";original.write_bytes(source);original.chmod(0o600)
  processor.scan_original(original);processor.validate_docx(source)
  if original.read_bytes()!=source or time.monotonic()>=deadline: raise ValueError()
  with ZipFile(io.BytesIO(source)) as archive:
   native=list(xml(archive.read("word/document.xml")).iter(W+"hyperlink"))
   relationships=list(xml(archive.read("word/_rels/document.xml.rels")))
  native_match=len(native)==1
  for link in native:
   matches=[r for r in relationships if r.get("Id")==link.get(R+"id")]
   native_match=native_match and len(matches)==1 and matches[0].get("Type")==R[1:-1]+"/hyperlink" and matches[0].get("TargetMode")=="External" and matches[0].get("Target")==URI
  if not native_match: raise ValueError()
  pdf=docx_renderer._render(source,directory/"pdf-render",deadline=min(deadline,time.monotonic()+10))
  output=directory/"flat";output.mkdir(mode=0o700)
  docx_renderer._run([docx_renderer.SOFFICE,"--headless","--nologo","--nodefault","--norestore",
   "-env:UserInstallation="+(directory/"flat-profile").as_uri(),"--convert-to",
   "fodt:OpenDocument Text Flat XML","--outdir",str(output),str(original)],deadline=min(deadline,time.monotonic()+10))
  flat=output/"input.fodt"
  if original.read_bytes()!=source or sorted(p.name for p in output.iterdir())!=["input.fodt"] or flat.is_symlink() or not flat.is_file() or not 0<flat.stat().st_size<=10*1024**2: raise ValueError()
  with flat.open("rb") as stream: imported_tree=xml(stream.read(10*1024**2+1))
  O="{urn:oasis:names:tc:opendocument:xmlns:office:1.0}"
  if imported_tree.tag!=O+"document" or imported_tree.get(O+"mimetype")!="application/vnd.oasis.opendocument.text": raise ValueError()
  imported=list(imported_tree.iter(T+"a"))
  link_count,uri_match,label_text,label_glyphs,label_in_page=pdf_observe(pdf)
  counts=(len(native),len(imported),link_count)
  if any(not 0<=n<=100 for n in counts): raise ValueError()
  variant=styled_copy(source);processor.validate_docx(variant)
  styled=directory/"styled.docx";styled.write_bytes(variant);styled.chmod(0o600)
  if deadline-time.monotonic()<35: raise ValueError()
  processor.scan_original(styled)
  if styled.read_bytes()!=variant or original.read_bytes()!=source or time.monotonic()>=deadline: raise ValueError()
  styled_pdf=docx_renderer._render(variant,directory/"styled-pdf-render",deadline=min(deadline,time.monotonic()+10))
  styled_count,styled_uri,styled_text,styled_glyphs,styled_in_page=pdf_observe(styled_pdf)
  if styled.read_bytes()!=variant or original.read_bytes()!=source or time.monotonic()>=deadline: raise ValueError()
  print(json.dumps(dict(zip(("native_hyperlinks","imported_hyperlinks","pdf_link_annotations","native_uri_match","imported_uri_match","pdf_uri_match",
   "imported_label_match","pdf_label_text_present","pdf_label_glyph_occurrences","pdf_label_in_page",
   "styled_pdf_link_annotations","styled_pdf_uri_match","styled_fixed_label_present"),
   (*counts,native_match,bool(imported) and all(n.get(X)==URI for n in imported),uri_match,
    len(imported)==1 and "".join(imported[0].itertext())=="Portfolio",label_text,label_glyphs,label_in_page,
    styled_count,styled_uri,styled_text and styled_glyphs>0 and styled_in_page),strict=True))))
try: main()
except Exception:
 print(json.dumps({"status":"unavailable"}));raise SystemExit(1) from None
'''


class DocumentImageError(RuntimeError):
    def __init__(self):
        super().__init__(ERROR)


class _ClamdFailure(NamedTuple):
    case: str
    phase: str
    chain: tuple
    truncated: bool


class _ClamdNativeRequirement(NamedTuple):
    case: str
    label: str


CLAMD_NATIVE_REQUIREMENTS = frozenset(('shutdown', 'pid', 'uid', 'signatures', 'log_cleanup', 'deadline', 'case_outcome', 'clean_metadata', 'negative_nonclean'))


def _native_requirement_value(value, phase, chain):
    if (type(value) is not str or value not in CLAMD_NATIVE_REQUIREMENTS or phase != "ownership"
            or type(chain) is not tuple or not chain or type(chain[0]) is not tuple
            or len(chain[0]) != 2 or chain[0][0] != "ValueError"):
        raise DocumentImageError()
    return value



def _clamd_diagnostic_pairs(pairs):
    result = {}
    for key, value in pairs:
        if key in result:
            raise DocumentImageError()
        result[key] = value
    return result


def _clamd_diagnostic_values(value):
    if (type(value) is not dict or set(value) != {"version", "status", "phase", "chain", "truncated"}
            or type(value["version"]) is not int or value["version"] != 1
            or value["status"] != "native_clamd_proof_unavailable"
            or type(value["phase"]) is not str or value["phase"] not in CLAMD_FAILURE_PHASES
            or type(value["truncated"]) is not bool or type(value["chain"]) is not list
            or not 1 <= len(value["chain"]) <= 8):
        raise DocumentImageError()
    chain = []
    for row in value["chain"]:
        if (type(row) is not dict or set(row) != {"class", "frames"}
                or type(row["class"]) is not str or row["class"] not in CLAMD_FAILURE_CLASSES
                or type(row["frames"]) is not list or len(row["frames"]) > 8):
            raise DocumentImageError()
        frames = []
        for frame in row["frames"]:
            if (type(frame) is not dict or set(frame) != {"function", "line"}
                    or type(frame["function"]) is not str or frame["function"] not in CLAMD_FAILURE_LINES
                    or type(frame["line"]) is not int
                    or frame["line"] not in CLAMD_FAILURE_LINES[frame["function"]]):
                raise DocumentImageError()
            frames.append((frame["function"], frame["line"]))
        chain.append((row["class"], tuple(frames)))
    return value["phase"], tuple(chain), value["truncated"]


def _clamd_diagnostic_object(phase, chain, truncated):
    return {"version": 1, "status": "native_clamd_proof_unavailable", "phase": phase,
            "chain": [{"class": category, "frames": [{"function": function, "line": line}
                for function, line in frames]} for category, frames in chain], "truncated": truncated}


class _ClamdStartupObservation(NamedTuple):
    case: str
    container_calls: int
    elapsed: float
    exit_code: int | None
    values: tuple | None
    native_requirement: str | None = None


def _clamd_startup_values(value):
    if (type(value) is not dict or set(value) != {"version", "status", "snapshot_final", "records", "truncated", "public_template_index_sha256"}
            or type(value["version"]) is not int or value["version"] != 1
            or type(value["status"]) is not str or value["status"] not in CLAMD_STARTUP_STATUSES
            or value["snapshot_final"] is not False or type(value["truncated"]) is not bool
            or type(value["public_template_index_sha256"]) is not str
            or value["public_template_index_sha256"] != CLAMD_STARTUP_INDEX_SHA
            or type(value["records"]) is not list or len(value["records"]) > 4):
        raise DocumentImageError()
    records = []
    for row in value["records"]:
        if (type(row) is not dict or set(row) != {"stream", "ordinal", "line_bytes", "line_sha256", "shape", "public_template_indexes"}
                or type(row["stream"]) is not str or row["stream"] not in {"stdout", "log"}
                or type(row["ordinal"]) is not int or not 0 <= row["ordinal"] <= 1023
                or type(row["line_bytes"]) is not int or not 0 <= row["line_bytes"] <= 2048
                or type(row["line_sha256"]) is not str or len(row["line_sha256"]) != 64
                or any(char not in "0123456789abcdef" for char in row["line_sha256"])
                or type(row["shape"]) is not str or row["shape"] not in CLAMD_STARTUP_SHAPES
                or type(row["public_template_indexes"]) is not list or len(row["public_template_indexes"]) > 8
                or any(type(index) is not int or not 0 <= index <= 103 for index in row["public_template_indexes"])
                or len(set(row["public_template_indexes"])) != len(row["public_template_indexes"])):
            raise DocumentImageError()
        records.append((row["stream"], row["ordinal"], row["line_bytes"], row["line_sha256"], row["shape"], tuple(row["public_template_indexes"])))
    if value["status"] == "observed_prefix" and not records:
        raise DocumentImageError()
    return value["status"], tuple(records), value["truncated"]


def _clamd_startup_object(status, records, truncated):
    return {"version": 1, "status": status, "snapshot_final": False, "records": [
        {"stream": stream, "ordinal": ordinal, "line_bytes": size, "line_sha256": digest,
         "shape": shape, "public_template_indexes": list(indexes)}
        for stream, ordinal, size, digest, shape, indexes in records],
        "truncated": truncated, "public_template_index_sha256": CLAMD_STARTUP_INDEX_SHA}


def _diagnose_clamd_startup(docker: list[str], image: str, platform: str, request: Path,
                            *, case: str) -> None:
    # Error-only derivative: the original failure and its acceptance are retained.
    diagnostic = _FAILURE.get()
    if (type(diagnostic) is not dict or case != "clamd_two_clean"
            or diagnostic.get("stage") != "case_validation" or diagnostic.get("case") != case
            or type(diagnostic.get("exit_code")) is not int or diagnostic["exit_code"] != 1
            or diagnostic.get("clamd_startup_observation") is not None):
        return
    original = diagnostic.get("clamd_failure")
    if (type(original) is not _ClamdFailure or original.case != case or original.phase != "source_scan"
            or not any(("_diagnostics", 594) in frames for _, frames in original.chain)):
        return
    # This marker precedes the only call; uncertainty never authorizes a retry.
    diagnostic["clamd_startup_observation"] = _ClamdStartupObservation(case, 1, 0.0, None, None)
    started = time.monotonic()
    child_code = None
    values = None
    requirement = None
    try:
        actual_code, raw = container(docker, image, platform, ["-c", CLAMD_STARTUP_DIAGNOSTIC_PROGRAM],
                                    [(request, "/tmp/request.json")])
        if type(actual_code) is int and 0 <= actual_code <= 255:
            child_code = actual_code
        if (child_code == 1 and type(raw) is bytes and 0 < len(raw) <= 8192 and raw.endswith(b"\n")):
            observed = json.loads(raw, object_pairs_hook=_clamd_diagnostic_pairs)
            required = {"version", "status", "phase", "chain", "truncated"}
            if (type(observed) is not dict or not required <= set(observed)
                    or not set(observed) <= required | {"startup_log_witness", "native_requirement"}):
                raise DocumentImageError()
            core = {key: item for key, item in observed.items()
                    if key not in {"startup_log_witness", "native_requirement"}}
            phase, chain, truncated = _clamd_diagnostic_values(core)
            requirement = _native_requirement_value(observed["native_requirement"], phase, chain) if "native_requirement" in observed else None
            witness = _clamd_startup_values(observed["startup_log_witness"]) if "startup_log_witness" in observed else None
            if raw != (json.dumps(observed, separators=(",", ":"), allow_nan=False) + "\n").encode("ascii"):
                raise DocumentImageError()
            values = phase, chain, truncated, witness
    except Exception:
        pass
    finally:
        # Do not call _progress: it would clear the genuine original failure.
        diagnostic["clamd_startup_observation"] = _ClamdStartupObservation(
            case, 1, round(time.monotonic() - started, 6), child_code, values, requirement)



def _note_clamd_failure(raw: bytes, code: int, *, case: str) -> None:
    diagnostic = _FAILURE.get()
    if type(diagnostic) is not dict:
        return
    diagnostic.pop("clamd_failure", None)
    diagnostic.pop("clamd_native_requirement", None)
    if (type(code) is not int or code != 1 or type(case) is not str
            or case not in _FAILURE_CASES or not case.startswith("clamd_")
            or type(raw) is not bytes or not 0 < len(raw) <= 8192 or not raw.endswith(b"\n")):
        return
    try:
        value = json.loads(raw, object_pairs_hook=_clamd_diagnostic_pairs)
        if type(value) is not dict or set(value) not in (
                {"version", "status", "phase", "chain", "truncated"},
                {"version", "status", "phase", "chain", "truncated", "native_requirement"}):
            return
        core = {key: item for key, item in value.items() if key != "native_requirement"}
        phase, chain, truncated = _clamd_diagnostic_values(core)
        label = _native_requirement_value(value["native_requirement"], phase, chain) if "native_requirement" in value else None
        if raw != (json.dumps(value, separators=(",", ":"), allow_nan=False) + "\n").encode("ascii"):
            return
    except Exception:
        return
    diagnostic["clamd_failure"] = _ClamdFailure(case, phase, chain, truncated)
    if label is not None:
        diagnostic["clamd_native_requirement"] = _ClamdNativeRequirement(case, label)


def _progress(stage: str, *, case: str | None = None, exit_code: int | None = None,
              completed_cases: int | None = None) -> None:
    diagnostic = _FAILURE.get()
    if type(diagnostic) is dict:
        diagnostic.pop("clamd_failure", None)
        diagnostic.pop("clamd_native_requirement", None)
        diagnostic.pop("clamd_startup_observation", None)
        diagnostic.pop("docx_refusal", None)
        diagnostic.pop("docx_fit_mismatch", None)
        diagnostic.pop("docx_link_observation", None)
        diagnostic.update(stage=stage, case=case, exit_code=exit_code)
        if completed_cases is not None:
            diagnostic["completed_cases"] = completed_cases


def _note_docx_refusal(result: object, binding: dict, code: int, *, case: str) -> None:
    """Keep only a typed finite projection of the exact bound exit1 refusal."""
    diagnostic = _FAILURE.get()
    if type(diagnostic) is not dict:
        return
    diagnostic.pop("docx_refusal", None)
    if (type(diagnostic.get("stage")) is not str or diagnostic["stage"] != "case_validation"
            or type(diagnostic.get("case")) is not str or type(case) is not str or diagnostic["case"] != case
            or type(diagnostic.get("exit_code")) is not int or diagnostic["exit_code"] != 1
            or case not in {"styled_multipage_fit", "rendered_reflow_refusal"} or type(code) is not int or code != 1
            or type(binding) is not dict or any(type(key) is not str for key in binding)
            or set(binding) != _DOCX_BINDING_KEYS
            or type(binding["version"]) is not int or binding["version"] != 1
            or type(binding["operation"]) is not str or binding["operation"] != "docx_layout"
            or type(binding["source_format"]) is not str or binding["source_format"] != "docx"
            or any(type(binding[key]) is not str or not re.fullmatch(r"[a-f0-9]{64}", binding[key])
                   for key in ("sha256", "candidate_sha256", "edits_sha256"))
            or any(type(binding[key]) is not int or not 0 < binding[key] <= 5242880
                   for key in ("size_bytes", "candidate_size_bytes"))
            or type(result) is not dict or any(type(key) is not str for key in result)
            or set(result) != _DOCX_BINDING_KEYS | {"error", "reason", "unit_id"}
            or any(type(result.get(key)) is not type(value) or result.get(key) != value for key, value in binding.items())
            or type(result["error"]) is not str or result["error"] != "docx_layout_refused"
            or type(result["reason"]) is not str or result["reason"] not in _DOCX_REASONS):
        return
    diagnostic["docx_refusal"] = _DocxRefusal(case, result["error"], result["reason"])


def _note_docx_fit_mismatch(result: object, binding: dict, code: int, *, case: str) -> None:
    """Project only existing failed fit-check names after an exact bound exit0."""
    diagnostic = _FAILURE.get()
    if type(diagnostic) is not dict:
        return
    diagnostic.pop("docx_fit_mismatch", None)
    if (type(diagnostic.get("stage")) is not str or diagnostic["stage"] != "case_validation"
            or type(diagnostic.get("case")) is not str or type(case) is not str or diagnostic["case"] != case
            or type(diagnostic.get("exit_code")) is not int or diagnostic["exit_code"] != 0
            or case != "styled_multipage_fit" or type(code) is not int or code != 0
            or type(binding) is not dict or any(type(key) is not str for key in binding)
            or set(binding) != _DOCX_BINDING_KEYS
            or type(binding["version"]) is not int or binding["version"] != 1
            or type(binding["operation"]) is not str or binding["operation"] != "docx_layout"
            or type(binding["source_format"]) is not str or binding["source_format"] != "docx"
            or any(type(binding[key]) is not str or not re.fullmatch(r"[a-f0-9]{64}", binding[key])
                   for key in ("sha256", "candidate_sha256", "edits_sha256"))
            or any(type(binding[key]) is not int or not 0 < binding[key] <= 5242880
                   for key in ("size_bytes", "candidate_size_bytes"))
            or type(result) is not dict or any(type(key) is not str for key in result)
            or set(result) != _DOCX_BINDING_KEYS | {"layout", "scan", "candidate_scan"}
            or any(type(result.get(key)) is not type(value) or result.get(key) != value for key, value in binding.items())
            or type(result["layout"]) is not dict or type(result["scan"]) is not dict):
        return
    layout = result["layout"]
    try:
        failed = (layout.get("validator") != "hirewiz-docx-layout-v1", layout.get("pages") != 2,
                  layout.get("unchanged_links") != 1, layout.get("outside_changed_regions_pixel_delta") != 0,
                  result["scan"].get("engine") != "ClamAV", result["candidate_scan"] != result["scan"])
        checks = tuple(name for name, mismatch in zip(_DOCX_FIT_CHECKS, failed, strict=True) if mismatch)
    except Exception:
        return
    if checks:
        pages, links = layout.get("pages"), layout.get("unchanged_links")
        diagnostic["docx_fit_mismatch"] = _DocxFitMismatch(case, checks,
            pages if type(pages) is int and 0 <= pages <= 20 else None,
            links if type(links) is int and 0 <= links <= 100 else None)


def _failure_report() -> dict:
    # Arbitrary context/native text cannot escape. A bound refusal contributes
    # only the independently allowlisted typed projection, never its body/unit.
    stored = _FAILURE.get()
    diagnostic = stored if type(stored) is dict else {}
    stage, case = diagnostic.get("stage"), diagnostic.get("case")
    code, completed = diagnostic.get("exit_code"), diagnostic.get("completed_cases")
    result: dict[str, object] = {"error": ERROR,
            "stage": stage if type(stage) is str and stage in _FAILURE_STAGES else "setup",
            "case": case if type(case) is str and case in _FAILURE_CASES else None,
            "exit_code": code if type(code) is int and 0 <= code <= 255 else None,
            "completed_cases": completed if type(completed) is int and 0 <= completed <= len(_FAILURE_CASES) else 0}
    refusal = diagnostic.get("docx_refusal")
    if (result["stage"] == "case_validation" and result["case"] in {"styled_multipage_fit", "rendered_reflow_refusal"}
            and result["exit_code"] == 1
            and type(refusal) is _DocxRefusal and type(refusal.case) is str and refusal.case == result["case"]
            and type(refusal.error) is str and refusal.error == "docx_layout_refused"
            and type(refusal.reason) is str and refusal.reason in _DOCX_REASONS):
        result.update(processor_error=refusal.error, processor_reason=refusal.reason)
    fit = diagnostic.get("docx_fit_mismatch")
    if (result["stage"] == "case_validation" and result["case"] == "styled_multipage_fit" and result["exit_code"] == 0
            and type(fit) is _DocxFitMismatch and type(fit.case) is str and fit.case == result["case"]
            and type(fit.failed_checks) is tuple and 0 < len(fit.failed_checks) <= len(_DOCX_FIT_CHECKS)
            and all(type(name) is str and name in _DOCX_FIT_CHECKS for name in fit.failed_checks)
            and len(set(fit.failed_checks)) == len(fit.failed_checks)):
        result["fit_failed_checks"] = list(fit.failed_checks)
        if type(fit.pages) is int and 0 <= fit.pages <= 20:
            result["fit_pages"] = fit.pages
        if type(fit.links) is int and 0 <= fit.links <= 100:
            result["fit_links"] = fit.links
    observation = diagnostic.get("docx_link_observation")
    if (result.get("fit_pages") == 2 and result.get("fit_links") == 0 and type(observation) is _DocxLinkObservation
            and type(observation.container_calls) is int and observation.container_calls in {0, 1}
            and type(observation.elapsed) is float and math.isfinite(observation.elapsed)
            and 0 <= observation.elapsed <= 150):
        result.update(link_diagnostic_container_calls=observation.container_calls, link_diagnostic_seconds=observation.elapsed,
                      link_diagnostic_status="unavailable")
        values = observation.values
        if (observation.container_calls == 1 and type(values) is tuple and len(values) == len(_LINK_FIELDS)
                and all(type(n) is int and 0 <= n <= 100 for n in (values[i] for i in _LINK_COUNT_INDEXES))
                and all(type(flag) is bool for flag in (values[i] for i in _LINK_BOOL_INDEXES))):
            result.update(link_diagnostic_status="observed", link_diagnostic=dict(zip(_LINK_FIELDS, values, strict=True)))
    clamd = diagnostic.get("clamd_failure")
    if (result["stage"] == "case_validation" and result["exit_code"] == 1
            and type(clamd) is _ClamdFailure and clamd.case == result["case"]
            and type(clamd.case) is str and clamd.case.startswith("clamd_")
            and clamd.case in _FAILURE_CASES):
        try:
            value = _clamd_diagnostic_object(clamd.phase, clamd.chain, clamd.truncated)
            _clamd_diagnostic_values(value)
        except Exception:
            pass
        else:
            result["clamd_failure"] = value
    requirement = diagnostic.get("clamd_native_requirement")
    if (requirement is not None and "clamd_failure" in result and type(requirement) is _ClamdNativeRequirement
            and requirement.case == result["case"] and type(clamd) is _ClamdFailure):
        try:
            result["native_requirement"] = _native_requirement_value(requirement.label, clamd.phase, clamd.chain)
        except Exception:
            pass
    startup = diagnostic.get("clamd_startup_observation")
    if (startup is not None and "clamd_failure" in result and result["case"] == "clamd_two_clean"
            and type(startup) is _ClamdStartupObservation and startup.case == result["case"]
            and type(startup.container_calls) is int and startup.container_calls == 1
            and type(startup.elapsed) is float and math.isfinite(startup.elapsed) and 0 <= startup.elapsed <= 150
            and (startup.exit_code is None or type(startup.exit_code) is int and 0 <= startup.exit_code <= 255)):
        projection = {"original_container_calls": 1, "diagnostic_container_calls": 1,
            "combined_case_container_calls": 2, "elapsed_seconds": startup.elapsed,
            "exit_code": startup.exit_code, "status": "returned_success" if startup.exit_code == 0 else "unavailable",
            "original_failure_retained": True, "acceptance_evidence": False}
        if startup.exit_code == 1 and type(startup.values) is tuple and len(startup.values) == 4:
            try:
                phase, chain, truncated, witness = startup.values
                core = _clamd_diagnostic_object(phase, chain, truncated)
                _clamd_diagnostic_values(core)
                requirement_label = _native_requirement_value(startup.native_requirement, phase, chain) if startup.native_requirement is not None else None
                observed = None
                if witness is not None:
                    if type(witness) is not tuple or len(witness) != 3:
                        raise DocumentImageError()
                    observed = _clamd_startup_object(*witness)
                    _clamd_startup_values(observed)
            except Exception:
                pass
            else:
                projection.update(status="error_receipt_observed", failure=core)
                if requirement_label is not None:
                    projection["native_requirement"] = requirement_label
                if observed is not None:
                    projection["startup_log_witness"] = observed
        result["clamd_startup_diagnostic"] = projection
    return result


def _diagnose_docx_links(docker: list[str], image: str, platform: str, request: Path,
                        result: object, binding: dict, code: int, *, case: str) -> None:
    _note_docx_fit_mismatch(result, binding, code, case=case)
    report = _failure_report()
    if report.get("fit_pages") != 2 or report.get("fit_links") != 0 or "unchanged_links" not in report.get("fit_failed_checks", []):
        return
    started = time.monotonic()
    container_calls = 0
    values = None
    try:
        expected = request.with_suffix(".diagnostic-binding.json")
        with expected.open("x", encoding="utf-8") as stream:
            json.dump(binding, stream, allow_nan=False)
        expected.chmod(0o444)
        container_calls = 1
        child_code, raw = container(docker, image, platform, ["-c", DOCX_LINK_PROGRAM],
            [(request, "/tmp/request.json"), (expected, "/tmp/diagnostic-binding.json")])
        observed = json.loads(raw)
        if (type(child_code) is int and child_code == 0 and type(observed) is dict
                and set(observed) == set(_LINK_FIELDS)
                and all(type(observed[key]) is int and 0 <= observed[key] <= 100 for key in (_LINK_FIELDS[i] for i in _LINK_COUNT_INDEXES))
                and all(type(observed[key]) is bool for key in (_LINK_FIELDS[i] for i in _LINK_BOOL_INDEXES))):
            values = tuple(observed[key] for key in _LINK_FIELDS)
    except Exception:
        pass
    diagnostic = _FAILURE.get()
    if type(diagnostic) is dict:
        diagnostic["docx_link_observation"] = _DocxLinkObservation(container_calls, round(time.monotonic() - started, 6), values)


def command(args: list[str], *, timeout: int = 20, limit: int = 1024 * 1024) -> tuple[int, bytes]:
    child = None
    try:
        child = subprocess.Popen(args, stdin=subprocess.DEVNULL, stdout=subprocess.PIPE,
                                 stderr=subprocess.DEVNULL, start_new_session=True)
        assert child.stdout is not None
        deadline = time.monotonic() + timeout
        output = bytearray()
        with selectors.DefaultSelector() as selector:
            selector.register(child.stdout, selectors.EVENT_READ)
            while selector.get_map():
                left = deadline - time.monotonic()
                if left <= 0:
                    raise DocumentImageError()
                for key, _ in selector.select(min(left, 1)):
                    chunk = os.read(key.fd, 65536)
                    if not chunk:
                        selector.unregister(key.fileobj)
                    else:
                        output.extend(chunk)
                        if len(output) > limit:
                            raise DocumentImageError()
            left = deadline - time.monotonic()
            if left <= 0:
                raise DocumentImageError()
            return child.wait(timeout=left), bytes(output)
    except Exception:
        raise DocumentImageError() from None
    finally:
        if child is not None:
            try:
                os.killpg(child.pid, signal.SIGKILL)
            except ProcessLookupError:
                pass
            child.wait()
            if child.stdout:
                child.stdout.close()


def successful(args: list[str], *, timeout: int = 20) -> bytes:
    code, output = command(args, timeout=timeout)
    if code:
        raise DocumentImageError()
    return output


def docker_command() -> list[str]:
    # Refuse a remote or caller-routed daemon. No context/account is changed.
    if os.environ.get("DOCKER_HOST") or os.environ.get("DOCKER_CONTEXT"):
        raise DocumentImageError()
    contexts = json.loads(successful(["docker", "context", "inspect"]))
    if not isinstance(contexts, list) or len(contexts) != 1:
        raise DocumentImageError()
    host = contexts[0].get("Endpoints", {}).get("docker", {}).get("Host", "")
    if not isinstance(host, str) or not re.fullmatch(r"unix:///[^\x00\r\n?#]+", host):
        raise DocumentImageError()
    return ["docker", "--host=" + host]


def source_hashes() -> dict[str, str]:
    return {path: hashlib.sha256((ROOT / path).read_bytes()).hexdigest() for path in SOURCES}


def build(docker: list[str], platform: str, hashes: dict[str, str]) -> str:
    with tempfile.TemporaryDirectory(prefix="hirewiz-document-build-") as directory:
        context = Path(directory) / "context"
        context.mkdir()
        for path in SOURCES:
            content = (ROOT / path).read_bytes()
            if hashlib.sha256(content).hexdigest() != hashes[path]:
                raise DocumentImageError()
            copied = context / path
            copied.parent.mkdir(parents=True, exist_ok=True)
            copied.write_bytes(content)
            copied.chmod(0o644)
        iid = Path(directory) / "image.id"
        args = docker + ["build", "--pull=false", "--platform", platform, "--iidfile", str(iid),
                         "--file", str(context / SOURCES[0]),
                         "--build-arg", "DOCUMENT_DEFINITIONS_NONCE=" + uuid.uuid4().hex]
        for path, digest in hashes.items():
            args += ["--label", LABEL + path + "=" + digest]
        args.append(str(context))
        successful(args, timeout=600)
        image = iid.read_text().strip()
    if not IMAGE_ID.fullmatch(image) or source_hashes() != hashes:
        raise DocumentImageError()
    return image


def image_metadata(docker: list[str], image: str, platform: str, hashes: dict[str, str]) -> None:
    if not IMAGE_ID.fullmatch(image):
        raise DocumentImageError()
    entries = json.loads(successful(docker + ["image", "inspect", image]))
    if not isinstance(entries, list) or len(entries) != 1:
        raise DocumentImageError()
    value = entries[0]
    config = value.get("Config", {})
    labels = config.get("Labels", {})
    if (value.get("Id") != image or platform != value.get("Os", "") + "/" + value.get("Architecture", "")
            or config.get("User") != "65534:65534"
            or config.get("Entrypoint") != ["/usr/local/bin/python3", "-I", "/opt/document/service.py"]
            or any(labels.get(LABEL + path) != digest for path, digest in hashes.items())):
        raise DocumentImageError()


def check_container(value: dict, image: str, owned: str, mounts: list[str], arguments: list[str], *, terminal: bool) -> None:
    config, host, state = value.get("Config", {}), value.get("HostConfig", {}), value.get("State", {})
    observed_mounts = value.get("Mounts", [])
    if (value.get("Name") != "/" + owned or value.get("Image") != image
            or config.get("Image") != image or config.get("User") != "65534:65534"
            or config.get("Entrypoint") != ["/usr/local/bin/python3"] or config.get("Cmd") != ["-I"] + arguments
            or host.get("NetworkMode") != "none" or host.get("ReadonlyRootfs") is not True
            or host.get("CapDrop") != ["ALL"] or host.get("CapAdd")
            or host.get("SecurityOpt") != ["no-new-privileges"]
            or host.get("Memory") != MEMORY or host.get("MemorySwap") != MEMORY
            or host.get("NanoCpus") != 1000000000 or host.get("PidsLimit") != 64
            or host.get("Tmpfs") != {"/tmp": TMPFS} or host.get("Privileged") is not False
            or len(observed_mounts) != len(mounts) or any(m.get("Type") != "bind" or m.get("RW") is not False
               or m.get("Source") + ":" + m.get("Destination") not in mounts for m in observed_mounts)
            or state.get("Running") is not False or state.get("OOMKilled") is not False
            or state.get("Dead") is not False or state.get("Error") != ""
            or state.get("Status") != ("exited" if terminal else "created")):
        raise DocumentImageError()


def container(docker: list[str], image: str, platform: str, arguments: list[str],
              mounts: list[tuple[Path, str]]) -> tuple[int, bytes]:
    owned = "hirewiz-document-test-" + uuid.uuid4().hex
    args = docker + ["create", "--pull=never", "--platform", platform, "--name", owned,
        "--network=none", "--read-only", "--cap-drop=ALL", "--security-opt=no-new-privileges",
        "--user=65534:65534", "--cpus=1", "--memory=2g", "--memory-swap=2g", "--pids-limit=64",
        "--tmpfs=/tmp:" + TMPFS, "--entrypoint=/usr/local/bin/python3"]
    expected = []
    for source, target in mounts:
        source = source.resolve(strict=True)
        args += ["--mount", "type=bind,source=" + str(source) + ",target=" + target + ",readonly"]
        expected.append(str(source) + ":" + target)
    args += [image, "-I"] + arguments
    try:
        successful(args)
        before = json.loads(successful(docker + ["inspect", owned]))
        if not isinstance(before, list) or len(before) != 1:
            raise DocumentImageError()
        check_container(before[0], image, owned, expected, arguments, terminal=False)
        code, output = command(docker + ["start", "--attach", owned], timeout=65)
        after = json.loads(successful(docker + ["inspect", owned]))
        if not isinstance(after, list) or len(after) != 1:
            raise DocumentImageError()
        check_container(after[0], image, owned, expected, arguments, terminal=True)
        if after[0]["State"].get("ExitCode") != code:
            raise DocumentImageError()
        return code, output
    finally:
        # Covers client timeout/unknown state; never remove somebody else's name.
        successful(docker + ["rm", "--force", owned], timeout=10)


def corpus() -> list[tuple[str, str, bytes, str]]:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from pypdf import PdfWriter

    def pdf(active=False):
        writer = PdfWriter()
        writer.add_blank_page(200, 200)
        if active:
            writer.add_js("this.print();")
        stream = io.BytesIO()
        writer.write(stream)
        return stream.getvalue()
    def docx(field=None):
        document = Document()
        document.add_paragraph("Candidate Example")
        document.add_paragraph("SKILLS")
        document.add_paragraph("Python, SQL")
        if field:
            node = OxmlElement("w:fldSimple")
            node.set(qn("w:instr"), field)
            document.add_paragraph()._p.append(node)
        stream = io.BytesIO()
        document.save(stream)
        return stream.getvalue()
    clean_docx = docx()
    infected = io.BytesIO()
    with ZipFile(io.BytesIO(clean_docx)) as original, ZipFile(infected, "w", ZIP_DEFLATED) as output:
        for entry in original.infolist():
            output.writestr(entry, original.read(entry))
        output.writestr("word/standalone-antivirus-test.txt", EICAR)
    if len(EICAR) != 68:
        raise DocumentImageError()
    return [("clean_pdf", "pdf", pdf(), "clean"), ("clean_docx", "docx", clean_docx, "clean"),
            ("page_counter", "docx", docx("PAGE"), "clean"),
            ("active_pdf", "pdf", pdf(True), "refused"),
            ("includetext", "docx", docx('INCLUDETEXT "https://fixture.invalid/candidate.docx"'), "refused"),
            ("eicar_exact_file", "pdf", EICAR, "detected"),
            ("eicar_zip_member", "docx", infected.getvalue(), "detected"),
            ("missing_definitions", "pdf", pdf(), "missing")]


def docx_layout_corpus(image: str) -> list[tuple[str, bytes, dict]]:
    """One genuine styled two-page fit and one genuine reflow refusal.

    The fixture is synthetic; the expected candidate is the exact source-pinned
    native patch, not a template or a purported Word render produced on the host.
    """
    from docx import Document
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
    spec = importlib.util.spec_from_file_location("document_native_docx_patch", ROOT / "backend/app/services/resume_layout.py")
    if spec is None or spec.loader is None:
        raise DocumentImageError()
    layout = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = layout
    spec.loader.exec_module(layout)
    original = "Built scalable systems with Python and SQL."
    document = Document()
    document.styles["Normal"].font.name = "Cambria"
    document.styles["Normal"].font.size = Pt(11)
    document.add_paragraph("Taylor Example")
    paragraph = document.add_paragraph()
    paragraph.add_run("Built scalable systems with ")
    paragraph.add_run("Python").bold = True
    paragraph.add_run(" and ")
    paragraph.add_run("SQL").italic = True
    paragraph.add_run(".")
    document.add_paragraph("Reduced API latency by 40% through focused profiling.")
    link_paragraph = document.add_paragraph()
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), document.part.relate_to("https://example.test/portfolio", RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "Portfolio"
    run.append(text)
    link.append(run)
    link_paragraph._p.append(link)
    document.add_page_break()
    document.add_paragraph("Original education and second-page running content.")
    stream = io.BytesIO()
    document.save(stream)
    source = stream.getvalue()
    unit = next(unit for unit in layout.extract_source_units(source, "docx") if unit["text"] == original)
    policy = hashlib.sha256((ROOT / "infra/document-processing/policy.json").read_bytes()).hexdigest()
    rows = []
    for name, replacement in (("styled_multipage_fit", "Built systems with Python and SQL."),
                              ("rendered_reflow_refusal", original + " reusable systems" * 40)):
        edits = [{"unit_id": unit["unit_id"], "original_text": original, "replacement_text": replacement}]
        candidate = layout._apply_docx(source, edits)
        expected = {"version": 1, "operation": "docx_layout", "sha256": hashlib.sha256(source).hexdigest(),
            "size_bytes": len(source), "source_format": "docx", "candidate_sha256": hashlib.sha256(candidate).hexdigest(),
            "candidate_size_bytes": len(candidate), "edits_sha256": hashlib.sha256(json.dumps(edits, sort_keys=True,
                separators=(",", ":"), ensure_ascii=True).encode()).hexdigest()}
        request = {"version": 1, "source_format": "docx", "sha256": expected["sha256"],
            "content_base64": base64.b64encode(source).decode(), "edits": edits,
            "candidate_sha256": expected["candidate_sha256"], "candidate_size_bytes": len(candidate),
            "worker_image": "local-document-image@" + image, "policy_sha256": policy}
        rows.append((name, json.dumps(request).encode(), expected))
    return rows


def clamd_corpus() -> list[tuple[str, bytes, str, bool]]:
    """New genuine daemon controls; ordinary clamscan corpus stays unchanged."""
    from docx import Document
    from pypdf import PdfWriter
    fixtures = {name: content for name, _, content, _ in corpus()}
    clean, docx = fixtures["clean_pdf"], fixtures["clean_docx"]
    spec = importlib.util.spec_from_file_location("document_native_clamd_patch", ROOT / "backend/app/services/resume_layout.py")
    if spec is None or spec.loader is None:
        raise DocumentImageError()
    layout = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = layout
    spec.loader.exec_module(layout)
    resident_text = "Developed Python automation tools for reliable test workflows."
    document = Document()
    document.add_paragraph("Candidate Example")
    document.add_paragraph(resident_text)
    resident_stream = io.BytesIO()
    document.save(resident_stream)
    resident_docx = resident_stream.getvalue()
    unit = next(value for value in layout.extract_source_units(resident_docx, "docx")
                if value["text"] == resident_text)
    resident_edits = [{"unit_id": unit["unit_id"], "original_text": unit["text"],
                       "replacement_text": "Built Python automation tools for reliable test workflows."}]
    resident_candidate = layout._apply_docx(resident_docx, resident_edits)
    def archive(rows):
        stream = io.BytesIO()
        with ZipFile(stream, "w", ZIP_DEFLATED) as zipped:
            for name, content in rows:
                zipped.writestr(name, content)
        return stream.getvalue()
    writer = PdfWriter()
    writer.add_blank_page(200, 200)
    writer.encrypt("public-synthetic-fixture-password")
    encrypted_pdf = io.BytesIO()
    writer.write(encrypted_pdf)
    encrypted_zip = bytearray(archive([("public.txt", b"public synthetic bytes")]))
    # Deliberately mark an owned synthetic ZIP entry encrypted. No password or
    # customer material is involved; real daemon refusal/unavailability required.
    for magic, offset in ((b"PK\x03\x04", 6), (b"PK\x01\x02", 8)):
        index = encrypted_zip.index(magic) + offset
        flags = int.from_bytes(encrypted_zip[index:index + 2], "little") | 1
        encrypted_zip[index:index + 2] = flags.to_bytes(2, "little")
    large = archive([("large.txt", b"x" * (6 * 1024**2 + 1))])
    total = archive([(str(i) + ".txt", b"x" * (5 * 1024**2)) for i in range(5)])
    files = archive([(str(i) + ".txt", b"x") for i in range(501)])
    nested = b"public synthetic bytes"
    for _ in range(18):
        nested = archive([("nested.zip", nested)])
    rows = [("clamd_two_clean", clean, docx, "clean", False),
        ("clamd_source_eicar", EICAR, docx, "source_refused", False),
        ("clamd_candidate_eicar", clean, EICAR, "candidate_refused", False),
        ("clamd_zip_member_eicar", clean, fixtures["eicar_zip_member"], "candidate_refused", False),
        ("clamd_encrypted_pdf", clean, encrypted_pdf.getvalue(), "candidate_nonclean", False),
        ("clamd_encrypted_zip", clean, bytes(encrypted_zip), "candidate_nonclean", False),
        ("clamd_max_filesize", clean, large, "candidate_nonclean", False),
        ("clamd_max_scansize", clean, total, "candidate_nonclean", False),
        ("clamd_max_files", clean, files, "candidate_nonclean", False),
        ("clamd_max_recursion", clean, nested, "candidate_nonclean", False),
        ("clamd_resident_parser", resident_docx, resident_candidate, "clean", True)]
    result = []
    for name, original, candidate, expected, resident in rows:
        if not 0 < len(original) <= 5 * 1024**2 or not 0 < len(candidate) <= 5 * 1024**2:
            raise DocumentImageError()
        payload = json.dumps({"case": name, "source": base64.b64encode(original).decode(),
            "candidate": base64.b64encode(candidate).decode(), "expected": expected,
            "resident_parser": resident, "candidate_sha256": hashlib.sha256(candidate).hexdigest(),
            "candidate_size_bytes": len(candidate), "edits": resident_edits if resident else None},
            separators=(",", ":")).encode()
        if len(payload) > 7 * 1024**2:
            raise DocumentImageError()
        result.append((name, payload, expected, resident))
    return result


def _validate_clamd_protocol(value: dict, scans: list, payload: bytes, outcome: str) -> None:
    """A completed AV reply is distinct from a file admitted as clean."""
    protocol = value.get("protocol_scans")
    attempts = value.get("instream_attempts")
    refusal = value.get("zip_admission_failure")
    if (type(protocol) is not list or not 1 <= len(protocol) <= 2
            or type(attempts) is not int or attempts != len(protocol)):
        raise DocumentImageError()
    fields = {"sha256", "size_bytes", "metadata", "reply"}
    for row in protocol:
        if (type(row) is not dict or set(row) != fields
                or type(row["sha256"]) is not str or not re.fullmatch(r"[a-f0-9]{64}", row["sha256"])
                or type(row["size_bytes"]) is not int or not 0 < row["size_bytes"] <= 5 * 1024**2
                or type(row["reply"]) is not str or row["reply"] not in {"OK", "FOUND", "OTHER"}
                or type(row["metadata"]) is not dict or set(row["metadata"]) != {"engine", "version", "definitions"}
                or row["metadata"].get("engine") != "ClamAV" or row["metadata"].get("version") != "1.4.3"
                or type(row["metadata"].get("definitions")) is not str
                or not re.fullmatch(r"[0-9]{1,10}/[0-9]{4}-[0-9]{2}-[0-9]{2}T[0-9]{2}:[0-9]{2}:[0-9]{2}\+00:00", row["metadata"]["definitions"])):
            raise DocumentImageError()
    if len(protocol) == 2 and protocol[1]["metadata"] != protocol[0]["metadata"]:
        raise DocumentImageError()
    requested = json.loads(payload)
    source = base64.b64decode(requested["source"], validate=True)
    identities = [(hashlib.sha256(source).hexdigest(), len(source)),
                  (requested["candidate_sha256"], requested["candidate_size_bytes"])]
    for index, row in enumerate(protocol):
        if (row["sha256"], row["size_bytes"]) != identities[index]:
            raise DocumentImageError()
    for index, row in enumerate(scans):
        recorded = protocol[index] if index < len(protocol) else None
        if (recorded is None or any(row[key] != recorded[key] for key in ("sha256", "size_bytes", "metadata"))
                or recorded["reply"] != {"CL_CLEAN": "OK", "CL_VIRUS": "FOUND"}[row["result"]]):
            raise DocumentImageError()
    replies = [row["reply"] for row in protocol]
    sequences = {"source_refused": ["FOUND"], "candidate_refused": ["OK", "FOUND"],
                 "clean": ["OK", "OK"]}
    if outcome != "candidate_nonclean":
        if replies != sequences[outcome] or refusal is not None:
            raise DocumentImageError()
        return
    if replies == ["OK", "OTHER"] and refusal is None:
        return
    if replies != ["OK", "OK"] or type(refusal) is not dict:
        raise DocumentImageError()
    if (set(refusal) != {"sha256", "size_bytes", "metadata", "reason"}
            or type(refusal["sha256"]) is not str or type(refusal["size_bytes"]) is not int
            or type(refusal["metadata"]) is not dict
            or type(refusal["reason"]) is not str or refusal["reason"] not in {
                "malformed", "unsupported", "encrypted", "member_bytes", "total_bytes", "members", "depth", "deadline"}
            or any(refusal[key] != protocol[1][key] for key in ("sha256", "size_bytes", "metadata"))):
        raise DocumentImageError()


def validate_clamd_native(value: dict, name: str, expected: str, resident: bool,
                          payload: bytes, platform: str) -> None:
    fields = {"case", "outcome", "native_installation", "owned_peer", "loaded_signatures", "scans",
        "daemon_returncode", "daemon_stopped_before_writer", "native_log_removed", "renderer_version_after_stop",
        "resident_parser_exercised", "resident_parser_cgroup_current_bytes", "cgroup_memory_max_bytes",
        "cgroup_memory_peak_bytes", "performance_proven", "cloud_run_isolation_proven",
        "protocol_scans", "instream_attempts", "zip_admission_failure"}
    if type(value) is not dict or set(value) != fields or value.get("case") != name:
        raise DocumentImageError()
    outcome = value.get("outcome")
    if outcome != expected and not (expected == "candidate_nonclean" and outcome == "candidate_refused"):
        raise DocumentImageError()
    manifest = value.get("native_installation")
    peer = value.get("owned_peer")
    scans = value.get("scans")
    machine, library_directory = {"linux/amd64": (62, "x86_64-linux-gnu"),
                                  "linux/arm64": (183, "aarch64-linux-gnu")}[platform]
    if (type(manifest) is not dict or set(manifest) != {"version", "engine", "files"}
            or type(manifest.get("version")) is not int or manifest["version"] != 1 or manifest.get("engine") != "1.4.3"
            or type(manifest.get("files")) is not list or len(manifest["files"]) != 3
            or type(peer) is not dict or set(peer) != {"pid", "uid", "gid"}
            or any(type(peer[k]) is not int or peer[k] < 0 for k in peer) or peer["pid"] == 0
            or type(value.get("loaded_signatures")) is not int or value["loaded_signatures"] < 1
            or type(scans) is not list or not 1 <= len(scans) <= 2
            or type(value.get("daemon_returncode")) is not int or value["daemon_returncode"] != 0
            or value.get("daemon_stopped_before_writer") is not True or value.get("native_log_removed") is not True
            or type(value.get("renderer_version_after_stop")) is not str
            or not re.fullmatch(r"LibreOffice [0-9]+\.[0-9]+\.[0-9]+(?:\.[0-9]+)*(?: [ -~]+)?", value["renderer_version_after_stop"])
            or value.get("resident_parser_exercised") is not resident
            or type(value.get("cgroup_memory_max_bytes")) is not int or value["cgroup_memory_max_bytes"] != MEMORY
            or type(value.get("cgroup_memory_peak_bytes")) is not int or not 0 < value["cgroup_memory_peak_bytes"] <= MEMORY
            or value.get("performance_proven") is not False or value.get("cloud_run_isolation_proven") is not False):
        raise DocumentImageError()
    paths = [row.get("path") for row in manifest["files"] if type(row) is dict]
    if (len(paths) != 3 or paths[:2] != ["/usr/sbin/clamd", "/usr/bin/clamscan"]
            or type(paths[2]) is not str or not re.fullmatch(
                "/usr/lib/" + library_directory + r"/libclamav\.so(?:\.[0-9]+)+", paths[2])):
        raise DocumentImageError()
    for row in manifest["files"]:
        if (type(row) is not dict or set(row) != {"path", "sha256", "size", "elf_machine"}
                or type(row["path"]) is not str or not row["path"].startswith("/usr/")
                or type(row["sha256"]) is not str or not re.fullmatch(r"[a-f0-9]{64}", row["sha256"])
                or type(row["size"]) is not int or not 20 <= row["size"] <= 128 * 1024**2
                or type(row["elf_machine"]) is not int or row["elf_machine"] != machine):
            raise DocumentImageError()
    for row in scans:
        if (type(row) is not dict or set(row) != {"sha256", "size_bytes", "metadata", "result"}
                or type(row["sha256"]) is not str or not re.fullmatch(r"[a-f0-9]{64}", row["sha256"])
                or type(row["size_bytes"]) is not int or not 0 < row["size_bytes"] <= 5 * 1024**2
                or row["result"] not in {"CL_CLEAN", "CL_VIRUS"}
                or type(row["metadata"]) is not dict or set(row["metadata"]) != {"engine", "version", "definitions"}
                or row["metadata"].get("engine") != "ClamAV" or row["metadata"].get("version") != "1.4.3"
                or type(row["metadata"].get("definitions")) is not str
                or not re.fullmatch(r"[0-9]{1,10}/[0-9]{4}-[0-9]{2}-[0-9]{2}T[0-9]{2}:[0-9]{2}:[0-9]{2}\+00:00", row["metadata"]["definitions"])):
            raise DocumentImageError()
    sequences = {"source_refused": ["CL_VIRUS"], "candidate_refused": ["CL_CLEAN", "CL_VIRUS"],
                 "candidate_nonclean": ["CL_CLEAN"], "clean": ["CL_CLEAN", "CL_CLEAN"]}
    if [row["result"] for row in scans] != sequences[outcome]:
        raise DocumentImageError()
    requested = json.loads(payload)
    source = base64.b64decode(requested["source"], validate=True)
    if scans[0]["sha256"] != hashlib.sha256(source).hexdigest() or scans[0]["size_bytes"] != len(source):
        raise DocumentImageError()
    if len(scans) == 2 and (scans[1]["sha256"] != requested["candidate_sha256"]
            or scans[1]["size_bytes"] != requested["candidate_size_bytes"]):
        raise DocumentImageError()
    if resident and (type(value.get("resident_parser_cgroup_current_bytes")) is not int
                     or not 0 < value["resident_parser_cgroup_current_bytes"] <= MEMORY):
        raise DocumentImageError()
    if not resident and value["resident_parser_cgroup_current_bytes"] is not None:
        raise DocumentImageError()
    _validate_clamd_protocol(value, scans, payload, outcome)


def check(platform: str, *, image: str | None = None) -> dict:
    _progress("daemon", completed_cases=0)
    docker = docker_command()
    _progress("source_hashes")
    hashes = source_hashes()
    _progress("build")
    image = image or build(docker, platform, hashes)
    _progress("image_metadata")
    image_metadata(docker, image, platform, hashes)
    _progress("runtime_sources")
    code, raw = container(docker, image, platform, ["-c", HASH_PROGRAM], [])
    _progress("runtime_sources", exit_code=code)
    if code or json.loads(raw) != {path: hashes[path] for path in RUNTIME}:
        raise DocumentImageError()
    observations = []
    layout_observations = []
    clamd_observations = []
    _progress("corpus")
    with tempfile.TemporaryDirectory(prefix="hirewiz-document-corpus-") as temporary:
        directory = Path(temporary)
        directory.chmod(0o755)
        empty = directory / "empty-definitions"
        empty.mkdir(mode=0o755)
        for name, kind, content, outcome in corpus():
            _progress("case_request", case=name)
            envelope = {"version": 1, "source_format": kind, "content_base64": base64.b64encode(content).decode(),
                        "sha256": hashlib.sha256(content).hexdigest()}
            request = directory / (name + ".json")
            request.write_text(json.dumps(envelope))
            request.chmod(0o444)
            mounts = [(request, "/tmp/request.json")]
            if outcome == "missing":
                mounts.append((empty, "/var/lib/clamav"))
            arguments = ["-c", DETECTION_PROGRAM] if outcome == "detected" else ["/opt/document/processor.py", "/tmp/request.json"]
            _progress("case_execution", case=name)
            code, raw = container(docker, image, platform, arguments, mounts)
            _progress("case_validation", case=name, exit_code=code)
            result = json.loads(raw)
            if outcome == "detected":
                if code or result != {"eicar_detected": True, "processor_refused": True}:
                    raise DocumentImageError()
            elif outcome in {"refused", "missing"}:
                expected = {"version": 1, "sha256": envelope["sha256"], "size_bytes": len(content),
                            "source_format": kind, "error": "document_inspection_unavailable" if outcome == "missing" else "document_inspection_refused"}
                if code != (2 if outcome == "missing" else 1) or result != expected:
                    raise DocumentImageError()
            elif (code or result.get("sha256") != envelope["sha256"] or result.get("size_bytes") != len(content)
                  or result.get("source_format") != kind or result.get("scan", {}).get("engine") != "ClamAV"
                  or (kind == "docx" and "Python" not in result.get("parsed", {}).get("skills", []))):
                raise DocumentImageError()
            observations.append({"case": name, "outcome": outcome, "passed": True,
                                 **({"scan": result["scan"]} if outcome == "clean" else {})})
            _progress("case_validation", case=name, exit_code=code, completed_cases=len(observations))
        _progress("clamd_corpus")
        for name, payload, expected_outcome, resident in clamd_corpus():
            _progress("case_request", case=name)
            request = directory / (name + ".json")
            request.write_bytes(payload)
            request.chmod(0o444)
            _progress("case_execution", case=name)
            code, raw = container(docker, image, platform, ["-c", CLAMD_NATIVE_PROGRAM], [(request, "/tmp/request.json")])
            _progress("case_validation", case=name, exit_code=code)
            if code:
                _note_clamd_failure(raw, code, case=name)
                _diagnose_clamd_startup(docker, image, platform, request, case=name)
                raise DocumentImageError()
            value = json.loads(raw)
            validate_clamd_native(value, name, expected_outcome, resident, payload, platform)
            clamd_observations.append(value)
        _progress("layout_corpus")
        for name, payload, binding in docx_layout_corpus(image):
            _progress("case_request", case=name)
            request = directory / (name + ".json")
            request.write_bytes(payload)
            request.chmod(0o444)
            _progress("case_execution", case=name)
            code, raw = container(docker, image, platform, ["-c", DOCX_LAYOUT_PROGRAM], [(request, "/tmp/request.json")])
            _progress("case_validation", case=name, exit_code=code)
            result = json.loads(raw)
            if any(type(result.get(key)) is not type(value) or result.get(key) != value for key, value in binding.items()):
                raise DocumentImageError()
            if name == "styled_multipage_fit":
                receipt = result.get("layout", {})
                if (code or receipt.get("validator") != "hirewiz-docx-layout-v1" or receipt.get("pages") != 2
                        or receipt.get("unchanged_links") != 1 or receipt.get("outside_changed_regions_pixel_delta") != 0
                        or result.get("scan", {}).get("engine") != "ClamAV" or result.get("candidate_scan") != result.get("scan")):
                    _note_docx_refusal(result, binding, code, case=name)
                    _note_docx_fit_mismatch(result, binding, code, case=name)
                    _diagnose_docx_links(docker, image, platform, request, result, binding, code, case=name)
                    raise DocumentImageError()
                layout_observations.append({"case": name, "passed": True, "layout": receipt, "scan": result["scan"]})
            # The core compares the complete link tuple before line bands.
            # This overlong paragraph moves the following link and is refused.
            elif (code != 1 or result.get("error") != "docx_layout_refused" or result.get("reason") not in {
                    "line_wrap_changed", "page_geometry_changed", "paragraph_page_changed", "paragraph_overflow",
                    "hyperlink_changed"}):
                _note_docx_refusal(result, binding, code, case=name)
                raise DocumentImageError()
            else:
                layout_observations.append({"case": name, "passed": True, "reason": result["reason"]})
            _progress("case_validation", case=name, exit_code=code,
                      completed_cases=len(observations) + len(layout_observations))
    _progress("source_recheck")
    if source_hashes() != hashes:
        raise DocumentImageError()
    return {"scope": "local_native_docker_only", "cloud_run_isolation_proven": False,
            "image_id": image, "platform": platform, "source_sha256": hashes,
            "runtime_source_bytes_verified": True, "restricted_container_metadata_verified": True,
            "corpus": observations, "docx_layout_corpus": layout_observations,
            "request_clamd_corpus": clamd_observations}


def main() -> None:
    os.umask(0o077)
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--platform", choices=("linux/amd64", "linux/arm64"), required=True)
    parser.add_argument("--image-id", help="Reuse only an immutable ID with exact current source labels/bytes.")
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    token = _FAILURE.set({"stage": "setup", "completed_cases": 0})
    try:
        report = json.dumps(check(args.platform, image=args.image_id), sort_keys=True, indent=2) + "\n"
        _progress("report_write")
        if args.output:
            with args.output.open("x", encoding="utf-8") as stream:
                stream.write(report)
        print(report, end="")
    except Exception:
        print(json.dumps(_failure_report()))
        raise SystemExit(1) from None
    finally:
        _FAILURE.reset(token)


if __name__ == "__main__":
    main()
