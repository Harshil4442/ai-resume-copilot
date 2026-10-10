"""Explicit production-Next/Chromium workflow against uniquely isolated local PG.

Not collected by pytest. Only this invocation's schema and processes are removed.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import os
import re
import selectors
import signal
import socket
import ssl
import subprocess
import sys
import time
from pathlib import Path
from typing import IO, cast
from uuid import uuid4

import httpx
from docx import Document
from sqlalchemy import create_engine, text
from sqlalchemy.engine import make_url

ROOT = Path(__file__).resolve().parents[2]
PG_DEFAULT = "postgresql+psycopg://hirewiz:hirewiz@127.0.0.1:55433/hirewiz_admission_test"


def _port():
    with socket.socket() as listener:
        listener.bind(("127.0.0.1", 0))
        return listener.getsockname()[1]


def _wait(url, process, seconds=90, *, verify=True):
    deadline = time.monotonic() + seconds
    while time.monotonic() < deadline:
        assert _owned_status(process) is None, "Owned server exited before readiness; inspect local log"
        try:
            if httpx.get(url, timeout=2, verify=verify).status_code == 200:
                return
        except httpx.HTTPError:
            pass
        time.sleep(0.2)
    raise AssertionError("Owned local server did not become ready within the deadline")


def _create_owned_tls(key: Path, certificate: Path) -> ssl.SSLContext:
    """Trust one invocation's loopback certificate without relaxing auth guards."""
    try:
        subprocess.run(["openssl", "req", "-x509", "-newkey", "rsa:2048", "-nodes", "-days", "1",
            "-keyout", str(key), "-out", str(certificate), "-subj", "/CN=127.0.0.1",
            "-addext", "subjectAltName=IP:127.0.0.1", "-addext", "basicConstraints=critical,CA:TRUE"],
            check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=10)
        key.chmod(0o600)
        return ssl.create_default_context(cafile=str(certificate))
    except BaseException:
        key.unlink(missing_ok=True)
        certificate.unlink(missing_ok=True)
        raise


# Never reap a session leader before signalling its group: an unreaped owned child
# pins its PID/PGID, preventing a recycled numeric ID from targeting another group.
_OWNED_GROUPS: dict[int, subprocess.Popen[bytes]] = {}
_STOPPED_GROUPS: dict[subprocess.Popen[bytes], dict[str, object]] = {}


def _spawn_owned(command, *, cwd, env, stdout, stderr) -> subprocess.Popen[bytes]:
    assert signal.getsignal(signal.SIGCHLD) == signal.SIG_DFL, "Owned leader requires default non-reaping SIGCHLD policy"
    process = subprocess.Popen(command, cwd=cwd, env=env, stdout=stdout,
                               stderr=stderr, start_new_session=True)
    _OWNED_GROUPS[process.pid] = process
    return process


def _owned_status(process: subprocess.Popen[bytes]) -> int | None:
    if process.returncode is not None:
        return process.returncode
    if hasattr(os, "waitid"):
        status = os.waitid(os.P_PID, process.pid, os.WEXITED | os.WNOHANG | os.WNOWAIT)
        return None if status is None else status.si_status
    # Some macOS Python distributions omit waitid. Inspect only the retained,
    # unreaped leader's state; never call Popen.poll/wait until group cleanup.
    states = [state for pid, state in _group_members(process.pid) if pid == process.pid]
    assert len(states) == 1, "Unreaped owned leader is unavailable; refusing numeric group reuse"
    return 0 if states[0].startswith("Z") else None


def _group_members(group_id) -> list[tuple[int, str]]:
    # PID/PGID/state only; never collect commands, environment or credentials.
    result = subprocess.run(["ps", "-A", "-o", "pid=,pgid=,stat="],
                            capture_output=True, text=True, check=True, timeout=2)
    return [(int(pid), state) for line in result.stdout.splitlines()
            for pid, group, state in [line.split()]
            if int(group) == group_id]


def _stop(process: subprocess.Popen[bytes]) -> dict[str, object]:
    if process in _STOPPED_GROUPS:
        return _STOPPED_GROUPS[process]  # A vanished/reused group is never signalled again.
    if _OWNED_GROUPS.get(process.pid) is not process:
        raise AssertionError("Refusing to signal a process group without retained ownership")
    if process.returncode is not None:
        raise AssertionError("Owned leader was reaped before group cleanup; refusing a reusable PGID")
    try:
        _owned_status(process)  # WNOWAIT proves this is still our unreaped child.
    except ChildProcessError as exc:
        raise AssertionError("Owned leader identity is unavailable; group was not signalled") from exc
    # _spawn_owned used start_new_session=True. macOS getsid rejects a zombie,
    # so the retained child identity (not a late numeric getsid) owns the group.
    def finished() -> dict[str, object] | None:
        members = _group_members(process.pid)
        # Zombies are already terminated. Retain our leader until all live
        # members are gone, then reap it without sending any further signal.
        if _owned_status(process) is not None and all(state.startswith("Z") for _, state in members):
            process.wait(timeout=1)
            _OWNED_GROUPS.pop(process.pid)
            result = {"group_id": process.pid, "verified_stopped": True,
                      "verification": "unreaped owned leader; no live group members"}
            _STOPPED_GROUPS[process] = result
            return result
        return None

    for sig, seconds in ((signal.SIGTERM, 5), (signal.SIGKILL, 3)):
        result = finished()
        if result is not None:
            return result
        try:
            os.killpg(process.pid, sig)
        except ProcessLookupError:
            pass  # Already vanished, while the unreaped leader still pins its ID.
        except PermissionError:
            # macOS can return EPERM when the last live member just exited.
            result = finished()
            if result is not None:
                return result
            raise
        deadline = time.monotonic() + seconds
        while True:
            result = finished()
            if result is not None:
                return result
            if time.monotonic() >= deadline:
                break
            time.sleep(0.05)
    raise AssertionError("Owned process group still has live members after bounded termination")


def _run_bounded(command, *, cwd, env, timeout):
    process = _spawn_owned(command, cwd=cwd, env=env, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    assert process.stdout is not None and process.stderr is not None
    streams: dict[IO[bytes], list[bytes]] = {process.stdout: [], process.stderr: []}
    deadline = time.monotonic() + timeout
    try:
        with selectors.DefaultSelector() as selector:
            for stream in streams:
                selector.register(stream, selectors.EVENT_READ)
            while selector.get_map() or _owned_status(process) is None:
                remaining = deadline - time.monotonic()
                if remaining <= 0:
                    raise AssertionError("Owned subprocess exceeded its bounded deadline")
                for key, _ in selector.select(min(remaining, 0.1)):
                    chunk = os.read(key.fd, 65536)
                    if chunk:
                        streams[cast(IO[bytes], key.fileobj)].append(chunk)
                    else:
                        selector.unregister(key.fileobj)
    finally:
        try:
            _stop(process)
        finally:
            for stream in streams:
                stream.close()
    return subprocess.CompletedProcess(command, process.returncode,
        b"".join(streams[process.stdout]).decode("utf-8", errors="replace"),
        b"".join(streams[process.stderr]).decode("utf-8", errors="replace"))



def _browser_failure_summary(report: Path) -> dict[str, object]:
    """Export bounded diagnostic categories/locations, never raw browser errors.

    Playwright errors can contain credentials, candidate content or request data.
    The raw report remains local; only fixed categories and this spec's numeric
    locations may enter the uploaded summary.
    """
    if not report.is_file():
        return {"report_state": "absent"}
    if report.stat().st_size > 2_000_000:
        return {"report_state": "too_large"}
    try:
        content = json.loads(report.read_text())
    except (OSError, UnicodeError, json.JSONDecodeError):
        return {"report_state": "unreadable"}
    failures: list[dict[str, object]] = []

    def visit(value: object) -> None:
        if len(failures) >= 8:
            return
        if isinstance(value, dict):
            errors = value.get("errors")
            if isinstance(errors, list):
                for error in errors[:8]:
                    if not isinstance(error, dict):
                        continue
                    message = error.get("message")
                    message = message if isinstance(message, str) else ""
                    if "Executable doesn't exist" in message:
                        category = "browser_not_installed"
                    elif "missing dependencies" in message:
                        category = "browser_dependencies"
                    elif "browserType.launch" in message:
                        category = "browser_launch"
                    elif "Test timeout" in message:
                        category = "test_timeout"
                    elif "page.goto" in message:
                        category = "navigation"
                    elif "expect(" in message:
                        category = "assertion"
                    elif "locator." in message:
                        category = "locator"
                    else:
                        category = "browser_error"
                    failure: dict[str, object] = {"category": category}
                    stack = error.get("stack")
                    stack = stack if isinstance(stack, str) else ""
                    location = re.search(r"cold-no-ai\.spec\.ts:(\d{1,6}):(\d{1,6})(?:\D|$)", stack)
                    if location is not None:
                        failure["spec_line"] = int(location[1])
                        failure["spec_column"] = int(location[2])
                    else:
                        # Playwright assertion errors can omit stack entirely.
                        # Keep only fixed-spec numeric coordinates, never its path.
                        point = error.get("location")
                        if isinstance(point, dict) and str(point.get("file", "")).replace("\\", "/").endswith("/journey/cold-no-ai.spec.ts"):
                            for source, target in (("line", "spec_line"), ("column", "spec_column")):
                                coordinate = point.get(source)
                                if type(coordinate) is int and 1 <= coordinate <= 999999:
                                    failure[target] = coordinate
                    if failure not in failures:
                        failures.append(failure)
                    if len(failures) >= 8:
                        break
            for key, item in value.items():
                if key != "errors":
                    visit(item)
        elif isinstance(value, list):
            for item in value:
                visit(item)

    visit(content)
    return {"report_state": "read", "failures": failures}


def _docx(path, label):
    document = Document()
    document.add_heading("Synthetic Candidate", 0)
    document.add_paragraph(label)
    document.add_heading("Skills", 1)
    document.add_paragraph("Python, PostgreSQL, Docker")
    document.add_heading("Experience", 1)
    document.add_paragraph("Built a Python service using PostgreSQL and documented Docker operations.")
    output = io.BytesIO()
    document.save(output)
    path.write_bytes(output.getvalue())


def _post(client, path, data, expected=200, headers=None):
    response = client.post(path, json=data, headers=headers)
    assert response.status_code == expected, "Actual local HTTP request did not produce expected status"
    return response.json()


def _seed_http(backend, *, verify=True):
    with httpx.Client(base_url=backend, timeout=10, verify=verify) as client:
        _post(client, "/api/auth/register", {"email": "operator.browser@example.com", "password": "synthetic-browser-password-123", "accepted_terms": True, "confirmed_age_18": True})
        login = _post(client, "/api/auth/login", {"email": "operator.browser@example.com", "password": "synthetic-browser-password-123"})
        headers = {"Authorization": "Bearer " + login["access_token"]}
        source = _post(client, "/api/v1/employer-jobs/sources", {
            "employer": "Synthetic Origin Employer", "employer_key": "synthetic-origin-employer",
            "platform": "greenhouse", "board_token": "syntheticbrowser",
            "careers_url": "https://employer.example/careers", "allowed_hosts": ["employer.example"],
            "verification_url": "https://employer.example/careers",
            "verification_note": "Synthetic official careers fixture, not a real employer permission or enrollment.",
        }, 201, headers)
        _post(client, f"/api/v1/employer-jobs/sources/{source['id']}/refresh", {}, 202, headers)
        deadline = time.monotonic() + 30
        while time.monotonic() < deadline:
            catalog = client.get("/api/v1/employer-jobs/catalog", headers=headers).json()
            if catalog["sources"][0]["status"] == "healthy":
                assert catalog["auto_submit_enabled"] is False
                return
            time.sleep(0.2)
        raise AssertionError("Actual durable refresh failed to index the finite synthetic feed")


def _sql_proof(engine):
    with engine.connect() as db:
        def scalar(sql):
            return db.execute(text(sql)).scalar_one()
        assert scalar("SELECT count(*) FROM model_call_events") == 0
        assert scalar("SELECT count(*) FROM analysis_runs") == 0
        assert scalar("SELECT count(*) FROM skill_coverage") == 0
        assert scalar("SELECT count(*) FROM users WHERE email='candidate.browser@example.com' AND terms_accepted_at IS NOT NULL AND age_confirmed_at IS NOT NULL") == 1
        assert scalar("SELECT count(*) FROM employer_application_attempts") == 0
        assert scalar("SELECT count(*) FROM dispatch_outbox WHERE topic='employer.apply'") == 0
        assert scalar("SELECT count(*) FROM service_credit_reservations WHERE operation='job_application'") == 0
        assert scalar("SELECT count(*) FROM entitlement_ledger") == 1
        assert scalar("SELECT count(*) FROM employer_postings") == 4
        assert scalar("SELECT count(*) FROM resumes") == 2
        assert scalar("SELECT count(*) FROM employer_applications WHERE status='manual_handoff' AND charged_credits=0") == 1
        assert scalar("SELECT count(*) FROM employer_application_approvals") == 2
        assert scalar("SELECT count(*) FROM employer_application_approvals WHERE revoked_at IS NOT NULL") == 1
        approvals = db.execute(text("SELECT allowed_actions FROM employer_application_approvals")).scalars().all()
        assert all(sorted(row) == ["fill", "upload"] for row in approvals)
        balance, units = db.execute(text("SELECT job_service_credits, ai_credits FROM users WHERE email='candidate.browser@example.com'")).one()
        assert balance == 96 and units == 52
        assert scalar("SELECT count(*) FROM usage_events WHERE source_type='payment_order' AND event_type='grant' AND amount=2") == 1
        assert scalar("SELECT count(*) FROM usage_events WHERE source_type='analysis_run'") == 0
        assert scalar("SELECT count(*) FROM service_credit_events WHERE event_type='grant' AND amount=100") == 1
        assert scalar("SELECT count(*) FROM payment_orders WHERE sku='starter_bundle' AND gross_amount_minor=64900 AND cost_policy_snapshot IS NOT NULL") == 1
        reservation = db.execute(text("SELECT reserved_amount, committed_amount, released_amount FROM service_credit_reservations WHERE operation='job_search'")).one()
        assert tuple(reservation) == (6, 4, 2)
        current = db.execute(text("SELECT review_snapshot FROM employer_application_approvals WHERE revoked_at IS NULL")).scalar_one()
        artifact_sha = scalar("SELECT sha256 FROM sealed_application_artifacts JOIN employer_applications ON sealed_application_artifacts.id=employer_applications.artifact_id")
        assert current["artifact_sha256"] == artifact_sha and current["resume_choice"] == "custom"
        receipt = scalar("SELECT receipt FROM employer_applications")
        assert receipt is None
        assert scalar("SELECT count(*) FROM user_profiles WHERE target_role='Python Engineer' AND preferred_location='Bengaluru'") == 1
        version = scalar("SELECT version_num FROM alembic_version")
        return {"model_calls": 0, "analysis_runs": 0, "analysis_units_after_purchase": 52, "bundle_analysis_units_granted": 2, "analysis_unit_operation_debits": 0, "posting_index": 4,
                "owned_docx_sources": 2, "service_credit_grants": 1, "service_balance": 96,
                "search_reserved": 6, "search_charged": 4, "search_released": 2,
                "application_fee": 0, "application_attempts": 0, "employer_apply_outbox": 0,
                "approval_records": 2, "revoked_approvals": 1, "final_submit_grants": 0,
                "current_exact_custom_artifact": True, "employer_receipt": None, "migration": version}


def run_journey(evidence: Path):
    assert not any(ROOT.glob("**/.env*")), "Run from a credential-free source archive"
    pg = make_url(os.getenv("HIREWIZ_TEST_POSTGRES_URL", PG_DEFAULT))
    assert (pg.host, pg.port, pg.database) == ("127.0.0.1", 55433, "hirewiz_admission_test"), "Only the dedicated loopback disposable database is accepted"
    schema = "cold_browser_" + uuid4().hex
    marker = evidence / ".hirewiz-cold-browser"
    if evidence.exists() and any(evidence.iterdir()) and not marker.is_file():
        raise RuntimeError("Evidence directory must be new or previously owned by this harness")
    evidence.mkdir(parents=True, exist_ok=True)
    marker.write_text("synthetic-cold-browser-workflow-v1\n")
    for name in ("summary.json", "provider.json", "browser.json", "playwright.json",
                 "build.log", "api.log", "next.log", "browser.log", "manual-handoff-390.png", "manual-handoff-1440.png"):
        (evidence / name).unlink(missing_ok=True)
    source_hashes = {name: hashlib.sha256((ROOT / name).read_bytes()).hexdigest() for name in (
        "frontend/components/EmployerApplicationPanel.tsx", "frontend/journey/cold-no-ai.spec.ts",
        "frontend/playwright.cold.config.ts", "frontend/app/api/backend/[...path]/route.ts",
        "frontend/tests/candidate-https-server.mjs", "frontend/app/api/account/[operation]/route.ts",
        "frontend/lib/accountTransportServer.ts", "frontend/lib/accountTransportClient.ts",
        "frontend/lib/candidateAuthServer.ts", "frontend/lib/auth.ts", "frontend/app/register/page.tsx",
        "frontend/lib/authOptions.ts", "frontend/package-lock.json", "backend/app/main.py",
        "backend/app/database.py", "backend/app/security.py", "backend/app/domains/employer/service.py",
        "backend/app/routers/resume.py", "backend/app/services/document_ingestion.py",
        "backend/app/domains/employer/credits.py", "backend/app/billing/service.py",
        "backend/scripts/run_cold_browser_journey.py", "backend/tests/fixtures/cold_browser_app.py",
        "backend/tests/test_fresh_user_no_ai_journey.py", "backend/tests/test_cold_browser_runner.py",
        "docs/delivery/FRESH_USER_NO_AI_BROWSER_JOURNEY_2026-10-09.md", "backend/uv.lock",
    )}
    summary = {"status": "failed", "source_hashes": source_hashes,
               "source_bundle_sha256": hashlib.sha256(json.dumps(source_hashes, sort_keys=True).encode()).hexdigest(),
               "scope": "Local synthetic connected workflow only", "production_proof": False}
    phase = "setup"
    env = {key: os.environ[key] for key in ("PATH", "HOME", "TMPDIR", "LANG") if key in os.environ}
    backend_port, frontend_port = _port(), _port()
    backend_url, frontend_url = f"https://127.0.0.1:{backend_port}", f"https://127.0.0.1:{frontend_port}"
    key, certificate = evidence / "loopback-key.pem", evidence / "loopback-cert.pem"
    database_url = pg.set(query={"options": f"-csearch_path={schema} -clock_timeout=4000 -cstatement_timeout=10000"}).render_as_string(hide_password=False)
    env.update({
        "APP_ENV": "test", "DATABASE_URL": database_url, "JWT_SECRET": "synthetic-browser-jwt-secret-over-thirty-two-characters",
        "FRONTEND_ORIGINS": frontend_url, "ADMIN_EMAILS": "operator.browser@example.com",
        "ANALYSIS_TASKS_MODE": "inline", "OPTIONAL_AI_GENERATION_ENABLED": "false",
        "EMPLOYER_DISCOVERY_ENABLED": "true", "EMPLOYER_AUTO_SUBMIT_ENABLED": "false",
        "EMPLOYER_SEARCH_CREDITS_PER_JOB": "2", "EMPLOYER_APPLY_CREDITS_PER_JOB": "20",
        "LIFECYCLE_EMAILS_ENABLED": "false", "RAZORPAY_CHECKOUT_ENABLED": "true",
        "RAZORPAY_ACCOUNT_APPROVED": "true", "PAYMENTS_GO_LIVE_REVIEW_COMPLETE": "true", "RAZORPAY_MODE": "test",
        "RAZORPAY_KEY_ID": "rzp_test_synthetic_browser", "RAZORPAY_KEY_SECRET": "synthetic-browser-order-secret",
        "RAZORPAY_WEBHOOK_SECRET": "synthetic-browser-webhook-secret",
        "NEXTAUTH_SECRET": "synthetic-nextauth-browser-secret-over-thirty-two-characters", "NEXTAUTH_URL": frontend_url,
        "BACKEND_URL": backend_url, "NEXT_TELEMETRY_DISABLED": "1", "PYTHONPATH": f"{ROOT}:{ROOT / 'backend'}",
        "NODE_EXTRA_CA_CERTS": str(certificate), "HIREWIZ_TEST_TLS_CERT": str(certificate),
        "HIREWIZ_TEST_TLS_KEY": str(key), "PORT": str(frontend_port),
        "COLD_BROWSER_BACKEND_PORT": str(backend_port), "COLD_BROWSER_BACKEND_URL": backend_url,
        "COLD_BROWSER_FRONTEND_URL": frontend_url, "COLD_BROWSER_PROVIDER_PROOF": str(evidence / "provider.json"),
        "COLD_BROWSER_BROWSER_PROOF": str(evidence / "browser.json"), "COLD_BROWSER_REPORT": str(evidence / "playwright.json"),
        "COLD_BROWSER_EVIDENCE": str(evidence), "COLD_BROWSER_ORIGINAL": str(evidence / "original.docx"),
        "COLD_BROWSER_CUSTOM": str(evidence / "custom.docx"),
    })
    sys.path.insert(0, str(ROOT))
    from backend.tests.expense_policy_fixtures import estimated_expense_policy
    expense = estimated_expense_policy() | {"review_status": "approved", "reviewed_by": "synthetic-local-browser-review"}
    env["HIREWIZ_EXPENSE_POLICY_JSON"] = json.dumps(expense)
    _docx(evidence / "original.docx", "Original source resume")
    _docx(evidence / "custom.docx", "Candidate-written custom resume")
    admin, processes, logs = create_engine(pg), [], []
    engine = create_engine(database_url)
    try:
        with admin.begin() as db:
            db.execute(text(f'CREATE SCHEMA "{schema}"'))
        phase = "migration"
        migration = _run_bounded([sys.executable, "-m", "alembic", "upgrade", "head"], cwd=ROOT / "backend", env=env, timeout=90)
        assert migration.returncode == 0, "Actual local migration failed"
        with engine.connect() as db:
            for table in ("users", "resumes", "employer_postings", "model_call_events", "employer_searches", "skill_coverage"):
                assert db.execute(text(f"SELECT count(*) FROM {table}")).scalar_one() == 0
        phase = "owned_loopback_tls"
        tls = _create_owned_tls(key, certificate)
        phase = "production_build"
        build = _run_bounded(["npm", "run", "build"], cwd=ROOT / "frontend", env=env, timeout=240)
        (evidence / "build.log").write_text(build.stdout + build.stderr)
        assert build.returncode == 0, "Archive production build failed; inspect local build log"
        for name, command, cwd in (
            ("api", [sys.executable, "-m", "backend.tests.fixtures.cold_browser_app"], ROOT),
            ("next", ["node", "tests/candidate-https-server.mjs"], ROOT / "frontend"),
        ):
            log = (evidence / f"{name}.log").open("w")
            logs.append(log)
            processes.append(_spawn_owned(command, cwd=cwd, env=env, stdout=log, stderr=log))
        phase = "api_startup"
        _wait(backend_url + "/api/health", processes[0], verify=tls)
        phase = "next_startup"
        _wait(frontend_url + "/register", processes[1], verify=tls)
        phase = "source_refresh"
        _seed_http(backend_url, verify=tls)
        with engine.connect() as db:
            assert db.execute(text("SELECT count(*) FROM employer_postings")).scalar_one() == 4, "Finite fixture feed must normalize four actual postings"
        phase = "browser"
        browser = _run_bounded(["npx", "playwright", "test", "-c", "playwright.cold.config.ts"], cwd=ROOT / "frontend", env=env, timeout=180)
        (evidence / "browser.log").write_text(browser.stdout + browser.stderr)
        if browser.returncode != 0:
            summary["browser_failure"] = _browser_failure_summary(evidence / "playwright.json")
        assert browser.returncode == 0, "Actual Chromium journey failed; inspect local browser log"
        phase = "sql_verification"
        proof = _sql_proof(engine)
        provider = json.loads((evidence / "provider.json").read_text())
        for name in ("forbidden_model_calls", "unexpected_external_requests", "unexpected_socket_egress", "dispatch_errors"):
            assert provider.get(name, 0) == 0
        assert provider["synthetic_checkout_orders"] == provider["synthetic_feed_gets"] == 1
        assert provider["synthetic_document_inspections"] == provider["synthetic_document_identity_requests"] == 2
        summary.update({"status": "passed", "scope": "Local production Next + Chromium + actual FastAPI/auth/SQL; finite synthetic external providers only",
                   "production_proof": False, "real_payment": False, "real_employer_contact": False,
                   "native_document_scan_and_sandbox_proof": False,
                   "cold_sql_and_processes": True, "redis_used": False, "sql": proof,
                   "browser": json.loads((evidence / "browser.json").read_text()), "provider_counts": provider})
    except BaseException as exc:
        summary.update({"status": "failed", "failed_phase": phase, "error_type": type(exc).__name__})
        raise
    finally:
        groups, cleanup_errors = [], []
        for process in reversed(processes):
            try:
                groups.append(_stop(process))
            except (AssertionError, OSError, subprocess.SubprocessError) as exc:
                cleanup_errors.append(type(exc).__name__)
        for log in logs:
            log.close()
        engine.dispose()
        with admin.begin() as db:
            db.execute(text(f'DROP SCHEMA IF EXISTS "{schema}" CASCADE'))
        admin.dispose()
        key.unlink(missing_ok=True)
        summary["cleanup"] = {"owned_processes_stopped": not cleanup_errors and len(groups) == len(processes),
                              "group_verification": groups, "errors": cleanup_errors, "own_schema_removed": True,
                              "owned_tls_private_key_removed": not key.exists()}
        if cleanup_errors:
            summary.update({"status": "failed", "failed_phase": "process_cleanup"})
        (evidence / "summary.json").write_text(json.dumps(summary, indent=2) + "\n")
        assert not cleanup_errors, "Owned process-group cleanup was not verified; inspect summary"


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--evidence", type=Path, required=True)
    parser.add_argument("--allow-disposable-postgres", action="store_true", required=True,
                        help="Allow only the unique-schema loopback disposable database contract")
    run_journey(parser.parse_args().evidence.resolve())
