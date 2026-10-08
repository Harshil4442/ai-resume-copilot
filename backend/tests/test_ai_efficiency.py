from __future__ import annotations

import io
from types import SimpleNamespace

import pytest
from backend.app.database import Base
from backend.app.domains.analysis import schemas, service, tasks
from backend.app.models import AnalysisRun, ModelCallEvent, Opportunity, Resume, User
from backend.app.routers.public_endpoints import OptimizeBulletRequest, optimize_bullet
from backend.app.services import llm_client, parsing
from backend.app.services.basic_matching import basic_match
from backend.app.services.interview_catalog import curated_interview_questions
from backend.app.services.market.skill_extractor import extract_skill_mentions
from docx import Document
from google import genai
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool
from starlette.requests import Request


def test_bounded_skills_preserve_source_spans_and_exclude_negated_mentions():
    text = 'Skills: JavaScript, C++, R, Python\nI go to work and render reports.\nNo experience with Java.\nWorked with Kubernetes.'
    mentions = extract_skill_mentions(text)
    skills = {item['skill'] for item in mentions}
    assert {'JavaScript', 'C++', 'R', 'Python', 'Kubernetes'} <= skills
    assert not {'Java', 'C', 'Go', 'Render'} & skills
    assert all(text[item['start']:item['end']] == item['text'] for item in mentions)


def test_nontechnical_catalog_and_unknown_terms_are_not_invented():
    result = basic_match(['Accounting', 'Excel'], 'Accounting position: financial analysis, Excel and unfamiliar bespoke ERP.', 'Accountant')
    assert result['matched_skills'] == ['Accounting', 'Microsoft Excel']
    assert result['missing_evidence'] == ['Financial Analysis']
    assert 'bespoke ERP' not in result['required_skills']
    assert result['score'] == pytest.approx(66.7)
    assert result['mode'] == 'basic'


def test_experience_uses_months_and_merges_concurrent_positions():
    sections = {'experience': 'Engineer Jan 2025 – Dec 2025\nIntern Jul 2025 – Dec 2025', 'education': 'Degree 2021 – 2025'}
    assert parsing.estimate_experience_years('\n'.join(sections.values()), sections) == 1.0
    assert parsing.estimate_experience_years('Intern Jan 2025 – Jun 2025', {'experience': 'Intern Jan 2025 – Jun 2025'}) == 0.5


def test_cold_resume_parse_has_no_external_generation(monkeypatch):
    def unexpected(*args, **kwargs):
        pytest.fail('Routine parsing must not call an external model')
    monkeypatch.setattr(llm_client, '_chat', unexpected)
    document = Document()
    document.add_paragraph('Jordan Candidate')
    document.add_heading('Technical Skills', 2)
    document.add_paragraph('Python, JavaScript, Docker')
    output = io.BytesIO()
    document.save(output)
    raw, sections, skills, years, contact = parsing.parse_resume_file(output.getvalue(), 'resume.docx')
    assert {'Python', 'JavaScript', 'Docker'} <= set(skills)
    assert 'Java' not in skills
    assert contact['name'] == 'Jordan Candidate'


@pytest.mark.parametrize('role,jd', [('Software Engineer', 'Python Docker systems'), ('Data Analyst', 'SQL financial analysis'), ('Accountant', 'Accounting Excel payroll'), ('New Role', '')])
def test_sparse_evidence_keeps_eight_curated_questions(role, jd):
    questions = curated_interview_questions(role, jd)
    assert len(questions) == 8
    assert len({item['question'] for item in questions}) == 8
    assert all(item['provenance'] == 'curated' and item['answer_state'] == 'evidence_needed' for item in questions)
    assert all(item['evidence_ids'] == [] for item in questions)


def test_public_diagnostics_never_invent_a_bullet_or_treat_a_version_as_impact():
    request = Request({'type': 'http', 'method': 'POST', 'path': '/', 'headers': [], 'client': ('127.0.0.1', 1), 'app': SimpleNamespace(state=SimpleNamespace())})
    original = 'Worked with Python 3.12 during 2025.'
    # Unwrap only the rate-limit decorator; the endpoint remains the actual implementation.
    result = optimize_bullet.__wrapped__(OptimizeBulletRequest(bullet_text=original), request)
    assert result.recommended_bullet == original
    assert result.metrics_present is False
    assert result.provenance == 'local_rules'
    quantified = optimize_bullet.__wrapped__(OptimizeBulletRequest(bullet_text='Reduced latency by 20%.'), request)
    assert quantified.metrics_present is True


def _database():
    engine = create_engine('sqlite://', connect_args={'check_same_thread': False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, autoflush=False)
    with factory() as db:
        db.add(User(id=1, email='synthetic@example.com', ai_credits=100, tier='free'))
        db.add(Resume(id=1, user_id=1, original_filename='resume.docx', raw_text='Python engineer', skills=['Python'], sections={'experience': 'Developed Python services'}, experience_years=1))
        db.add(Opportunity(id='opp_test', user_id=1, resume_id=1, title='Software Engineer', company='Example', job_description='Build Python and Docker services for our customers.', job_snapshot={}))
        db.commit()
    return engine, factory


def _create(factory, operation='job_match', key='request-key-001', mode=None):
    input_payload = {'resume_id': 1, 'job_description': 'Build Python and Docker services for our customers.'} if operation == 'job_match' else {}
    if mode:
        input_payload['mode'] = mode
    with factory() as db:
        run, created = service.create_run(db, user_id=1, payload=schemas.AnalysisRunCreate(operation=operation, opportunity_id='opp_test', input=input_payload), header_idempotency_key=key)
        return run.id, created


def test_basic_and_curated_run_with_empty_cache_and_llm_disabled(monkeypatch):
    engine, factory = _database()
    monkeypatch.setattr(tasks, 'SessionLocal', factory)
    monkeypatch.setattr(llm_client, '_chat', lambda *args: pytest.fail('Unexpected generation'))
    try:
        match_id, _ = _create(factory)
        interview_id, _ = _create(factory, 'interview_questions', 'request-interview-001')
        assert tasks.process_analysis_run(match_id) == 'succeeded'
        assert tasks.process_analysis_run(interview_id) == 'succeeded'
        with factory() as db:
            match = db.get(AnalysisRun, match_id)
            assert match.result_payload['mode'] == 'basic'
            assert match.result_payload['dimensions'] == []
            assert match.result_payload['true_gaps'] == ['Docker']
            assert match.generation_attempt_count == 0
            assert len(db.get(AnalysisRun, interview_id).result_payload['questions']) == 8
            assert db.query(ModelCallEvent).count() == 0
    finally:
        engine.dispose()


def test_dependency_reuse_avoids_charge_and_invalidates_changed_resume(monkeypatch):
    engine, factory = _database()
    try:
        first, created = _create(factory)
        assert created
        same, created = _create(factory, key='request-key-other')
        assert (same, created) == (first, False)
        with factory() as db:
            assert db.get(User, 1).ai_credits == 99
            resume = db.get(Resume, 1)
            resume.skills = ['Python', 'Docker']
            db.commit()
        changed, created = _create(factory, key='request-changed-001')
        assert created and changed != first
        with factory() as db:
            assert db.get(User, 1).ai_credits == 98
    finally:
        engine.dispose()


def test_total_attempt_budget_covers_fallback_and_a_second_logical_call(monkeypatch, persisted_model_budget):
    calls = []
    class Models:
        def generate_content(self, *, model, **kwargs):
            calls.append(model)
            if len(calls) in {1, 3}:
                raise RuntimeError('503 temporarily unavailable')
            return SimpleNamespace(text='result')
    monkeypatch.setenv('LLM_API_KEY', 'synthetic')
    monkeypatch.setattr(llm_client, 'LLM_MODEL', 'gemini-3.6-flash')
    monkeypatch.setattr(genai, 'Client', lambda **kwargs: SimpleNamespace(models=Models()))
    assert llm_client._chat([{'role': 'user', 'content': 'synthetic'}]) == 'result'
    with pytest.raises(llm_client.LLMProviderError) as error:
        llm_client._chat([{'role': 'user', 'content': 'repair synthetic'}])
    assert len(calls) == 3
    assert tasks._retryable(error.value) is False


def test_persisted_attempts_survive_domain_failure_and_worker_redelivery(monkeypatch):
    engine, factory = _database()
    calls = []
    class Models:
        def generate_content(self, *, model, **kwargs):
            calls.append(model)
            raise RuntimeError('503 temporary outage')
    monkeypatch.setattr(tasks, 'SessionLocal', factory)
    monkeypatch.setenv('LLM_API_KEY', 'synthetic')
    monkeypatch.setattr(llm_client, 'LLM_MODEL', 'gemini-3.6-flash')
    monkeypatch.setattr(genai, 'Client', lambda **kwargs: SimpleNamespace(models=Models()))
    try:
        run_id, _ = _create(factory, mode='enhanced')
        assert tasks.process_analysis_run(run_id) == 'failed'
        assert tasks.process_analysis_run(run_id) == 'failed'
        assert len(calls) == 3
        with factory() as db:
            run = db.get(AnalysisRun, run_id)
            assert run.generation_attempt_count == 3
            events = db.query(ModelCallEvent).filter_by(analysis_run_id=run_id).all()
            assert len(events) == 3
            assert all(event.status == 'failed' and event.tokens_estimated for event in events)
            assert run.model_cost_reserved_micros == sum(event.reserved_cost_micros for event in events) > 0
            assert run.model_cost_settled_micros == 0
            assert db.get(User, 1).ai_credits == 100
    finally:
        engine.dispose()


def test_reused_client_key_remains_bound_to_its_original_inputs():
    from fastapi import HTTPException

    engine, factory = _database()
    try:
        original, _ = _create(factory)
        assert _create(factory, key='coalesced-client-key') == (original, False)
        with factory() as db:
            resume = db.get(Resume, 1)
            resume.skills = ['Python', 'Docker']
            db.commit()
        with pytest.raises(HTTPException) as error:
            _create(factory, key='coalesced-client-key')
        assert error.value.status_code == 409
    finally:
        engine.dispose()


def test_enhanced_mode_preserves_generation_and_records_actual_model_and_tokens(monkeypatch):
    import json

    engine, factory = _database()
    calls = []
    output = {
        'extracted_jd_skills': ['python', 'docker'],
        'skill_analysis': [{'via_skill': 'python', 'jd_skill': 'python', 'coverage': 1.0}],
        'dimensions': [{'name': 'Experience Level Fit', 'score': 70, 'feedback': 'Review role expectations.'}],
        'fit_summary': 'Python work is documented; Docker evidence is missing. Review duties and seniority.',
        'improvement_tips': ['Provide a verified Docker example if you have one.'],
    }
    class Models:
        def generate_content(self, *, model, contents, config):
            calls.append(model)
            assert 'EDUCATION:' in contents[-1]['parts'][0]['text']
            if len(calls) == 1:
                raise RuntimeError('404 model not found')
            return SimpleNamespace(text=json.dumps(output), usage_metadata=SimpleNamespace(prompt_token_count=123, candidates_token_count=45, thoughts_token_count=0))
    monkeypatch.setattr(tasks, 'SessionLocal', factory)
    monkeypatch.setenv('LLM_API_KEY', 'synthetic')
    monkeypatch.setattr(llm_client, 'LLM_MODEL', 'gemini-3.6-flash')
    monkeypatch.setattr(genai, 'Client', lambda **kwargs: SimpleNamespace(models=Models()))
    try:
        run_id, _ = _create(factory, mode='enhanced')
        assert tasks.process_analysis_run(run_id) == 'succeeded'
        assert len(calls) == 2
        with factory() as db:
            run = db.get(AnalysisRun, run_id)
            assert run.result_payload['mode'] == 'enhanced'
            assert run.model == calls[-1]
            assert run.generation_attempt_count == 2
            events = db.query(ModelCallEvent).filter_by(analysis_run_id=run_id).order_by(ModelCallEvent.attempt_number).all()
            assert len(events) == 2
            assert events[-1].model == calls[-1]
            assert events[-1].input_tokens == 123 and events[-1].output_tokens == 45
            assert events[-1].tokens_estimated is False
    finally:
        engine.dispose()


@pytest.mark.parametrize('outcome,expected_units,expected_state', [('success', 1, 'completed'), ('same', 0, 'unchanged'), ('failure', 0, 'failed')])
def test_optional_upload_enrichment_is_explicit_refundable_and_preserves_original(monkeypatch, outcome, expected_units, expected_state):
    import asyncio
    import json

    from backend.app.routers.resume import parse_resume
    from starlette.datastructures import Headers, UploadFile

    engine, factory = _database()
    calls = []
    class Models:
        def generate_content(self, **kwargs):
            calls.append(kwargs)
            if outcome == 'failure':
                raise RuntimeError('400 invalid provider request')
            skills = ['Python', 'BespokeCRM'] if outcome == 'success' else ['Python']
            return SimpleNamespace(text=json.dumps(skills), usage_metadata=SimpleNamespace(prompt_token_count=40, candidates_token_count=10, thoughts_token_count=0))
    monkeypatch.setenv('LLM_API_KEY', 'synthetic')
    monkeypatch.setattr(llm_client, 'LLM_MODEL', 'gemini-3.6-flash')
    monkeypatch.setattr(genai, 'Client', lambda **kwargs: SimpleNamespace(models=Models()))
    document = Document()
    document.add_paragraph('Jordan Candidate')
    document.add_heading('Skills', 2)
    document.add_paragraph('Python, BespokeCRM')
    output = io.BytesIO()
    document.save(output)
    original = output.getvalue()
    request = Request({'type': 'http', 'method': 'POST', 'path': '/resume/parse', 'headers': [], 'client': ('127.0.0.1', 1)})
    try:
        with factory() as db:
            upload = UploadFile(io.BytesIO(original), filename='resume.docx', headers=Headers({'content-type': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'}))
            response = asyncio.run(parse_resume.__wrapped__(request, upload, True, db, db.get(User, 1)))
            assert response.enrichment_state == expected_state
            assert response.enrichment_units == expected_units
            assert db.get(User, 1).ai_credits == 100 - expected_units
            assert db.get(Resume, response.resume_id).source_document == original
            assert ('BespokeCRM' in response.skills) == (outcome == 'success')
            run = db.query(AnalysisRun).filter_by(operation='resume_enrichment').one()
            assert run.generation_attempt_count == 1
            assert db.query(ModelCallEvent).filter_by(analysis_run_id=run.id).count() == 1
            if outcome != 'success':
                assert response.warnings
    finally:
        engine.dispose()


def test_legacy_and_enrichment_usage_runs_remain_readable_in_history():
    engine, factory = _database()
    try:
        from backend.app.services.guardrails import billable_operation
        with factory() as db:
            with billable_operation(user_id=1, db=db, operation='resume_enrichment', amount=1) as run:
                run_id = run.id
        with factory() as db:
            response = schemas.AnalysisRunResponse.model_validate(db.get(AnalysisRun, run_id))
            assert response.operation == 'resume_enrichment'
            assert response.status == 'succeeded'
    finally:
        engine.dispose()
