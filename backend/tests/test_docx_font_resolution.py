"""Declared/installed/rendered font controls, not native Linux rendering proof."""
from __future__ import annotations

import hashlib
import importlib.util
import json
import sys
import tempfile
import types
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

ROOT = Path(__file__).resolve().parents[2]
SERVICES = ROOT / "backend/app/services"
PREFIX = "hirewiz_font_control"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def load(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


def saved(document):
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def source(text="Candidate Example", family="Courier Prime"):
    document = Document()
    document.styles["Normal"].font.name = family
    document.add_paragraph(text)
    return document


def alter_part(content, name, change):
    output = BytesIO()
    with ZipFile(BytesIO(content)) as original, ZipFile(output, "w") as archive:
        for entry in original.infolist():
            raw = original.read(entry)
            if entry.filename == name:
                tree = etree.fromstring(raw)
                change(tree)
                raw = etree.tostring(tree)
            archive.writestr(entry, raw)
    return output.getvalue()


class FontControls(unittest.TestCase):
    def setUp(self):
        self.modules = patch.dict(sys.modules)
        self.modules.start()
        package = types.ModuleType(PREFIX)
        package.__path__ = [str(SERVICES)]
        sys.modules[PREFIX] = package
        self.core = load(PREFIX + ".docx_layout_validation", SERVICES / "docx_layout_validation.py")
        self.fonts = load(PREFIX + ".docx_fonts", SERVICES / "docx_fonts.py")
        self.renderer = load("font_control_renderer", ROOT / "infra/document-processing/docx_renderer.py")
        sys.modules["app.services.docx_fonts"] = self.fonts
        sys.modules["app.services.docx_layout_validation"] = self.core

    def tearDown(self):
        self.modules.stop()

    def fonts_of(self, document):
        return self.fonts.document_fonts(saved(document))

    def assert_refused(self, action):
        with self.assertRaises(self.core.DocxLayoutError) as error:
            action()
        self.assertEqual(error.exception.reason, "unsupported_source_font")

    def test_only_effective_used_fonts_are_required(self):
        document = source()
        unused = document.styles.add_style("Unused missing face", WD_STYLE_TYPE.CHARACTER)
        unused.font.name = "Missing Font"
        use = self.fonts_of(document)
        self.assertEqual(set(use.required), {("Courier Prime", False, False)})
        self.assertIn(ord("C"), use.required[("Courier Prime", False, False)])

    def test_paragraph_style_chain_and_direct_character_style(self):
        document = source()
        base = document.styles.add_style("Custom base", WD_STYLE_TYPE.PARAGRAPH)
        base.base_style = document.styles["Normal"]
        base.font.bold = True
        derived = document.styles.add_style("Custom paragraph", WD_STYLE_TYPE.PARAGRAPH)
        derived.base_style = base
        derived.font.italic = True
        character = document.styles.add_style("Custom character", WD_STYLE_TYPE.CHARACTER)
        character.font.name = "Liberation Serif"
        paragraph = document.paragraphs[0]
        paragraph.style = derived
        paragraph.runs[0].style = character
        use = self.fonts_of(document)
        self.assertEqual(set(use.required), {("Liberation Serif", True, True)})
        paragraph.runs[0].bold = False
        self.assertEqual(set(self.fonts_of(document).required), {("Liberation Serif", False, True)})

    def test_theme_and_direct_font_override(self):
        document = Document()
        run = document.add_paragraph("Candidate").runs[0]
        self.assertEqual(set(self.fonts_of(document).required), {("Cambria", False, False)})
        run.font.name = "Courier Prime"
        self.assertEqual(set(self.fonts_of(document).required), {("Courier Prime", False, False)})
        run._r.rPr.rFonts.set(qn("w:asciiTheme"), "majorAscii")
        self.assertEqual(set(self.fonts_of(document).required), {("Calibri", False, False)})

    def test_header_footer_and_table_use_inherited_fonts(self):
        document = source()
        document.sections[0].header.paragraphs[0].text = "Header"
        document.sections[0].footer.paragraphs[0].text = "Footer"
        document.add_table(1, 1).cell(0, 0).text = "Table text"
        use = self.fonts_of(document)
        self.assertEqual(set(use.required), {("Courier Prime", False, False)})
        text = ["".join(char for char, _ in p) for p in use.paragraphs]
        self.assertCountEqual(text, ["CandidateExample", "Header", "Footer", "Tabletext"])

    def test_bold_style_toggle_and_direct_italic(self):
        document = source()
        document.styles["Normal"].font.bold = True
        child = document.styles.add_style("Toggle", WD_STYLE_TYPE.PARAGRAPH)
        child.base_style = document.styles["Normal"]
        child.font.bold = True
        document.paragraphs[0].style = child
        document.paragraphs[0].runs[0].italic = True
        self.assertEqual(set(self.fonts_of(document).required), {("Courier Prime", False, True)})

    def test_missing_theme_style_cycle_and_query_syntax_refuse(self):
        document = source()
        document.styles["Normal"].base_style = document.styles["Normal"]
        self.assert_refused(lambda: self.fonts_of(document))
        self.assert_refused(lambda: self.fonts_of(source(family="Courier Prime:weight=bold")))
        document = Document()
        document.add_paragraph("Hello")
        raw = alter_part(saved(document), "word/_rels/document.xml.rels", lambda tree: [tree.remove(row) for row in list(tree) if row.get("Type", "").endswith("/theme")])
        self.assert_refused(lambda: self.fonts.document_fonts(raw))

    def test_ambiguous_visible_variants_and_generated_text_refuse(self):
        document = source()
        document.paragraphs[0].runs[0].font.small_caps = True
        self.assert_refused(lambda: self.fonts_of(document))
        document = source()
        document.add_paragraph("Numbered", "List Number")
        self.assert_refused(lambda: self.fonts_of(document))
        document = source()
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        document.paragraphs[0]._p.append(field)
        self.assert_refused(lambda: self.fonts_of(document))

    def test_effective_transform_is_checked_after_direct_run_override(self):
        for name in ("caps", "smallCaps", "vanish", "webHidden", "specVanish"):
            with self.subTest(property=name):
                document = source("Ordinary mixed case")
                inherited = OxmlElement("w:" + name)
                document.styles["Normal"].element.get_or_add_rPr().append(inherited)
                run = document.paragraphs[0].runs[0]
                override = OxmlElement("w:" + name)
                override.set(qn("w:val"), "false")
                run._r.get_or_add_rPr().append(override)
                self.assertEqual(set(self.fonts_of(document).required), {("Courier Prime", False, False)})
                run._r.rPr.remove(override)
                self.assert_refused(lambda document=document: self.fonts_of(document))

    def test_style_transform_toggles_resolve_before_effective_refusal(self):
        for attribute in ("all_caps", "small_caps", "hidden"):
            with self.subTest(property=attribute):
                document = source("Ordinary mixed case")
                setattr(document.styles["Normal"].font, attribute, True)
                child = document.styles.add_style("Toggle transformation", WD_STYLE_TYPE.PARAGRAPH)
                child.base_style = document.styles["Normal"]
                setattr(child.font, attribute, True)
                document.paragraphs[0].style = child
                self.assertEqual(set(self.fonts_of(document).required), {("Courier Prime", False, False)})
                # An off style toggle keeps the inherited transformation on;
                # only direct run false values are absolute overrides.
                setattr(child.font, attribute, False)
                self.assert_refused(lambda document=document: self.fonts_of(document))

    def test_inactive_table_condition_font_is_not_a_required_face(self):
        for named in (True, False):
            with self.subTest(named_look=named):
                document = source()
                style = document.styles.add_style("First-row font", WD_STYLE_TYPE.TABLE)
                condition = OxmlElement("w:tblStylePr")
                condition.set(qn("w:type"), "firstRow")
                properties = OxmlElement("w:rPr")
                fonts = OxmlElement("w:rFonts")
                fonts.set(qn("w:ascii"), "Missing Font")
                properties.append(fonts)
                condition.append(properties)
                style.element.append(condition)
                table = document.add_table(1, 1)
                table.style = style
                table.cell(0, 0).text = "Ordinary supported text"
                look = table._tbl.tblPr.find(qn("w:tblLook"))
                look.attrib.clear()
                look.set(qn("w:val"), "0020" if named else "0000")
                if named:
                    # Named attributes supersede the legacy bitmask.
                    look.set(qn("w:firstRow"), "0")
                self.assertEqual(set(self.fonts_of(document).required), {("Courier Prime", False, False)})
                look.set(qn("w:firstRow"), "1")
                self.assert_refused(lambda document=document: self.fonts_of(document))
                look.set(qn("w:firstRow"), "0")
                override = OxmlElement("w:cnfStyle")
                override.set(qn("w:firstRow"), "1")
                table.cell(0, 0).paragraphs[0]._p.get_or_add_pPr().append(override)
                # A per-cell/row/paragraph override needs a separate resolver.
                self.assert_refused(lambda document=document: self.fonts_of(document))

    def test_disabled_table_bands_and_unknown_named_flags(self):
        document = source()
        table = document.add_table(1, 1)
        look = table._tbl.tblPr.find(qn("w:tblLook"))
        look.attrib.clear()
        look.set(qn("w:val"), "0600")
        self.assertTrue({"firstRow", "lastRow", "firstCol", "lastCol", "band1Horz", "band2Horz", "band1Vert", "band2Vert"}
                        <= self.fonts._inactive_table_conditions(table._tbl))
        look.set(qn("w:noHBand"), "true")
        self.assertEqual(self.fonts._inactive_table_conditions(table._tbl), {"band1Horz", "band2Horz"})

    def test_source_glyph_slots_and_missing_character_mapping(self):
        document = source("Résumé • 2026")
        self.assertEqual(set(self.fonts_of(document).required), {("Courier Prime", False, False)})
        self.assert_refused(lambda: self.fonts_of(source("Candidate 😀")))

    def test_equal_wrong_rendered_font_is_rejected(self):
        use = self.fonts_of(source("ABC"))
        expected = {("Courier Prime", False, False): "CourierPrime-Regular"}
        glyphs = [{"text": char, "font": "LiberationMono"} for char in "ABC"]
        for _before_or_after in range(2):
            self.assert_refused(lambda: self.fonts.verify_rendered_fonts(use, glyphs, expected))
        for row in glyphs:
            row["font"] = "CourierPrime-Regular"
        self.fonts.verify_rendered_fonts(use, glyphs, expected)

    def test_complete_layout_validator_rejects_same_fallback_in_both_pdfs(self):
        fixtures = load("font_control_pdf_fixtures", ROOT / "backend/tests/test_docx_layout_validation.py")
        document = source("Taylor Example")
        document.add_paragraph(fixtures.ORIGINAL)
        document.add_paragraph(fixtures.PROTECTED)
        original = saved(document)
        candidate = alter_part(original, "word/document.xml", lambda tree: setattr(list(tree.iter(W + "t"))[1], "text", fixtures.REPLACEMENT))
        edit = {"unit_id": "docx:" + hashlib.sha256(original).hexdigest()[:16] + ":p1:" + hashlib.sha256(fixtures.ORIGINAL.encode()).hexdigest()[:12],
                "original_text": fixtures.ORIGINAL, "replacement_text": fixtures.REPLACEMENT}
        identity = self.core.RendererIdentity("registry.test/worker@sha256:" + "a" * 64, "b" * 64, "Synthetic test only", "c" * 64)
        contract = (self.fonts.document_fonts(original), self.fonts.document_fonts(candidate),
                    {("Courier Prime", False, False): "CourierPrime-Regular"})
        self.assert_refused(lambda: self.core.validate_docx_layout(
            source_docx=original, candidate_docx=candidate, edits=[edit],
            before_pdf=fixtures.pdf(), after_pdf=fixtures.pdf(text=fixtures.REPLACEMENT),
            renderer=identity, font_contract=contract))

    def test_actual_bold_face_and_unmapped_extra_text_are_checked(self):
        document = source("ABC")
        document.paragraphs[0].add_run("D").bold = True
        use = self.fonts_of(document)
        expected = {("Courier Prime", False, False): "CourierPrime-Regular", ("Courier Prime", True, False): "CourierPrime-Bold"}
        glyphs = [{"text": char, "font": "CourierPrime-Regular"} for char in "ABCD"]
        self.assert_refused(lambda: self.fonts.verify_rendered_fonts(use, glyphs, expected))
        glyphs[-1]["font"] = "CourierPrime-Bold"
        self.fonts.verify_rendered_fonts(use, glyphs, expected)
        glyphs.append({"text": "X", "font": "CourierPrime-Regular"})
        self.assert_refused(lambda: self.fonts.verify_rendered_fonts(use, glyphs, expected))

    def test_duplicate_paragraphs_with_different_faces_are_ambiguous(self):
        document = source("Same")
        document.add_paragraph("Same").runs[0].font.name = "Liberation Serif"
        use = self.fonts_of(document)
        faces = {("Courier Prime", False, False): "CourierPrime-Regular",
                 ("Liberation Serif", False, False): "LiberationSerif"}
        for order in (("LiberationSerif", "CourierPrime-Regular"),
                      ("CourierPrime-Regular", "LiberationSerif")):
            glyphs = [{"text": char, "font": font} for font in order for char in "Same"]
            # Even the unswapped case refuses until part/occurrence placement
            # can be bound; finding both expected faces somewhere is not proof.
            self.assert_refused(lambda glyphs=glyphs: self.fonts.verify_rendered_fonts(use, glyphs, faces))

    def test_identical_font_repeats_preserve_source_occurrences(self):
        document = source("Same")
        document.add_paragraph("Same")
        use = self.fonts_of(document)
        faces = {("Courier Prime", False, False): "CourierPrime-Regular"}
        glyphs = [{"text": char, "font": "CourierPrime-Regular"} for char in "SameSame"]
        self.fonts.verify_rendered_fonts(use, glyphs, faces)
        self.assert_refused(lambda: self.fonts.verify_rendered_fonts(use, glyphs[:4], faces))
        self.fonts.verify_rendered_fonts(use, glyphs + glyphs[:4], faces)

    def test_exact_installed_face_and_character_coverage(self):
        filename = "/usr/share/fonts/truetype/hirewiz-courier-prime/CourierPrime-Regular.ttf"
        required = {("Courier Prime", False, False): {ord("A"), ord("é")}}
        response = f"Courier Prime\nCourierPrime-Regular\n{filename}\n80\n0\n20-7e a0-ff\n".encode()
        with tempfile.TemporaryDirectory() as folder:
            manifest = Path(folder) / "manifest.json"
            manifest.write_text(json.dumps({"files": [{"path": filename}]}))
            with patch.object(self.renderer, "MANIFEST", manifest), patch.object(self.renderer, "_font_query", return_value=response) as query:
                self.assertEqual(self.renderer._resolve_fonts(required, deadline=10), {("Courier Prime", False, False): "CourierPrime-Regular"})
                self.assertEqual(query.call_args.args, ("Courier Prime:weight=regular:slant=roman",))
            for mutated in (response.replace(b"Courier Prime\n", b"Liberation Mono\n"),
                            response.replace(b"80\n0\n", b"200\n0\n"),
                            response.replace(b"80\n0\n", b"80\n100\n"),
                            response.replace(b" a0-ff", b""),
                            response.replace(filename.encode(), b"/tmp/injected.ttf")):
                with patch.object(self.renderer, "MANIFEST", manifest), patch.object(self.renderer, "_font_query", return_value=mutated):
                    self.assert_refused(lambda: self.renderer._resolve_fonts(required, deadline=10))

    def test_missing_source_font_stops_before_writer_and_cannot_change_original(self):
        original = saved(source())
        before = bytes(original)
        with patch.object(self.renderer, "font_manifest", return_value=("fixture", "a" * 64)), \
                patch.object(self.renderer, "_resolve_fonts", side_effect=self.core.DocxLayoutError("unsupported_source_font")), \
                patch.object(self.renderer, "_render", side_effect=AssertionError("Writer must not start")):
            self.assert_refused(lambda: self.renderer.validate_rendered_edit(original, original, [], image="fixture", policy="fixture", deadline=10))
        self.assertEqual(original, before)

    def test_candidate_native_corpus_preserves_fallback_negative(self):
        builder = load("font_control_image_builder", ROOT / "backend/scripts/document_test_image.py")
        rows = builder.candidate_docx_layout_corpus("sha256:" + "a" * 64)
        self.assertEqual([row[0] for row in rows], ["styled_multipage_fit", "rendered_reflow_refusal", "unsupported_source_font_refusal"])
        import base64
        families = []
        for _name, payload, _binding in rows:
            raw = base64.b64decode(json.loads(payload)["content_base64"])
            use = self.fonts.document_fonts(raw)
            families.append({key[0] for key in use.required})
            self.assertEqual(len(use.required), 3)  # regular, bold and italic
        self.assertEqual(families, [{"Courier Prime"}, {"Courier Prime"}, {"Cambria"}])


if __name__ == "__main__":
    unittest.main()
