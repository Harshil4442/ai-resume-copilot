from __future__ import annotations

from contextlib import closing
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pdfplumber
import pypdfium2 as pdfium
import pypdfium2.raw as raw
import pytest
import reportlab
from backend.app.services.resume_layout import (
    ResumeLayoutError,
    _apply_docx,
    apply_source_edits,
    extract_source_units,
)
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from lxml import etree
from PIL import Image, ImageChops, ImageDraw
from pypdf import PdfReader, PdfWriter
from pypdf.generic import (
    ArrayObject,
    DecodedStreamObject,
    DictionaryObject,
    NameObject,
    NumberObject,
)
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas

ORIGINAL = "Built scalable systems with Python and SQL."
REPLACEMENT = "Built systems with Python and SQL."
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _pdf(*, subset: bool = False, two_pages: bool = False) -> bytes:
    output = BytesIO()
    canvas = Canvas(output, pagesize=(612, 792), invariant=True)
    canvas.setFont("Times-Bold", 18)
    canvas.drawCentredString(306, 752, "Taylor Example")
    canvas.setFont("Times-Roman", 10)
    canvas.setFillColorRGB(0.9, 0.01, 0.55)
    canvas.drawCentredString(306, 732, "taylor@example.com | Bengaluru")
    canvas.linkURL("https://example.test/taylor", (190, 728, 422, 741), thickness=0)
    canvas.setFillColorRGB(0, 0, 0)
    canvas.setFont("Times-Bold", 12)
    canvas.drawString(48, 704, "EXPERIENCE")
    canvas.line(48, 698, 564, 698)
    canvas.drawString(48, 682, "Example Company")
    canvas.setFont("Times-Italic", 10)
    canvas.drawRightString(564, 682, "2020 - 2024")
    font = "Times-Roman"
    if subset:
        font = "LayoutTestSubset"
        if font not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(
                TTFont(font, str(Path(reportlab.__file__).parent / "fonts" / "Vera.ttf"))
            )
    canvas.setFont(font, 11)
    canvas.drawString(60, 662, ORIGINAL)
    canvas.drawString(60, 644, "Reduced API latency by 40% through focused profiling.")
    canvas.saveState()
    canvas.translate(500, 550)
    canvas.rotate(90)
    canvas.drawString(0, 0, "Built scalable systems with Python and SQL.")
    canvas.restoreState()
    if two_pages:
        canvas.showPage()
        canvas.setFont("Times-Roman", 11)
        canvas.drawString(48, 662, "Delivered reliable services with automated quality checks.")
        canvas.setFont("Times-Italic", 12)
        canvas.setFillColorRGB(0.15, 0.4, 0.65)
        canvas.drawString(48, 640, "Original portfolio link")
        canvas.linkURL("https://example.test/portfolio", (48, 636, 190, 652), thickness=0)
        canvas.drawImage(ImageReader(Image.new("RGB", (20, 20), "#485c11")), 540, 700, 20, 20)
    canvas.save()
    return output.getvalue()


def _edit(unit: dict, replacement: str = REPLACEMENT) -> dict:
    return {
        "unit_id": unit["unit_id"],
        "original_text": unit["text"],
        "replacement_text": replacement,
        "evidence_ids": ["synthetic-evidence"],
        "reason": "Concise wording using the same source facts.",
    }


def _first_unit(source: bytes, kind: str) -> dict:
    return next(unit for unit in extract_source_units(source, kind) if unit["text"] == ORIGINAL)


def _pdf_snapshot(source: bytes) -> list[dict]:
    pages = []
    with pdfium.PdfDocument(source) as document:
        for index in range(len(document)):
            with closing(document[index]) as page, closing(page.get_textpage()) as textpage:
                with closing(page.render(scale=2)) as bitmap:
                    image = bitmap.to_pil().convert("RGB").copy()
                objects = []
                for obj in page.get_objects(filter=[raw.FPDF_PAGEOBJ_TEXT], textpage=textpage):
                    objects.append(
                        (
                            obj.extract(),
                            obj.get_bounds(),
                            obj.get_font().get_base_name(),
                            obj.get_font_size(),
                            obj.get_matrix().get(),
                        )
                    )
                pages.append(
                    {
                        "text": textpage.get_text_range(),
                        "image": image,
                        "objects": objects,
                        "size": page.get_size(),
                    }
                )
    return pages


def test_pdf_native_edit_preserves_font_geometry_page_count_and_unedited_pixels():
    source = _pdf(two_pages=True)
    unit = _first_unit(source, "pdf")
    result = apply_source_edits(source, "pdf", [_edit(unit)])
    before, after = _pdf_snapshot(source), _pdf_snapshot(result)
    assert len(before) == len(after) == 2
    assert REPLACEMENT in after[0]["text"]
    # The rotated, protected source occurrence remains; the edited occurrence
    # is replaced, not concealed under an ATS-visible overlay.
    assert after[0]["text"].count(ORIGINAL) == before[0]["text"].count(ORIGINAL) - 1
    old = next(obj for obj in before[0]["objects"] if obj[0] == ORIGINAL)
    new = next(obj for obj in after[0]["objects"] if obj[0] == REPLACEMENT)
    assert new[2:] == old[2:]
    assert new[1][2] - new[1][0] <= old[1][2] - old[1][0] + 0.1
    assert new[1][3] - new[1][1] <= old[1][3] - old[1][1] + 0.1
    difference = ImageChops.difference(before[0]["image"], after[0]["image"])
    left, bottom, right, top = old[1]
    draw = ImageDraw.Draw(difference)
    draw.rectangle(
        (
            int(left * 2) - 3,
            int((792 - top) * 2) - 3,
            int(right * 2) + 3,
            int((792 - bottom) * 2) + 3,
        ),
        fill=(0, 0, 0),
    )
    assert difference.getbbox() is None
    assert before[1]["text"] == after[1]["text"]
    assert before[1]["image"].tobytes() == after[1]["image"].tobytes()
    with pdfplumber.open(BytesIO(result)) as saved:
        assert [annotation["uri"] for page in saved.pages for annotation in page.annots] == [
            "https://example.test/taylor",
            "https://example.test/portfolio",
        ]


def test_pdf_candidates_protect_headers_contact_dates_and_rotated_text():
    units = extract_source_units(_pdf(), "pdf")
    assert [unit["text"] for unit in units] == [
        ORIGINAL,
        "Reduced API latency by 40% through focused profiling.",
    ]
    assert all(unit["preserve_numbers"] for unit in units)
    assert units[0]["font"] == "Times-Roman"
    assert units[0]["width"] > 0
    assert extract_source_units(_pdf(), "pdf") == units


def test_pdf_subset_font_roundtrips_supported_text_and_rejects_missing_glyph():
    source = _pdf(subset=True)
    unit = _first_unit(source, "pdf")
    result = apply_source_edits(source, "pdf", [_edit(unit)])
    assert REPLACEMENT in _pdf_snapshot(result)[0]["text"]
    with pytest.raises(ResumeLayoutError, match="glyph"):
        apply_source_edits(source, "pdf", [_edit(unit, "Built systems with Ж Python and SQL.")])


def test_pdf_native_character_positions_retain_source_word_spacing(monkeypatch):
    import backend.app.services.resume_layout as layout

    source = _pdf()
    original_width = raw.FPDFFont_GetGlyphWidth

    def zero_space_width(font, glyph, size, width):
        result = original_width(font, glyph, size, width)
        if glyph == ord(" "):
            width._obj.value = 0.0
        return result

    # Simulate a TeX Type1 subset without a space glyph. The real native text
    # writer and position APIs execute and the saved PDF is re-read.
    monkeypatch.setattr(raw, "FPDFFont_GetGlyphWidth", zero_space_width)
    result = apply_source_edits(source, "pdf", [_edit(_first_unit(source, "pdf"))])
    with pdfium.PdfDocument(result) as document:
        with closing(document[0]) as page, closing(page.get_textpage()) as textpage:
            obj = next(
                obj
                for obj in page.get_objects(filter=[raw.FPDF_PAGEOBJ_TEXT], textpage=textpage)
                if obj.extract() == REPLACEMENT
            )
            assert all(gap > 2 for gap in layout._word_gaps(obj, textpage))


def test_pdf_rejects_overflow_and_leaves_original_unchanged():
    source = _pdf()
    unit = _first_unit(source, "pdf")
    with pytest.raises(ResumeLayoutError, match="slot") as rejected:
        apply_source_edits(
            source, "pdf", [_edit(unit, "Built " + "wide systems " * 20)]
        )
    assert rejected.value.unit_id == unit["unit_id"]
    assert "Native width limit is" in rejected.value.repair_hint
    assert "Character count alone cannot prove font fit" in rejected.value.repair_hint
    assert ORIGINAL not in rejected.value.repair_hint
    assert _first_unit(source, "pdf")["text"] == ORIGINAL


def _pdf_line(text: str, *, obstruction: bool = False) -> bytes:
    output = BytesIO()
    canvas = Canvas(output, pagesize=(612, 792), invariant=True)
    canvas.setFont("Times-Roman", 11)
    canvas.drawString(60, 620, text)
    if obstruction:
        # Below the original ink, but inside the font's descender envelope.
        canvas.rect(60, 617.35, 200, 0.4, stroke=0, fill=1)
    canvas.save()
    return output.getvalue()


@pytest.mark.parametrize(
    "original,replacement,changed_edge",
    [
        ("Created stable interfaces with SQL.", "Built stable interfaces with SQL.", "left"),
        ("Built stable interfaces with SQL.", "Built easy interfaces with SQL.", "bottom"),
    ],
)
def test_pdf_safe_edits_use_font_envelope_instead_of_original_ink_bounds(
    original, replacement, changed_edge
):
    source = _pdf_line(original)
    unit = extract_source_units(source, "pdf")[0]
    before = _pdf_snapshot(source)[0]
    result = apply_source_edits(source, "pdf", [_edit(unit, replacement)])
    after = _pdf_snapshot(result)[0]
    old, new = before["objects"][0], after["objects"][0]
    assert old[0] == original and new[0] == replacement
    assert new[2:] == old[2:]  # Font, font size and baseline/matrix are unchanged.
    assert new[1][2] <= old[1][2] + 0.1
    if changed_edge == "left":
        # B has a smaller left ink bearing than C at the exact same baseline.
        assert new[1][0] < old[1][0] - 0.1
    else:
        assert "y" not in original and "y" in replacement
        assert new[1][1] < old[1][1] - 0.1
    assert after["text"].count(replacement) == 1
    assert original not in after["text"]


def test_pdf_font_envelope_does_not_allow_a_new_overlap_with_neighboring_content():
    original = "Built stable interfaces with SQL."
    replacement = "Built easy interfaces with SQL."
    unobstructed = _pdf_line(original)
    safe_unit = extract_source_units(unobstructed, "pdf")[0]
    assert (
        replacement
        in _pdf_snapshot(apply_source_edits(unobstructed, "pdf", [_edit(safe_unit, replacement)]))[
            0
        ]["text"]
    )

    source = _pdf_line(original, obstruction=True)
    unit = extract_source_units(source, "pdf")[0]
    before = _pdf_snapshot(source)[0]
    assert before["objects"][0][1][1] > 617.75  # No original overlap with the rectangle.
    with pytest.raises(ResumeLayoutError, match="overlaps neighboring") as error:
        apply_source_edits(source, "pdf", [_edit(unit, replacement)])
    assert error.value.unit_id == unit["unit_id"]
    after = _pdf_snapshot(source)[0]
    assert after["text"] == before["text"]
    assert after["image"].tobytes() == before["image"].tobytes()


def _pdf_same_font_name_with_bullet_encoding() -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(612, 792)
    body_font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Times-Roman"),
            NameObject("/Encoding"): NameObject("/WinAnsiEncoding"),
        }
    )
    bullet_font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Times-Roman"),
            NameObject("/Encoding"): DictionaryObject(
                {
                    NameObject("/Type"): NameObject("/Encoding"),
                    NameObject("/BaseEncoding"): NameObject("/WinAnsiEncoding"),
                    NameObject("/Differences"): ArrayObject(
                        [NumberObject(1), NameObject("/bullet")]
                    ),
                }
            ),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {
                    NameObject("/FBody"): writer._add_object(body_font),
                    NameObject("/FBullet"): writer._add_object(bullet_font),
                }
            )
        }
    )
    stream = DecodedStreamObject()
    stream.set_data(
        b"BT /FBody 11 Tf 60 620 Td (Created stable interfaces with SQL.) Tj ET\n"
        b"BT /FBullet 11 Tf 48 620 Td <01> Tj ET\n"
    )
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def test_pdf_edit_retains_untouched_bullet_resource_with_same_base_font_name():
    source = _pdf_same_font_name_with_bullet_encoding()
    unit = extract_source_units(source, "pdf")[0]
    replacement = "Built stable interfaces with SQL."
    before = _pdf_snapshot(source)[0]
    result = apply_source_edits(source, "pdf", [_edit(unit, replacement)])
    after = _pdf_snapshot(result)[0]
    assert after["objects"][0][0] == replacement
    assert before["objects"][1][0].strip() == "•"
    assert after["objects"][1] == before["objects"][1]
    bullet_crop = (90, 325, 108, 346)  # Twice the PDF coordinates around the bullet.
    assert after["image"].crop(bullet_crop).tobytes() == before["image"].crop(bullet_crop).tobytes()
    assert before["image"].crop(bullet_crop).getextrema() != ((255, 255),) * 3
    with BytesIO(source) as original_stream, BytesIO(result) as result_stream:
        original_page = PdfReader(original_stream).pages[0]
        result_page = PdfReader(result_stream).pages[0]
        original_fonts = original_page["/Resources"]["/Font"]
        result_fonts = result_page["/Resources"]["/Font"]
        assert set(result_fonts) == set(original_fonts) == {"/FBody", "/FBullet"}
        assert original_fonts["/FBody"]["/BaseFont"] == original_fonts["/FBullet"]["/BaseFont"]
        for name in original_fonts:
            assert result_fonts[name].get_object() == original_fonts[name].get_object()
        assert result_fonts["/FBody"]["/Encoding"] != result_fonts["/FBullet"]["/Encoding"]


def test_pdf_candidates_include_complete_line_context_and_supported_font_characters():
    source = _pdf_same_font_name_with_bullet_encoding()
    units = extract_source_units(source, "pdf")
    assert len(units) == 1
    unit = units[0]
    assert unit["text"] == "Created stable interfaces with SQL."
    assert unit["line_context"] == "• Created stable interfaces with SQL."
    assert set(unit["text"]) <= set(unit["allowed_characters"])
    assert "y" in unit["allowed_characters"] and "y" not in unit["text"]
    assert "Ж" not in unit["allowed_characters"]
    assert extract_source_units(source, "pdf") == units

    subset_unit = _first_unit(_pdf(subset=True), "pdf")
    assert set(subset_unit["text"]) <= set(subset_unit["allowed_characters"])
    assert "Ж" not in subset_unit["allowed_characters"]


def test_pdf_context_keeps_section_boundaries_and_prioritizes_summary_over_skills():
    output = BytesIO()
    canvas = Canvas(output, pagesize=(612, 792), invariant=True)
    passages = [
        ("SUMMARY", 660, "Built Python services for customers with reliable releases."),
        ("EXPERIENCE", 610, "Developed backend services with automated quality checks."),
        (
            "TECHNICAL SKILLS",
            560,
            "Python services, deployment tools, profiling, and release automation.",
        ),
    ]
    for heading, y, text in passages:
        canvas.setFont("Times-Bold", 12)
        canvas.drawString(60, y, heading)
        canvas.setFont("Times-Roman", 11)
        canvas.drawString(60, y - 20, text)
    canvas.save()
    units = extract_source_units(output.getvalue(), "pdf")
    assert [unit["section"] for unit in units] == ["summary", "experience", "technical skills"]
    assert [unit["generation_priority"] for unit in units] == [1, 2, 8]
    for index, unit in enumerate(units):
        assert unit["text"] in unit["paragraph_context"]
        assert all(
            other[2] not in unit["paragraph_context"]
            for other_index, other in enumerate(passages)
            if other_index != index
        )


def test_pdf_fragment_boundaries_protect_sentence_join_before_emphasized_metric():
    output = BytesIO()
    canvas = Canvas(output, pagesize=(612, 792), invariant=True)
    canvas.setFont("Times-Bold", 12)
    canvas.drawString(48, 704, "EXPERIENCE")
    x = 60.0
    pieces = [
        ("Times-Bold", "Developed Python automation tools "),
        ("Times-Roman", "for ARM architecture validation, enabling "),
        ("Times-Bold", "100+ engineers"),
    ]
    for font, text in pieces:
        canvas.setFont(font, 10)
        canvas.drawString(x, 662, text)
        x += pdfmetrics.stringWidth(text, font, 10)
    canvas.save()
    source = output.getvalue()
    unit = next(
        unit for unit in extract_source_units(source, "pdf") if unit["text"].startswith("for ARM")
    )
    assert unit["required_prefix"] == "for ARM"
    assert unit["required_suffix"] == "validation, enabling"
    with pytest.raises(ResumeLayoutError, match="adjoining fragment boundaries") as rejected:
        apply_source_edits(
            source, "pdf", [_edit(unit, "for ARM validation, enabling reliability ")]
        )
    assert rejected.value.unit_id == unit["unit_id"]
    accepted = apply_source_edits(source, "pdf", [_edit(unit, "for ARM validation, enabling ")])
    with (
        pdfium.PdfDocument(accepted) as document,
        closing(document[0]) as page,
        closing(page.get_textpage()) as textpage,
    ):
        assert "for ARM validation, enabling" in textpage.get_text_range()
        assert "100+ engineers" in textpage.get_text_range()


def test_pdf_fragment_boundaries_leave_complete_sentences_bullets_and_columns_free():
    output = BytesIO()
    canvas = Canvas(output, pagesize=(612, 792), invariant=True)
    canvas.setFont("Times-Roman", 10)
    canvas.drawString(48, 662, "•")
    canvas.drawString(60, 662, ORIGINAL)
    canvas.drawString(400, 662, "Another isolated column with professional details.")
    canvas.save()
    units = extract_source_units(output.getvalue(), "pdf")
    assert len(units) == 2
    assert all(
        not unit.get("required_prefix") and not unit.get("required_suffix") for unit in units
    )


def _docx() -> bytes:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.65)
    section.left_margin = section.right_margin = Inches(0.65)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    document.add_paragraph("Taylor Example", "Title")
    document.add_paragraph("taylor@example.com | Bengaluru")
    document.add_paragraph("Experience", "Heading 1")
    paragraph = document.add_paragraph()
    paragraph.add_run("Built scalable systems with ")
    paragraph.add_run("Python").bold = True
    paragraph.add_run(" and ")
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), "skills")
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    italic = OxmlElement("w:i")
    properties.append(italic)
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = "SQL"
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    paragraph.add_run(".")
    document.add_paragraph("Reduced API latency by 40% through focused profiling.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Example Company"
    table.cell(0, 1).text = "2020 - 2024"
    section.header.paragraphs[0].text = "Original running header"
    section.footer.paragraphs[0].text = "Original running footer"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _without_text(xml: bytes) -> bytes:
    tree = etree.fromstring(xml)
    for node in tree.iter(f"{W}t"):
        node.text = ""
        node.attrib.pop("{http://www.w3.org/XML/1998/namespace}space", None)
    return etree.tostring(tree, method="c14n")


def test_docx_edit_retains_runs_bold_hyperlinks_tables_headers_footers_and_zip_parts():
    # Native patch invariants are separate from actual private render validation.
    source = _docx()
    result = _apply_docx(source, [_edit(_first_unit(source, "docx"))])
    with ZipFile(BytesIO(source)) as before, ZipFile(BytesIO(result)) as after:
        assert before.namelist() == after.namelist()
        for name in before.namelist():
            if name != "word/document.xml":
                assert before.read(name) == after.read(name)
        assert _without_text(before.read("word/document.xml")) == _without_text(
            after.read("word/document.xml")
        )
        tree = etree.fromstring(after.read("word/document.xml"))
        paragraph = list(tree.find(f"{W}body").iter(f"{W}p"))[3]
        assert "".join(node.text or "" for node in paragraph.iter(f"{W}t")) == REPLACEMENT
        bold_run = next(
            run for run in paragraph.iter(f"{W}r") if run.find(f"{W}rPr/{W}b") is not None
        )
        assert bold_run.find(f"{W}t").text == "Python"
        hyperlink = paragraph.find(f"{W}hyperlink")
        assert hyperlink.get(f"{W}anchor") == "skills"
        assert hyperlink.find(f"{W}r/{W}t").text == "SQL"


def test_docx_native_patch_cannot_certify_render_fit_or_edit_nonbody_parts(monkeypatch):
    source = _docx()
    units = extract_source_units(source, "docx")
    assert all("header" not in unit["text"] and "footer" not in unit["text"] for unit in units)
    edit = _edit(_first_unit(source, "docx"), ORIGINAL + " with API monitoring")
    assert _apply_docx(source, [edit]) != source
    # A syntactically valid native patch is not a fit receipt. With no configured
    # renderer, the public path must refuse before any identity/provider call.
    monkeypatch.setenv("DOCUMENT_WORKER_URL", "")
    with pytest.raises(ResumeLayoutError, match="private DOCX renderer"):
        apply_source_edits(source, "docx", [edit])


def test_docx_native_font_allows_new_ascii_letters_without_changing_styles():
    source = _docx()
    replacement = "Built robust systems with Python and SQL."
    assert "r" in replacement and "r" not in ORIGINAL
    result = _apply_docx(source, [_edit(_first_unit(source, "docx"), replacement)])
    with ZipFile(BytesIO(result)) as archive:
        tree = etree.fromstring(archive.read("word/document.xml"))
        paragraph = list(tree.find(f"{W}body").iter(f"{W}p"))[3]
        assert "".join(node.text or "" for node in paragraph.iter(f"{W}t")) == replacement


@pytest.mark.parametrize("kind,fixture", [("pdf", _pdf), ("docx", _docx)])
def test_noop_is_byte_identical_and_unknown_stale_duplicate_or_fact_edits_fail(kind, fixture):
    source = fixture()
    unit = _first_unit(source, kind)
    assert apply_source_edits(source, kind, []) is source
    assert apply_source_edits(source, kind, [_edit(unit, ORIGINAL)]) is source
    for mutation in (
        {"unit_id": "unknown"},
        {"original_text": "stale text"},
        {"replacement_text": "Built systems with Python and SQL for 20 users."},
        {"replacement_text": "Built systems\nwith Python and SQL."},
    ):
        edit = _edit(unit)
        edit.update(mutation)
        with pytest.raises(ResumeLayoutError):
            apply_source_edits(source, kind, [edit])
    with pytest.raises(ResumeLayoutError, match="repeated"):
        apply_source_edits(source, kind, [_edit(unit), _edit(unit)])
    numeric = next(unit for unit in extract_source_units(source, kind) if "40%" in unit["text"])
    with pytest.raises(ResumeLayoutError, match="numbers"):
        apply_source_edits(source, kind, [_edit(numeric, numeric["text"].replace("40%", "80%"))])


@pytest.mark.parametrize(
    "kind,source", [("pdf", b"not a PDF"), ("docx", b"not a DOCX"), ("txt", b"plain text")]
)
def test_unsupported_sources_fail_clearly(kind, source):
    with pytest.raises(ResumeLayoutError):
        extract_source_units(source, kind)


def test_docx_manual_breaks_and_fields_are_not_editable():
    source = _docx()
    with ZipFile(BytesIO(source)) as archive:
        tree = etree.fromstring(archive.read("word/document.xml"))
        paragraph = list(tree.find(f"{W}body").iter(f"{W}p"))[3]
        paragraph.find(f"{W}r").append(etree.Element(f"{W}br"))
        result = BytesIO()
        with ZipFile(result, "w") as output:
            for info in archive.infolist():
                output.writestr(
                    info,
                    etree.tostring(tree)
                    if info.filename == "word/document.xml"
                    else archive.read(info.filename),
                )
    assert ORIGINAL not in [
        unit["text"] for unit in extract_source_units(result.getvalue(), "docx")
    ]


def test_employer_role_and_emphasized_identity_lines_remain_protected():
    output = BytesIO()
    canvas = Canvas(output, pagesize=(612, 792), invariant=True)
    canvas.setFont("Times-Roman", 11)
    canvas.drawString(48, 650, "ARM Embedded Technologies Pvt. Ltd. Bengaluru")
    canvas.drawString(48, 630, "Senior Software Engineer Platform Systems 2020 - 2024")
    canvas.setFont("Times-Bold", 11)
    canvas.drawString(48, 610, "Example Technical Research and Development Company")
    canvas.setFont("Times-Italic", 11)
    canvas.drawString(48, 590, "Principal Platform Technology Research Associate")
    canvas.setFont("Times-Roman", 11)
    canvas.drawString(48, 570, ORIGINAL)
    canvas.save()
    assert [unit["text"] for unit in extract_source_units(output.getvalue(), "pdf")] == [ORIGINAL]

    document = Document()
    document.add_paragraph("ARM Embedded Technologies Pvt. Ltd. Bengaluru")
    document.add_paragraph("Senior Software Engineer Platform Systems 2020 - 2024")
    document.add_paragraph().add_run(
        "Example Technical Research and Development Company"
    ).bold = True
    paragraph = document.add_paragraph()
    paragraph.add_run("Built scalable systems with ")
    paragraph.add_run("Python").bold = True
    paragraph.add_run(" and SQL.")
    data = BytesIO()
    document.save(data)
    assert [unit["text"] for unit in extract_source_units(data.getvalue(), "docx")] == [ORIGINAL]


# Link-association regressions use actual synthetic native documents, not provider calls.
def _docx_hyperlink_source(*, external=True, split_label=False, second_paragraph=False):
    from docx.opc.constants import RELATIONSHIP_TYPE

    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Built scalable systems with Python and ")
    hyperlink = OxmlElement("w:hyperlink")
    if external:
        identity = paragraph.part.relate_to(
            "https://example.invalid/sql-guide", RELATIONSHIP_TYPE.HYPERLINK, is_external=True
        )
        hyperlink.set(qn("r:id"), identity)
    else:
        hyperlink.set(qn("w:anchor"), "skills")
    for label in (["SQ", "L"] if split_label else ["SQL"]):
        run = OxmlElement("w:r")
        properties = OxmlElement("w:rPr")
        properties.append(OxmlElement("w:i"))
        run.append(properties)
        text = OxmlElement("w:t")
        text.text = label
        run.append(text)
        hyperlink.append(run)
    paragraph._p.append(hyperlink)
    paragraph.add_run(".")
    if second_paragraph:
        document.add_paragraph("Developed automated quality checks for reliable releases.")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


@pytest.mark.parametrize("external", [True, False])
@pytest.mark.parametrize("split_label", [True, False])
@pytest.mark.parametrize("replacement", [
    "Built systems with Python and Go.",
    "Built systems with Python and robust SQL.",
])
def test_docx_link_labels_and_styled_boundaries_cannot_receive_replacement_text(
    external, split_label, replacement
):
    source = _docx_hyperlink_source(external=external, split_label=split_label)
    unit = _first_unit(source, "docx")
    with pytest.raises(ResumeLayoutError, match="hyperlink label") as rejected:
        apply_source_edits(source, "docx", [_edit(unit, replacement)])
    assert rejected.value.unit_id == unit["unit_id"]
    assert "Keep existing hyperlink labels exactly" in rejected.value.repair_hint
    assert ORIGINAL not in rejected.value.repair_hint
    assert _first_unit(source, "docx")["text"] == ORIGINAL
    assert apply_source_edits(source, "docx", []) is source


@pytest.mark.parametrize("external", [True, False])
@pytest.mark.parametrize("split_label", [True, False])
def test_docx_surrounding_body_edits_preserve_exact_link_nodes_and_relationships(
    external, split_label
):
    source = _docx_hyperlink_source(external=external, split_label=split_label)
    result = _apply_docx(source, [_edit(_first_unit(source, "docx"))])
    with ZipFile(BytesIO(source)) as before, ZipFile(BytesIO(result)) as after:
        assert before.namelist() == after.namelist()
        for name in before.namelist():
            if name != "word/document.xml":
                assert before.read(name) == after.read(name)
        original = etree.fromstring(before.read("word/document.xml"))
        saved = etree.fromstring(after.read("word/document.xml"))
        assert etree.tostring(next(original.iter(f"{W}hyperlink")), method="c14n") == etree.tostring(
            next(saved.iter(f"{W}hyperlink")), method="c14n"
        )
        assert _without_text(before.read("word/document.xml")) == _without_text(after.read("word/document.xml"))
        assert "".join(node.text or "" for node in saved.iter(f"{W}t")) == REPLACEMENT


def test_docx_link_refusal_is_atomic_and_attributed_for_retaining_another_safe_edit():
    source = _docx_hyperlink_source(second_paragraph=True)
    units = extract_source_units(source, "docx")
    linked = next(unit for unit in units if unit["text"] == ORIGINAL)
    other = next(unit for unit in units if unit["text"].startswith("Developed"))
    edits = [_edit(other, "Developed quality checks for reliable releases."),
             _edit(linked, "Built systems with Python and robust SQL.")]
    with pytest.raises(ResumeLayoutError, match="hyperlink label") as refused:
        apply_source_edits(source, "docx", edits)
    assert refused.value.unit_id == linked["unit_id"]
    assert extract_source_units(source, "docx") == units
    retained = [edit for edit in edits if edit["unit_id"] != refused.value.unit_id]
    result = _apply_docx(source, retained)
    assert [unit["text"] for unit in extract_source_units(result, "docx")] == [
        ORIGINAL, "Developed quality checks for reliable releases."
    ]


def _pdf_hyperlink_source(text, *, extra=False):
    output = BytesIO()
    canvas = Canvas(output, pagesize=(612, 792), invariant=True)
    canvas.setFont("Times-Roman", 11)
    canvas.drawString(60, 620, text)
    left = 60 + pdfmetrics.stringWidth(text[:text.index("SQL")], "Times-Roman", 11)
    right = left + pdfmetrics.stringWidth("SQL", "Times-Roman", 11)
    canvas.linkURL("https://example.invalid/sql-guide", (left, 616, right, 633), thickness=0)
    if extra:
        canvas.drawString(60, 580, "Developed automated quality checks for reliable releases.")
    canvas.save()
    return output.getvalue()


def test_pdf_edit_after_link_preserves_clickable_text_and_original_destination():
    original = "Built SQL systems with reliable Python tools."
    replacement = "Built SQL systems with Python tools."
    source = _pdf_hyperlink_source(original)
    unit = extract_source_units(source, "pdf")[0]
    result = apply_source_edits(source, "pdf", [_edit(unit, replacement)])
    import backend.app.services.resume_layout as layout

    with pdfium.PdfDocument(source) as before, pdfium.PdfDocument(result) as after:
        with closing(before[0]) as old_page, closing(after[0]) as new_page:
            assert layout._pdf_link_text(old_page) == layout._pdf_link_text(new_page)
            assert layout._pdf_link_text(old_page)[0][2] == "SQL "
            assert layout._pdf_link_text(new_page)[0][2] == "SQL "
            assert layout._page_structure(old_page, before) == layout._page_structure(new_page, after)
    assert replacement in _pdf_snapshot(result)[0]["text"]


def test_pdf_changed_link_label_is_refused_without_a_template_or_overlay():
    source = _pdf_hyperlink_source(ORIGINAL)
    unit = _first_unit(source, "pdf")
    with pytest.raises(ResumeLayoutError, match="original PDF link") as refused:
        apply_source_edits(source, "pdf", [_edit(unit, "Built systems with Python and Go.")])
    assert refused.value.unit_id == unit["unit_id"]
    assert "original clickable rectangle" in refused.value.repair_hint
    assert ORIGINAL not in refused.value.repair_hint
    assert _pdf_snapshot(source)[0]["text"] == ORIGINAL


def test_pdf_link_movement_refusal_preserves_an_unrelated_useful_edit():
    source = _pdf_hyperlink_source(ORIGINAL, extra=True)
    units = extract_source_units(source, "pdf")
    linked = next(unit for unit in units if unit["text"] == ORIGINAL)
    other = next(unit for unit in units if unit["text"].startswith("Developed"))
    edits = [_edit(linked), _edit(other, "Developed quality checks for reliable releases.")]
    with pytest.raises(ResumeLayoutError, match="original PDF link") as refused:
        apply_source_edits(source, "pdf", edits)
    assert refused.value.unit_id == linked["unit_id"]
    assert extract_source_units(source, "pdf") == units
    result = apply_source_edits(source, "pdf", [edit for edit in edits if edit["unit_id"] != refused.value.unit_id])
    saved_text = _pdf_snapshot(result)[0]["text"]
    assert ORIGINAL in saved_text and "Developed quality checks for reliable releases." in saved_text
