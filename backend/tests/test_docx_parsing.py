import io

from backend.app.services.parsing import extract_text_from_docx
from docx import Document


def test_docx_extraction_keeps_tables_in_document_order_and_deduplicates_merged_cells():
    document = Document()
    document.add_paragraph("Before the table")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Acme Engineer"
    table.cell(0, 1).text = "2023-2025"
    merged = table.cell(1, 0).merge(table.cell(1, 1))
    merged.text = "Reduced API latency by 40 percent."
    document.add_paragraph("After the table")
    output = io.BytesIO()
    document.save(output)

    extracted = extract_text_from_docx(output.getvalue())

    assert extracted.splitlines() == [
        "Before the table",
        "Acme Engineer\t2023-2025",
        "Reduced API latency by 40 percent.",
        "After the table",
    ]


def test_docx_extraction_preserves_content_from_a_table_only_resume():
    document = Document()
    table = document.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "Jordan Candidate"
    table.cell(1, 0).text = "Experience\nBuilt reliable Python services."
    output = io.BytesIO()
    document.save(output)

    assert extract_text_from_docx(output.getvalue()).splitlines() == [
        "Jordan Candidate", "Experience", "Built reliable Python services."
    ]
