"""Synthetic monetary policy invariants; these rates are not real model prices."""
from __future__ import annotations

import copy
import json
from datetime import UTC, datetime
from types import SimpleNamespace

import httpx
import pytest
from backend.app import models
from backend.app.domains.analysis import tasks
from backend.app.services import llm_client
from backend.app.services import model_cost_policy as cost_policy
from backend.app.services.generation_budget import (
    GenerationBudgetExhausted,
    persistent_run_budget,
)
from backend.app.services.guardrails import billable_operation
from backend.app.services.model_cost_policy import (
    ModelCostPolicy,
    ModelCostUnavailable,
    freeze_run_quote,
    input_token_estimate,
    quoted_cost,
)
from backend.tests.model_cost_fixtures import synthetic_policy
from backend.tests.test_ai_efficiency import _create, _database
from fastapi import HTTPException
from google import genai

MESSAGES = [{"role": "user", "content": "Synthetic Python role"}]
GOOGLE = "https://generativelanguage.googleapis.com"


def new_run(factory, *, ceiling=10_000_000, attempts=3, frozen=True, operation="job_match", run_id="run_cost"):
    with factory() as db:
        run = models.AnalysisRun(
            id=run_id, user_id=1, operation=operation, status="running",
            idempotency_key="cost-test-key-" + run_id, input_fingerprint="c" * 64,
            input_payload={}, generation_attempt_limit=attempts,
        )
        if frozen:
            freeze_run_quote(run)
            if ceiling != 10_000_000:
                quote = copy.deepcopy(run.model_cost_quote)
                quote["operation_limits_micros"][operation] = ceiling
                run.model_cost_quote = quote
                run.model_cost_ceiling_micros = ceiling
        db.add(run)
        db.commit()
    return run_id


@pytest.fixture
def context():
    engine, factory = _database()
    try:
        yield factory
    finally:
        engine.dispose()


def admit(budget, messages=None):
    return budget.admit("google", "gemini-3.6-flash", messages or MESSAGES, api_base=GOOGLE)


def usage(prompt=10, candidates=4, thoughts=3):
    return SimpleNamespace(text="synthetic", usage_metadata=SimpleNamespace(
        prompt_token_count=prompt, candidates_token_count=candidates,
        thoughts_token_count=thoughts, total_token_count=prompt + candidates + thoughts,
    ))


def test_no_policy_deterministic_and_mixed_legacy_work_without_generation(context, monkeypatch):
    monkeypatch.delenv("LLM_MODEL_COST_POLICY_JSON")
    run_id, _ = _create(context)
    monkeypatch.setattr(tasks, "SessionLocal", context)
    assert tasks.process_analysis_run(run_id) == "succeeded"
    with context() as db:
        with billable_operation(user_id=1, db=db, operation="match_question", amount=0, input_payload={"mode": "direct"}) as run:
            legacy_id = run.id
    with context() as db:
        assert db.get(models.AnalysisRun, legacy_id).model_cost_quote is None
        assert db.query(models.ModelCallEvent).count() == 0


def test_missing_policy_rejects_optional_generation_before_units_or_run(context, monkeypatch):
    monkeypatch.delenv("LLM_MODEL_COST_POLICY_JSON")
    with pytest.raises(HTTPException) as error:
        _create(context, mode="enhanced")
    assert error.value.status_code == 503
    with context() as db:
        assert db.get(models.User, 1).ai_credits == 100
        assert db.query(models.AnalysisRun).count() == 0
        with pytest.raises(HTTPException) as error:
            with billable_operation(user_id=1, db=db, operation="rewrite_bullets", amount=1):
                pytest.fail("No execution without quote")
        assert error.value.status_code == 503


def test_unscoped_chat_cannot_make_external_attempt(monkeypatch):
    monkeypatch.setenv("LLM_API_KEY", "synthetic")
    monkeypatch.setattr(llm_client, "LLM_MODEL", "gemini-3.6-flash")
    monkeypatch.setattr(genai, "Client", lambda **kwargs: pytest.fail("Must admit before SDK creation"))
    with pytest.raises(llm_client.LLMProviderError, match="persisted operation") as error:
        llm_client._chat(MESSAGES)
    assert error.value.budget_exhausted is True


@pytest.mark.parametrize("mutation", [
    lambda p: p.update(currency="INR"),
    lambda p: p["operation_limits_micros"].update(job_match=0),
    lambda p: p["operation_limits_micros"].update(job_match=True),
    lambda p: p["pricing_quotes"][0].update(input_rate_micros_per_million=0),
    lambda p: p["pricing_quotes"][0].pop("output_rate_micros_per_million"),
    lambda p: p["pricing_quotes"][0].update(api_base="https://user:secret@example.com"),
    lambda p: p["pricing_quotes"][0].update(model="*"),
    lambda p: p["pricing_quotes"].append(p["pricing_quotes"][0]),
])
def test_policy_has_no_free_unknown_ambiguous_or_wildcard_prices(mutation):
    policy = synthetic_policy()
    mutation(policy)
    with pytest.raises(ValueError):
        ModelCostPolicy.model_validate(policy)


def test_reservation_survives_crash_and_new_context_without_reset(context):
    price = synthetic_policy()["pricing_quotes"][0]
    hold = quoted_cost(price, input_token_estimate(MESSAGES), price["max_output_tokens"])
    run_id = new_run(context, ceiling=hold, attempts=6)
    with persistent_run_budget(context, run_id) as budget:
        record = admit(budget)  # simulate process death before finish/network outcome
    with context() as db:
        run = db.get(models.AnalysisRun, run_id)
        assert run.model_cost_reserved_micros == hold
        assert run.model_cost_settled_micros == 0
        assert db.get(models.ModelCallEvent, record["event_id"]).cost_state == "reserved"
    with persistent_run_budget(context, run_id) as budget:
        with pytest.raises(GenerationBudgetExhausted, match="estimated spend"):
            admit(budget)
    with context() as db:
        assert db.get(models.AnalysisRun, run_id).generation_attempt_count == 1


def test_unknown_and_missing_usage_keep_hold_then_settle_once(context):
    run_id = new_run(context)
    with persistent_run_budget(context, run_id) as budget:
        record = admit(budget)
        budget.finish(record, latency_ms=5, error=RuntimeError("secret-provider-body private@example.com"))
        with context() as db:
            held = db.get(models.AnalysisRun, run_id).model_cost_reserved_micros
            event = db.get(models.ModelCallEvent, record["event_id"])
            assert event.error_code == "provider_error"
            assert event.cost_state == "outcome_unknown"
        budget.finish(record, latency_ms=7, response=usage())
        budget.finish(record, latency_ms=8, response=usage())
        budget.finish(record, latency_ms=9, error=RuntimeError("late duplicate error"))
    with context() as db:
        run = db.get(models.AnalysisRun, run_id)
        event = db.get(models.ModelCallEvent, record["event_id"])
        assert held > 17
        assert run.model_cost_reserved_micros == 0
        assert run.model_cost_settled_micros == event.settled_cost_micros == 17
        assert event.cost_state == "settled" and event.error_code is None
        assert event.input_tokens == 10 and event.output_tokens == 7
        assert event.tokens_estimated is False
        assert event.settled_at is not None
        assert "private@example.com" not in json.dumps(event.pricing_quote)


@pytest.mark.parametrize("response", [
    SimpleNamespace(text="synthetic"),
    SimpleNamespace(usage_metadata=SimpleNamespace(prompt_token_count=10, candidates_token_count=4)),
    SimpleNamespace(usage_metadata=SimpleNamespace(prompt_token_count=10, candidates_token_count=4, total_token_count=16, thoughts_token_count=3)),
    {"usage": {"prompt_tokens": 10, "completion_tokens": -1}},
])
def test_incomplete_or_invalid_usage_is_not_zero_cost(context, response):
    run_id = new_run(context)
    with persistent_run_budget(context, run_id) as budget:
        record = admit(budget)
        budget.finish(record, latency_ms=1, response=response)
    with context() as db:
        run = db.get(models.AnalysisRun, run_id)
        event = db.get(models.ModelCallEvent, record["event_id"])
        assert run.model_cost_reserved_micros == event.reserved_cost_micros > 0
        assert event.settled_cost_micros is None and event.settled_at is None
        assert event.cost_state == "usage_unavailable" and event.tokens_estimated is True


def test_usage_overrun_is_not_clipped_and_blocks_repairs(context):
    run_id = new_run(context)
    with persistent_run_budget(context, run_id) as budget:
        record = admit(budget)
        budget.finish(record, latency_ms=1, response=usage(candidates=20_000))
        with pytest.raises(GenerationBudgetExhausted, match="exceeded"):
            admit(budget)
    with context() as db:
        run = db.get(models.AnalysisRun, run_id)
        event = db.get(models.ModelCallEvent, record["event_id"])
        assert run.model_cost_state == event.cost_state == "overrun"
        assert run.model_cost_settled_micros == 20_013 > event.reserved_cost_micros
        assert run.generation_attempt_count == 1


def test_current_tighter_ceiling_denies_without_rewriting_quote(context, monkeypatch):
    run_id = new_run(context)
    with context() as db:
        original = copy.deepcopy(db.get(models.AnalysisRun, run_id).model_cost_quote)
    monkeypatch.setenv("LLM_MODEL_COST_POLICY_JSON", json.dumps(synthetic_policy(ceiling=1)))
    with persistent_run_budget(context, run_id) as budget:
        with pytest.raises(GenerationBudgetExhausted, match="estimated spend"):
            admit(budget)
    with context() as db:
        run = db.get(models.AnalysisRun, run_id)
        assert run.model_cost_quote == original
        assert run.model_cost_ceiling_micros == 10_000_000
        assert run.generation_attempt_count == 0


@pytest.mark.parametrize("change", ["remove", "rate", "version", "endpoint"])
def test_revoked_or_changed_exact_price_denies_future_attempt(context, monkeypatch, change):
    run_id = new_run(context)
    policy = synthetic_policy()
    if change == "remove":
        policy["pricing_quotes"].pop(0)
    else:
        key, value = {"rate": ("input_rate_micros_per_million", 9), "version": ("version", "v2"), "endpoint": ("api_base", "https://other.example")}[change]
        policy["pricing_quotes"][0][key] = value
    monkeypatch.setenv("LLM_MODEL_COST_POLICY_JSON", json.dumps(policy))
    with persistent_run_budget(context, run_id) as budget:
        with pytest.raises(ModelCostUnavailable):
            admit(budget)
    with context() as db:
        assert db.get(models.AnalysisRun, run_id).generation_attempt_count == 0


def test_current_output_tightening_is_frozen_on_attempt(context, monkeypatch):
    run_id = new_run(context)
    monkeypatch.setenv("LLM_MODEL_COST_POLICY_JSON", json.dumps(synthetic_policy(output_limit=12)))
    with persistent_run_budget(context, run_id) as budget:
        record = admit(budget)
    assert record["max_output_tokens"] == 12
    with context() as db:
        event = db.get(models.ModelCallEvent, record["event_id"])
        run = db.get(models.AnalysisRun, run_id)
        assert event.output_token_limit == 12
        assert run.model_cost_quote["pricing_quotes"][0]["max_output_tokens"] == 1024


def test_exact_prompt_bytes_and_input_limit_reject_before_admission(context, monkeypatch):
    messages = [{"role": "user", "content": "🙂" * 100}]
    assert input_token_estimate(messages) > 400
    policy = synthetic_policy()
    policy["pricing_quotes"][0]["max_input_tokens"] = 20
    monkeypatch.setenv("LLM_MODEL_COST_POLICY_JSON", json.dumps(policy))
    run_id = new_run(context)
    with persistent_run_budget(context, run_id) as budget:
        with pytest.raises(GenerationBudgetExhausted, match="input estimate"):
            admit(budget, messages)
    with context() as db:
        assert db.query(models.ModelCallEvent).count() == 0


def test_legacy_unquoted_attempt_history_cannot_gain_fresh_budget(context):
    run_id = new_run(context, frozen=False)
    with context() as db:
        run = db.get(models.AnalysisRun, run_id)
        run.generation_attempt_count = 1
        db.commit()
    with persistent_run_budget(context, run_id) as budget:
        with pytest.raises(ModelCostUnavailable, match="legacy"):
            admit(budget)
    with context() as db:
        assert db.get(models.AnalysisRun, run_id).model_cost_quote is None


def test_legacy_first_generation_freezes_price_and_hold_before_network(context, monkeypatch):
    monkeypatch.setenv("LLM_API_KEY", "synthetic")
    monkeypatch.setattr(llm_client, "LLM_MODEL", "gemini-3.6-flash")
    seen = []
    class Models:
        def generate_content(self, *, config, **kwargs):
            with context() as db:
                run = db.query(models.AnalysisRun).one()
                seen.append(run.id)
                assert run.model_cost_quote and run.model_cost_reserved_micros > 0
                assert db.query(models.ModelCallEvent).one().output_token_limit == config.max_output_tokens == 1024
                assert config.candidate_count == 1
            return usage()
    def client(**kwargs):
        assert kwargs["http_options"].retry_options.attempts == 1
        return SimpleNamespace(models=Models())
    monkeypatch.setattr(genai, "Client", client)
    with context() as db:
        with billable_operation(user_id=1, db=db, operation="job_match_legacy", amount=1):
            assert llm_client._chat(MESSAGES) == "synthetic"
    assert len(seen) == 1


@pytest.mark.parametrize("provider,model,base,parameter", [
    ("openai", "gpt-4o-mini", "https://api.openai.com/v1", "max_completion_tokens"),
    ("openai_compatible", "synthetic-compatible", "https://provider.example/v1", "max_tokens"),
])
def test_http_provider_output_bound_and_complete_usage(context, monkeypatch, provider, model, base, parameter):
    run_id = new_run(context)
    monkeypatch.setenv("LLM_API_KEY", "synthetic")
    monkeypatch.setattr(llm_client, "LLM_MODEL", model)
    monkeypatch.setattr(llm_client, "LLM_API_BASE", base)
    def post(url, *, json, **kwargs):
        assert url == base + "/chat/completions"
        assert json[parameter] == 1024
        with context() as db:
            assert db.get(models.AnalysisRun, run_id).model_cost_reserved_micros > 0
            assert db.query(models.ModelCallEvent).one().provider == provider
        return httpx.Response(200, request=httpx.Request("POST", url), json={
            "choices": [{"message": {"content": "synthetic"}}],
            "usage": {"prompt_tokens": 10, "completion_tokens": 8,
                      "completion_tokens_details": {"reasoning_tokens": 3}},
        })
    monkeypatch.setattr(httpx, "post", post)
    with persistent_run_budget(context, run_id):
        assert llm_client._chat(MESSAGES) == "synthetic"
    with context() as db:
        assert db.get(models.AnalysisRun, run_id).model_cost_settled_micros == 18


def test_nested_different_operation_cannot_charge_wrong_owner_budget(context):
    run_id = new_run(context)
    with persistent_run_budget(context, run_id):
        with pytest.raises(GenerationBudgetExhausted, match="distinct operations"):
            with persistent_run_budget(context, "different_run"):
                pytest.fail("No cross-operation budget reuse")


def test_expired_immutable_quote_denies_future_calls_without_reset(context, monkeypatch):
    run_id = new_run(context)
    with persistent_run_budget(context, run_id) as budget:
        record = admit(budget)
        budget.finish(record, latency_ms=1, response=usage())
    class FutureClock:
        @staticmethod
        def now(tz):
            return datetime(2100, 1, 1, tzinfo=UTC)
    monkeypatch.setattr(cost_policy, "datetime", FutureClock)
    with persistent_run_budget(context, run_id) as budget:
        with pytest.raises(ModelCostUnavailable, match="expired"):
            admit(budget)
    with context() as db:
        run = db.get(models.AnalysisRun, run_id)
        assert run.generation_attempt_count == 1
        assert run.model_cost_settled_micros == 17
        assert run.model_cost_quote["pricing_quotes"][0]["expires_at"].startswith("2099-")


@pytest.mark.parametrize("valid_from,expires_at", [
    ("2020-01-01T00:00:00", "2099-01-01T00:00:00Z"),
    ("2020-01-01T00:00:00Z", "2020-01-01T00:00:00Z"),
    ("2099-01-01T00:00:00Z", "2020-01-01T00:00:00Z"),
])
def test_price_validity_must_be_finite_ordered_and_timezone_aware(valid_from, expires_at):
    policy = synthetic_policy()
    policy["pricing_quotes"][0].update(valid_from=valid_from, expires_at=expires_at)
    with pytest.raises(ValueError):
        ModelCostPolicy.model_validate(policy)


@pytest.mark.parametrize("outcome,expected_calls", [("fallback", 2), ("permanent_error", 1), ("success", 1)])
def test_google_closes_every_attempt_and_pins_endpoint(context, monkeypatch, outcome, expected_calls):
    run_id = new_run(context)
    monkeypatch.setenv("LLM_API_KEY", "synthetic")
    monkeypatch.setenv("GOOGLE_GENAI_USE_VERTEXAI", "true")
    monkeypatch.setattr(llm_client, "LLM_MODEL", "gemini-3.6-flash")
    created, closed = [], []
    def client(**kwargs):
        index = len(created) + 1
        created.append(index)
        assert kwargs["vertexai"] is False
        options = kwargs["http_options"]
        assert options.base_url == GOOGLE
        assert options.timeout == 90_000
        assert options.retry_options.attempts == 1
        assert options.client_args["follow_redirects"] is False
        class Models:
            def generate_content(self, **kwargs):
                if outcome == "permanent_error":
                    raise RuntimeError("400 private-provider-body")
                if outcome == "fallback" and index == 1:
                    raise RuntimeError("503 private-provider-body")
                return usage()
        return SimpleNamespace(models=Models(), close=lambda: closed.append(index))
    monkeypatch.setattr(genai, "Client", client)
    with persistent_run_budget(context, run_id):
        if outcome == "permanent_error":
            with pytest.raises(llm_client.LLMProviderError) as error:
                llm_client._chat(MESSAGES)
            assert "private-provider-body" not in str(error.value)
            assert error.value.__suppress_context__ is True
        else:
            assert llm_client._chat(MESSAGES) == "synthetic"
    assert created == closed == list(range(1, expected_calls + 1))


def _upload_request():
    import io

    from docx import Document
    from starlette.datastructures import Headers, UploadFile
    from starlette.requests import Request

    document = Document()
    document.add_paragraph("Synthetic Candidate")
    document.add_heading("Skills", 2)
    document.add_paragraph("Python")
    output = io.BytesIO()
    document.save(output)
    original = output.getvalue()
    request = Request({"type": "http", "method": "POST", "path": "/resume/parse", "headers": [], "client": ("127.0.0.1", 1)})
    upload = UploadFile(io.BytesIO(original), filename="synthetic.docx", headers=Headers({"content-type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"}))
    return request, upload, original


def test_optional_upload_without_policy_returns_one_saved_original_and_no_charge(context, monkeypatch):
    import asyncio

    from backend.app.routers.resume import parse_resume

    monkeypatch.delenv("LLM_MODEL_COST_POLICY_JSON")
    monkeypatch.setattr(genai, "Client", lambda **kwargs: pytest.fail("No generation without policy"))
    request, upload, original = _upload_request()
    with context() as db:
        initial = db.query(models.Resume).count()
        response = asyncio.run(parse_resume.__wrapped__(request, upload, True, db, db.get(models.User, 1)))
        assert response.enrichment_state == "unavailable"
        assert response.extraction_mode == "deterministic"
        assert response.enrichment_units == 0
        assert any("model-cost authorization" in warning for warning in response.warnings)
        assert db.get(models.Resume, response.resume_id).source_document == original
        assert db.query(models.Resume).count() == initial + 1
        assert db.query(models.ModelCallEvent).count() == db.query(models.AnalysisRun).count() == 0
        assert db.get(models.User, 1).ai_credits == 100


def test_unrelated_upload_503_is_not_swallowed_as_policy_unavailability(context, monkeypatch):
    import asyncio
    from contextlib import contextmanager

    from backend.app.routers.resume import parse_resume
    from backend.app.services import guardrails

    @contextmanager
    def unrelated_failure(**kwargs):
        raise HTTPException(status_code=503, detail="Unrelated outage")
        yield  # pragma: no cover
    monkeypatch.setattr(guardrails, "billable_operation", unrelated_failure)
    request, upload, _ = _upload_request()
    with context() as db:
        with pytest.raises(HTTPException, match="Unrelated outage") as error:
            asyncio.run(parse_resume.__wrapped__(request, upload, True, db, db.get(models.User, 1)))
        assert error.value.status_code == 503


def test_foreign_operation_settlement_cannot_move_or_release_another_hold(context):
    first = new_run(context)
    second = new_run(context, run_id="run_other")
    with persistent_run_budget(context, first) as budget:
        record = admit(budget)
    with persistent_run_budget(context, second) as budget:
        with pytest.raises(ModelCostUnavailable, match="does not belong"):
            budget.finish(record, latency_ms=1, response=usage())
    with context() as db:
        assert db.get(models.AnalysisRun, first).model_cost_reserved_micros > 0
        assert db.get(models.AnalysisRun, second).model_cost_reserved_micros == 0
        assert db.get(models.ModelCallEvent, record["event_id"]).settled_at is None
