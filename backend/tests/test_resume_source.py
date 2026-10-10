from __future__ import annotations

import io

import pytest
from backend.app.database import Base, get_db
from backend.app.domains.analysis import schemas as analysis_schemas
from backend.app.domains.analysis import service as analysis_service
from backend.app.models import AnalysisRun, EvidenceItem, Opportunity, Resume, UsageEvent, User
from backend.app.rate_limiter import limiter
from backend.app.routers import auth
from backend.app.routers import resume as resume_router
from backend.app.security import create_access_token, get_current_user
from docx import Document
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from reportlab.pdfgen.canvas import Canvas
from slowapi import _rate_limit_exceeded_handler
from slowapi.errors import RateLimitExceeded
from sqlalchemy import create_engine, inspect
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool


def _source_bytes(source_format: str) -> bytes:
    output = io.BytesIO()
    if source_format == "pdf":
        document = Canvas(output)
        document.setFont("Times-Bold", 12)
        document.drawString(50, 750, "Original Resume")
        document.save()
    else:
        document = Document()
        document.add_paragraph().add_run("Original Resume").bold = True
        document.save(output)
    return output.getvalue()


def _client():
    limiter._storage.reset()
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    factory = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    with factory() as db:
        db.add_all(
            [
                User(id=1, email="owner@example.com", ai_credits=50, tier="free"),
                User(id=2, email="other@example.com", ai_credits=50, tier="free"),
                Resume(id=10, user_id=1, original_filename="legacy.pdf", raw_text="Original text"),
                Resume(
                    id=20,
                    user_id=2,
                    original_filename="private.pdf",
                    source_document=_source_bytes("pdf"),
                    source_format="pdf",
                ),
            ]
        )
        db.commit()

    app = FastAPI()
    app.state.limiter = limiter
    app.add_exception_handler(RateLimitExceeded, _rate_limit_exceeded_handler)
    app.include_router(resume_router.router, prefix="/api")
    app.include_router(auth.router, prefix="/api")

    def override_db():
        with factory() as db:
            yield db

    def override_user():
        with factory() as db:
            return db.get(User, 1)

    app.dependency_overrides[get_db] = override_db
    app.dependency_overrides[get_current_user] = override_user
    return engine, factory, TestClient(app)


@pytest.mark.parametrize("source_format", ["pdf", "docx"])
def test_upload_retains_exact_source_and_exposes_private_metadata(monkeypatch, source_format):
    original = _source_bytes(source_format)
    media_type = (
        "application/pdf"
        if source_format == "pdf"
        else "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )

    def parse(data, *, source_format):
        assert data == original
        assert source_format in {"pdf", "docx"}
        return "Original text", {"experience": "Original experience"}, ["Python"], 2.0, {}

    monkeypatch.setattr(resume_router, "inspect_resume_document", parse)
    engine, factory, client = _client()
    try:
        uploaded = client.post(
            "/api/resume/parse",
            files={"file": (f"folder/résumé.{source_format}", original, media_type)},
        )
        assert uploaded.status_code == 200, uploaded.text
        resume_id = uploaded.json()["resume_id"]
        assert uploaded.json()["source_available"] is True
        assert uploaded.json()["source_format"] == source_format

        metadata = client.get(f"/api/resume/{resume_id}")
        assert metadata.status_code == 200
        assert metadata.json()["source_available"] is True
        assert metadata.json()["source_format"] == source_format
        listed = client.get("/api/resume/list").json()["resumes"]
        assert {item["id"] for item in listed} == {10, resume_id}
        assert next(item for item in listed if item["id"] == resume_id)["source_available"] is True

        with factory() as db:
            row = db.get(Resume, resume_id)
            assert row.source_available is True
            assert "source_document" in inspect(row).unloaded

        downloaded = client.get(f"/api/resume/{resume_id}/source")
        assert downloaded.status_code == 200
        assert downloaded.content == original
        assert downloaded.headers["content-type"] == media_type
        assert downloaded.headers["cache-control"] == "private, no-store"
        assert downloaded.headers["x-content-type-options"] == "nosniff"
        assert "filename*=UTF-8''r%C3%A9sum%C3%A9." in downloaded.headers["content-disposition"]

        account = client.get("/api/auth/export-account")
        assert account.status_code == 200, account.text
        exported = next(item for item in account.json()["resumes"] if item["id"] == resume_id)
        assert "source_document" not in exported
        assert exported["source_available"] is True
        assert exported["source_format"] == source_format

        deleted = client.post(
            "/api/auth/delete-account",
            headers={"Authorization": f"Bearer {create_access_token(subject='1')}"},
        )
        assert deleted.status_code == 200, deleted.text
        with factory() as db:
            assert db.get(Resume, resume_id) is None
            assert db.get(Resume, 10) is None
            assert db.get(Resume, 20).source_available is True
    finally:
        client.close()
        engine.dispose()


def test_source_download_handles_legacy_ownership_and_authentication():
    engine, _, client = _client()
    try:
        metadata = client.get("/api/resume/10")
        assert metadata.json()["source_available"] is False
        assert metadata.json()["source_format"] is None
        unavailable = client.get("/api/resume/10/source")
        assert unavailable.status_code == 409
        assert "Upload" in unavailable.json()["detail"]
        assert client.get("/api/resume/20/source").status_code == 404
        assert client.get("/api/resume/999/source").status_code == 404

        client.app.dependency_overrides.pop(get_current_user)
        assert client.get("/api/resume/20/source").status_code == 401
    finally:
        client.close()
        engine.dispose()


def test_tailoring_without_source_fails_before_usage_reservation():
    engine, factory, client = _client()
    try:
        with factory() as db:
            db.add_all(
                [
                    Opportunity(
                        id="opp_legacy_source",
                        user_id=1,
                        resume_id=10,
                        title="Platform Engineer",
                        company="Example Co",
                        job_description="Build reliable Python services.",
                    ),
                    EvidenceItem(
                        id="evd_legacy_source",
                        user_id=1,
                        resume_id=10,
                        category="experience",
                        title="Platform work",
                        evidence_text="Built reliable Python services.",
                        approval_state="approved",
                    ),
                ]
            )
            db.commit()
        with factory() as db, pytest.raises(HTTPException) as raised:
            analysis_service.create_run(
                db,
                user_id=1,
                payload=analysis_schemas.AnalysisRunCreate(
                    operation="resume_tailor", opportunity_id="opp_legacy_source", input={}
                ),
                header_idempotency_key="legacy-source-tailor-001",
            )
        assert raised.value.status_code == 409
        assert "original resume" in raised.value.detail
        with factory() as db:
            assert db.get(User, 1).ai_credits == 50
            assert db.query(UsageEvent).count() == 0
            assert db.query(AnalysisRun).count() == 0
    finally:
        client.close()
        engine.dispose()
