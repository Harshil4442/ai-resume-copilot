from __future__ import annotations

import json
import logging
from io import BytesIO

import pytest
from backend.app.database import Base
from backend.app.domains.analysis import schemas, tasks
from backend.app.domains.analysis import service as analysis_service
from backend.app.models import (
    AnalysisRun,
    EvidenceItem,
    ModelCallEvent,
    Opportunity,
    Resume,
    ResumeVersion,
    UsageEvent,
    User,
)
from backend.app.services import llm_client
from backend.app.services.resume_artifacts import render_resume_version
from backend.app.services.resume_layout import (
    ResumeLayoutError,
    apply_source_edits,
    extract_source_units,
)
from docx import Document as DocxDocument
from fastapi import HTTPException
from pypdf import PdfReader
from reportlab.pdfgen.canvas import Canvas
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

SOURCE_EXPERIENCE = "Developed reliable Python services for customers and improved system performance."
TAILORED_EXPERIENCE = "Built reliable Python services for customers and improved system performance."
PDF_OVERFLOW_ORIGINAL = "Built stable Python interfaces for customers."
PDF_OVERFLOW_REPLACEMENT = "Made stable Python interfaces for customers."


def _native_source():
    document = DocxDocument()
    document.styles["Normal"].font.name = "Times New Roman"
    document.add_heading("Experience", level=2)
    document.add_paragraph(SOURCE_EXPERIENCE)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _database():
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    factory = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    with factory() as db:
        db.add_all(
            [
                User(id=1, email="owner@example.com", ai_credits=50, tier="free"),
                User(id=2, email="other@example.com", ai_credits=50, tier="free"),
                Resume(
                    id=10,
                    user_id=1,
                    original_filename="owner.docx",
                    source_document=_native_source(),
                    source_format="docx",
                    raw_text="Python engineer",
                    skills=["python"],
                    sections={"experience": "Built Python services"},
                ),
                Resume(
                    id=20,
                    user_id=2,
                    original_filename="other.pdf",
                    raw_text="Private resume",
                    skills=["java"],
                    sections={"experience": "Private"},
                ),
            ]
        )
        db.commit()
    return engine, factory


def _payload(*, resume_id: int = 10, title: str = "Platform Engineer"):
    return schemas.AnalysisRunCreate(
        operation="job_match",
        input={
            "resume_id": resume_id,
            "job_title": title,
            "company": "Example Co",
            "job_description": "Build and operate reliable Python services for global customers.",
        },
    )


def test_ownership_is_validated_before_usage_is_reserved():
    engine, factory = _database()
    try:
        with factory() as db:
            try:
                analysis_service.create_run(
                    db,
                    user_id=1,
                    payload=_payload(resume_id=20),
                    header_idempotency_key="ownership-check-001",
                )
                raise AssertionError("Expected ownership failure")
            except HTTPException as exc:
                assert exc.status_code == 404
        with factory() as db:
            assert db.get(User, 1).ai_credits == 50
            assert db.query(UsageEvent).count() == 0
            assert db.query(AnalysisRun).count() == 0
    finally:
        engine.dispose()


def test_create_run_is_idempotent_and_rejects_key_reuse_with_new_input():
    engine, factory = _database()
    try:
        with factory() as db:
            first, created = analysis_service.create_run(
                db,
                user_id=1,
                payload=_payload(),
                header_idempotency_key="match-create-001",
            )
            first_id = first.id
            assert created is True
        with factory() as db:
            duplicate, created = analysis_service.create_run(
                db,
                user_id=1,
                payload=_payload(),
                header_idempotency_key="match-create-001",
            )
            assert created is False
            assert duplicate.id == first_id
            assert db.get(User, 1).ai_credits == 49
            assert db.query(UsageEvent).count() == 1
        with factory() as db:
            try:
                analysis_service.create_run(
                    db,
                    user_id=1,
                    payload=_payload(title="Different role"),
                    header_idempotency_key="match-create-001",
                )
                raise AssertionError("Expected idempotency conflict")
            except HTTPException as exc:
                assert exc.status_code == 409
    finally:
        engine.dispose()


def test_duplicate_task_delivery_commits_usage_once(monkeypatch):
    engine, factory = _database()
    monkeypatch.setattr(tasks, "SessionLocal", factory)
    monkeypatch.setattr(
        tasks,
        "execute_operation",
        lambda db, run: {"match_id": 123, "match_score": 88},
    )
    try:
        with factory() as db:
            run, _ = analysis_service.create_run(
                db,
                user_id=1,
                payload=_payload(),
                header_idempotency_key="duplicate-task-001",
            )
            run_id = run.id

        assert tasks.process_analysis_run(run_id) == "succeeded"
        assert tasks.process_analysis_run(run_id) == "succeeded"

        with factory() as db:
            run = db.get(AnalysisRun, run_id)
            assert run.status == "succeeded"
            assert run.committed_units == 1
            assert db.get(User, 1).ai_credits == 49
            assert [event.event_type for event in db.query(UsageEvent).order_by(UsageEvent.created_at)] == [
                "reserve",
                "commit",
            ]
    finally:
        engine.dispose()


def test_terminal_failure_releases_reserved_usage(monkeypatch):
    engine, factory = _database()
    monkeypatch.setattr(tasks, "SessionLocal", factory)

    def fail(db, run):
        raise ValueError("non-retryable malformed provider response")

    monkeypatch.setattr(tasks, "execute_operation", fail)
    try:
        with factory() as db:
            run, _ = analysis_service.create_run(
                db,
                user_id=1,
                payload=_payload(),
                header_idempotency_key="failed-task-001",
            )
            run_id = run.id

        assert tasks.process_analysis_run(run_id) == "failed"
        with factory() as db:
            run = db.get(AnalysisRun, run_id)
            assert run.status == "failed"
            assert run.usage_state == "released"
            assert db.get(User, 1).ai_credits == 50
            assert [event.event_type for event in db.query(UsageEvent).order_by(UsageEvent.created_at)] == [
                "reserve",
                "release",
            ]
    finally:
        engine.dispose()


def test_repaired_interview_output_saves_eight_questions_and_charges_once(monkeypatch):
    engine, factory = _database()
    monkeypatch.setattr(tasks, "SessionLocal", factory)
    questions = [
        {
            "question": f"How would you improve the reliability of platform service {index}?",
            "coaching_angle": "Describe a concrete method and a verified result.",
            "evidence_ids": [],
        }
        for index in range(1, 9)
    ]
    provider_responses = iter(
        [json.dumps(questions[:1]), f"```json\n{json.dumps(questions)}\n```"]
    )
    chat_calls = []

    def response(messages):
        chat_calls.append(len(messages))
        return next(provider_responses)

    monkeypatch.setattr(llm_client, "_chat", response)
    try:
        with factory() as db:
            db.add(
                Opportunity(
                    id="opp_repaired_interview",
                    user_id=1,
                    resume_id=10,
                    title="Platform Engineer",
                    company="Example Co",
                    job_description="Build and operate reliable Python services for global customers.",
                    job_snapshot={},
                )
            )
            db.commit()
        with factory() as db:
            run, created = analysis_service.create_run(
                db,
                user_id=1,
                payload=schemas.AnalysisRunCreate(
                    operation="interview_questions",
                    opportunity_id="opp_repaired_interview",
                    input={"num_questions": 8},
                ),
                header_idempotency_key="repaired-interview-output-001",
            )
            run_id = run.id
            assert created is True
            assert run.usage_state == "reserved"
            assert db.get(User, 1).ai_credits == 49

        assert tasks.process_analysis_run(run_id) == "succeeded"
        assert tasks.process_analysis_run(run_id) == "succeeded"
        assert len(chat_calls) == 2

        with factory() as db:
            run = db.get(AnalysisRun, run_id)
            assert run.status == "succeeded"
            assert run.usage_state == "committed"
            assert run.committed_units == 1
            assert run.attempt_count == 1
            assert run.error_code is None
            assert run.result_payload["opportunity_id"] == "opp_repaired_interview"
            stored_questions = run.result_payload["questions"]
            assert len(stored_questions) == 8
            assert len({item["question"] for item in stored_questions}) == 8
            assert db.get(User, 1).ai_credits == 49
            events = db.query(UsageEvent).order_by(UsageEvent.created_at).all()
            assert [event.event_type for event in events] == ["reserve", "commit"]
            assert [event.amount for event in events] == [-1, 0]
    finally:
        engine.dispose()


@pytest.mark.parametrize(
    "provider_response",
    [
        "I could not format the interview questions as JSON.",
        '[{"question":"How did you improve a Python service?",'
        '"coaching_angle":"Explain the verified result.","evidence_ids":[]}]',
    ],
    ids=["malformed-response", "incomplete-question-set"],
)
def test_invalid_interview_output_fails_without_charging_or_saving_a_result(
    monkeypatch, provider_response
):
    engine, factory = _database()
    monkeypatch.setattr(tasks, "SessionLocal", factory)
    chat_calls = []

    def invalid_response(*args, **kwargs):
        chat_calls.append(args)
        return provider_response

    monkeypatch.setattr(llm_client, "_chat", invalid_response)
    try:
        with factory() as db:
            db.add(
                Opportunity(
                    id="opp_interview",
                    user_id=1,
                    resume_id=10,
                    title="Platform Engineer",
                    company="Example Co",
                    job_description="Build and operate reliable Python services for global customers.",
                    job_snapshot={},
                )
            )
            db.commit()
        with factory() as db:
            run, created = analysis_service.create_run(
                db,
                user_id=1,
                payload=schemas.AnalysisRunCreate(
                    operation="interview_questions",
                    opportunity_id="opp_interview",
                    input={"num_questions": 8},
                ),
                header_idempotency_key="invalid-interview-output-001",
            )
            run_id = run.id
            assert created is True
            assert run.estimated_units == 1
            assert run.usage_state == "reserved"
            assert db.get(User, 1).ai_credits == 49

        assert tasks.process_analysis_run(run_id) == "failed"
        assert tasks.process_analysis_run(run_id) == "failed"
        assert len(chat_calls) == 2

        with factory() as db:
            run = db.get(AnalysisRun, run_id)
            assert run.status == "failed"
            assert run.usage_state == "released"
            assert run.committed_units == 0
            assert run.result_payload is None
            assert run.completed_at is not None
            assert run.error_code == "InterviewOutputError"
            assert run.attempt_count == 1
            assert db.get(User, 1).ai_credits == 50
            events = db.query(UsageEvent).order_by(UsageEvent.created_at).all()
            assert [event.event_type for event in events] == ["reserve", "release"]
            assert [event.amount for event in events] == [-1, 1]
    finally:
        engine.dispose()


def test_queued_cancellation_releases_usage():
    engine, factory = _database()
    try:
        with factory() as db:
            run, _ = analysis_service.create_run(
                db,
                user_id=1,
                payload=_payload(),
                header_idempotency_key="cancel-task-001",
            )
            run_id = run.id
        with factory() as db:
            cancelled = analysis_service.cancel_run(db, 1, run_id)
            assert cancelled.status == "cancelled"
        with factory() as db:
            assert db.get(User, 1).ai_credits == 50
            assert [event.event_type for event in db.query(UsageEvent).order_by(UsageEvent.created_at)] == [
                "reserve",
                "release",
            ]
    finally:
        engine.dispose()


def test_active_premium_run_is_audited_without_deduction():
    engine, factory = _database()
    try:
        with factory() as db:
            user = db.get(User, 1)
            user.tier = "premium"
            user.premium_until = None
            db.commit()
        with factory() as db:
            run, _ = analysis_service.create_run(
                db,
                user_id=1,
                payload=_payload(),
                header_idempotency_key="premium-task-001",
            )
            assert run.usage_state == "waived"
            assert db.get(User, 1).ai_credits == 50
            assert db.query(UsageEvent).one().event_type == "waive"
    finally:
        engine.dispose()


def test_evidence_tailoring_creates_a_traceable_version_and_commits_once(monkeypatch):
    engine, factory = _database()
    monkeypatch.setattr(tasks, "SessionLocal", factory)
    calls = []

    def tailor(**kwargs):
        calls.append(kwargs)
        return _source_edit_response(kwargs["source_units"])

    monkeypatch.setattr(llm_client, "tailor_resume_from_evidence", tailor)
    try:
        run_id = _create_tailoring_run(factory)

        assert tasks.process_analysis_run(run_id) == "succeeded"
        assert tasks.process_analysis_run(run_id) == "succeeded"
        assert len(calls) == 1
        with factory() as db:
            run = db.get(AnalysisRun, run_id)
            version = db.query(ResumeVersion).one()
            resume = db.get(Resume, 10)
            assert run.status == "succeeded"
            assert run.usage_state == "committed"
            assert run.committed_units == 10
            assert run.attempt_count == 1
            assert run.prompt_version == "resume-source-v4"
            assert db.get(User, 1).ai_credits == 40
            assert version.generation_run_id == run_id
            assert version.evidence_ids == ["evd_approved"]
            assert version.structured_content["evidence_policy"] == "approved_only"
            assert version.structured_content["format_preservation"] == "source"
            assert version.structured_content["source_format"] == "docx"
            artifact = render_resume_version(version, resume, "docx")
            edited = DocxDocument(BytesIO(artifact.content))
            assert edited.paragraphs[-1].text == TAILORED_EXPERIENCE
            assert DocxDocument(BytesIO(resume.source_document)).paragraphs[-1].text == SOURCE_EXPERIENCE
            events = db.query(UsageEvent).order_by(UsageEvent.created_at).all()
            assert [event.event_type for event in events] == ["reserve", "commit"]
            assert [event.amount for event in events] == [-10, 0]
    finally:
        engine.dispose()


def _source_edit_response(units, *, replacement=TAILORED_EXPERIENCE):
    unit = next(item for item in units if item["text"] == SOURCE_EXPERIENCE)
    return {
        "source_edits": [
            {
                "unit_id": unit["unit_id"],
                "original_text": unit["text"],
                "replacement_text": replacement,
                "evidence_ids": ["evd_approved"],
                "reason": "Use concise wording supported by the approved platform experience.",
            }
        ],
        "evidence_needed": [],
        "evidence_policy": "approved_only",
    }


def _create_tailoring_run(factory):
    with factory() as db:
        db.add_all(
            [
                Opportunity(
                    id="opp_tailor",
                    user_id=1,
                    resume_id=10,
                    title="Platform Engineer",
                    company="Example Co",
                    job_description="Build and operate reliable Python services for global customers.",
                    job_snapshot={},
                ),
                EvidenceItem(
                    id="evd_approved",
                    user_id=1,
                    resume_id=10,
                    category="experience",
                    title="Platform work",
                    evidence_text=SOURCE_EXPERIENCE,
                    skills=["python"],
                    metrics={},
                    approval_state="approved",
                ),
            ]
        )
        db.commit()
    with factory() as db:
        run, _ = analysis_service.create_run(
            db,
            user_id=1,
            payload=schemas.AnalysisRunCreate(
                operation="resume_tailor", opportunity_id="opp_tailor", input={}
            ),
            header_idempotency_key="native-tailor-request-001",
        )
        assert run.usage_state == "reserved"
        assert db.get(User, 1).ai_credits == 40
        return run.id


def _configure_pdf_tailoring_source(factory):
    output = BytesIO()
    canvas = Canvas(output, pagesize=(612, 792), invariant=True)
    canvas.setFont("Times-Roman", 11)
    canvas.drawString(60, 650, SOURCE_EXPERIENCE)
    canvas.drawString(60, 620, PDF_OVERFLOW_ORIGINAL)
    canvas.save()
    source = output.getvalue()
    with factory() as db:
        resume = db.get(Resume, 10)
        resume.source_document = source
        resume.source_format = "pdf"
        resume.original_filename = "owner.pdf"
        db.commit()
    units = extract_source_units(source, "pdf")
    assert len(units) == 2
    return source, {unit["text"]: unit for unit in units}


def _pdf_tailoring_edit(unit, replacement):
    return {
        "unit_id": unit["unit_id"],
        "original_text": unit["text"],
        "replacement_text": replacement,
        "evidence_ids": ["evd_approved"],
        "reason": "Use a concise action verb for the approved Python experience.",
    }


def test_mixed_pdf_tailoring_saves_only_the_proven_safe_edits_and_charges_once(monkeypatch):
    engine, factory = _database()
    monkeypatch.setattr(tasks, "SessionLocal", factory)
    calls = []
    try:
        source, units = _configure_pdf_tailoring_source(factory)
        safe = _pdf_tailoring_edit(units[SOURCE_EXPERIENCE], TAILORED_EXPERIENCE)
        unsafe = _pdf_tailoring_edit(units[PDF_OVERFLOW_ORIGINAL], PDF_OVERFLOW_REPLACEMENT)
        unsafe["evidence_ids"] = ["evd_overflow"]
        with factory() as db:
            db.add(
                EvidenceItem(
                    id="evd_overflow", user_id=1, resume_id=10, category="experience",
                    title="Interface work", evidence_text=PDF_OVERFLOW_ORIGINAL,
                    skills=["python"], metrics={}, approval_state="approved",
                )
            )
            db.commit()
        # Fewer characters can still overflow a fixed native-font text slot.
        assert len(unsafe["replacement_text"]) < len(unsafe["original_text"])

        def response(messages):
            calls.append(messages)
            return json.dumps({"source_edits": [safe, unsafe], "evidence_needed": []})

        monkeypatch.setattr(llm_client, "_chat", response)
        run_id = _create_tailoring_run(factory)
        assert tasks.process_analysis_run(run_id) == "succeeded"
        assert tasks.process_analysis_run(run_id) == "succeeded"
        assert len(calls) == 1
        with factory() as db:
            run = db.get(AnalysisRun, run_id)
            resume = db.get(Resume, 10)
            version = db.query(ResumeVersion).one()
            assert run.status == "succeeded" and run.attempt_count == 1
            assert run.usage_state == "committed" and run.committed_units == 10
            assert run.prompt_version == "resume-source-v4"
            assert db.get(User, 1).ai_credits == 40
            assert version.generation_run_id == run_id
            assert version.structured_content["source_edits"] == [safe]
            assert run.result_payload["content"]["source_edits"] == [safe]
            assert version.evidence_ids == ["evd_approved"]
            assert resume.source_document == source
            artifact = render_resume_version(version, resume, "pdf")
            exported_text = PdfReader(BytesIO(artifact.content)).pages[0].extract_text()
            assert TAILORED_EXPERIENCE in exported_text
            assert SOURCE_EXPERIENCE not in exported_text
            assert PDF_OVERFLOW_ORIGINAL in exported_text
            assert PDF_OVERFLOW_REPLACEMENT not in exported_text
            events = db.query(UsageEvent).order_by(UsageEvent.created_at).all()
            assert [event.event_type for event in events] == ["reserve", "commit"]
            assert [event.amount for event in events] == [-10, 0]
            assert db.query(ModelCallEvent).one().status == "succeeded"
    finally:
        engine.dispose()


def test_all_pdf_tailoring_edits_overflow_twice_without_saving_and_refund_once(monkeypatch):
    engine, factory = _database()
    monkeypatch.setattr(tasks, "SessionLocal", factory)
    calls = []
    try:
        source, units = _configure_pdf_tailoring_source(factory)
        unsafe = _pdf_tailoring_edit(units[PDF_OVERFLOW_ORIGINAL], PDF_OVERFLOW_REPLACEMENT)

        def response(messages):
            calls.append(messages)
            return json.dumps({"source_edits": [unsafe], "evidence_needed": []})

        monkeypatch.setattr(llm_client, "_chat", response)
        run_id = _create_tailoring_run(factory)
        assert tasks.process_analysis_run(run_id) == "failed"
        assert tasks.process_analysis_run(run_id) == "failed"
        assert len(calls) == 2
        assert unsafe["unit_id"] in calls[1][1]["content"]
        assert "does not fit its original PDF text slot" in calls[1][1]["content"]
        with factory() as db:
            run = db.get(AnalysisRun, run_id)
            assert run.status == "failed" and run.attempt_count == 1
            assert run.usage_state == "released" and run.committed_units == 0
            assert run.error_code == "TailoringOutputError"
            assert run.result_payload is None
            assert db.query(ResumeVersion).count() == 0
            assert db.query(ModelCallEvent).count() == 0
            assert db.get(Resume, 10).source_document == source
            assert db.get(User, 1).ai_credits == 50
            events = db.query(UsageEvent).order_by(UsageEvent.created_at).all()
            assert [event.event_type for event in events] == ["reserve", "release"]
            assert [event.amount for event in events] == [-10, 10]
    finally:
        engine.dispose()


def test_mixed_pdf_joining_edit_preserves_the_sentence_and_charges_only_once(monkeypatch):
    engine, factory = _database()
    monkeypatch.setattr(tasks, "SessionLocal", factory)
    joining_original = "for CPU architecture validation, enabling "
    joining_replacement = "for CPU validation, enabling reliability "
    calls = []
    try:
        output = BytesIO()
        canvas = Canvas(output, pagesize=(612, 792), invariant=True)
        canvas.setFont("Times-Roman", 11)
        canvas.drawString(60, 650, SOURCE_EXPERIENCE)
        leading = "Built checks "
        canvas.drawString(60, 620, leading)
        cursor = 60 + canvas.stringWidth(leading, "Times-Roman", 11)
        canvas.drawString(cursor, 620, joining_original)
        cursor += canvas.stringWidth(joining_original, "Times-Roman", 11)
        canvas.setFont("Times-Bold", 11)
        canvas.drawString(cursor, 620, "50+ engineers")
        canvas.save()
        source = output.getvalue()
        units = {unit["text"].strip(): unit for unit in extract_source_units(source, "pdf")}
        joining_unit = units[joining_original.strip()]
        assert joining_unit["required_prefix"] == "for CPU"
        assert joining_unit["required_suffix"] == "validation, enabling"
        safe = _pdf_tailoring_edit(units[SOURCE_EXPERIENCE], TAILORED_EXPERIENCE)
        unsafe = _pdf_tailoring_edit(joining_unit, joining_replacement)
        unsafe["evidence_ids"] = ["evd_joining"]
        with factory() as db:
            resume = db.get(Resume, 10)
            resume.source_document = source
            resume.source_format = "pdf"
            resume.original_filename = "owner.pdf"
            db.add(
                EvidenceItem(
                    id="evd_joining", user_id=1, resume_id=10, category="experience",
                    title="Validation work", evidence_text=joining_original,
                    skills=["CPU"], metrics={}, approval_state="approved",
                )
            )
            db.commit()

        def response(messages):
            calls.append(messages)
            return json.dumps({"source_edits": [safe, unsafe], "evidence_needed": []})

        monkeypatch.setattr(llm_client, "_chat", response)
        run_id = _create_tailoring_run(factory)
        assert tasks.process_analysis_run(run_id) == "succeeded"
        assert tasks.process_analysis_run(run_id) == "succeeded"
        assert len(calls) == 1
        with factory() as db:
            run = db.get(AnalysisRun, run_id)
            resume = db.get(Resume, 10)
            version = db.query(ResumeVersion).one()
            assert run.status == "succeeded" and run.attempt_count == 1
            assert run.usage_state == "committed" and run.committed_units == 10
            assert run.result_payload["content"]["source_edits"] == [safe]
            assert version.structured_content["source_edits"] == [safe]
            assert version.evidence_ids == ["evd_approved"]
            assert "rejected_source_edits" not in version.structured_content
            assert resume.source_document == source
            artifact = render_resume_version(version, resume, "pdf")
            exported_text = PdfReader(BytesIO(artifact.content)).pages[0].extract_text()
            assert TAILORED_EXPERIENCE in exported_text
            assert joining_original.strip() in exported_text
            assert joining_replacement.strip() not in exported_text
            assert "50+ engineers" in exported_text
            assert db.get(User, 1).ai_credits == 40
            events = db.query(UsageEvent).order_by(UsageEvent.created_at).all()
            assert [event.event_type for event in events] == ["reserve", "commit"]
            assert [event.amount for event in events] == [-10, 0]
    finally:
        engine.dispose()


def test_all_invalid_llm_edit_repair_hint_reaches_second_request_and_charges_once(monkeypatch):
    engine, factory = _database()
    monkeypatch.setattr(tasks, "SessionLocal", factory)
    calls = []
    errors = []
    real_tailor = llm_client.tailor_resume_from_evidence

    def capture_hint(**kwargs):
        try:
            return real_tailor(**kwargs)
        except llm_client.TailoringOutputError as error:
            errors.append(error)
            raise

    monkeypatch.setattr(llm_client, "tailor_resume_from_evidence", capture_hint)
    try:
        run_id = _create_tailoring_run(factory)
        with factory() as db:
            units = extract_source_units(db.get(Resume, 10).source_document, "docx")
        invalid = _source_edit_response(units, replacement=SOURCE_EXPERIENCE)
        valid = _source_edit_response(units)
        responses = iter([json.dumps(invalid), json.dumps(valid)])

        def response(messages):
            calls.append(messages)
            return next(responses)

        monkeypatch.setattr(llm_client, "_chat", response)
        assert tasks.process_analysis_run(run_id) == "succeeded"
        assert tasks.process_analysis_run(run_id) == "succeeded"
        assert len(calls) == 2 and len(errors) == 1
        repair_hint = errors[0].repair_hint
        assert repair_hint
        assert repair_hint in calls[1][1]["content"]
        assert "previous proposed changes could not be applied" not in calls[0][1]["content"]
        with factory() as db:
            run = db.get(AnalysisRun, run_id)
            assert run.usage_state == "committed" and run.committed_units == 10
            assert run.attempt_count == 1
            assert db.query(ResumeVersion).one().structured_content["source_edits"] == valid["source_edits"]
            assert db.get(User, 1).ai_credits == 40
            events = db.query(UsageEvent).order_by(UsageEvent.created_at).all()
            assert [event.event_type for event in events] == ["reserve", "commit"]
            assert [event.amount for event in events] == [-10, 0]
    finally:
        engine.dispose()


def test_empty_tailoring_output_leaves_one_content_repair_and_charges_once(monkeypatch):
    engine, factory = _database()
    monkeypatch.setattr(tasks, "SessionLocal", factory)
    calls = []
    errors = []
    real_tailor = llm_client.tailor_resume_from_evidence

    def capture_hint(**kwargs):
        try:
            return real_tailor(**kwargs)
        except llm_client.TailoringOutputError as error:
            errors.append(error)
            raise

    monkeypatch.setattr(llm_client, "tailor_resume_from_evidence", capture_hint)
    try:
        run_id = _create_tailoring_run(factory)
        with factory() as db:
            source = db.get(Resume, 10).source_document
            units = extract_source_units(source, "docx")
        overlength = _source_edit_response(units, replacement=TAILORED_EXPERIENCE * 2)
        valid = _source_edit_response(units)
        responses = iter([
            json.dumps({"source_edits": [], "evidence_needed": []}),
            json.dumps(overlength),
            json.dumps(valid),
        ])

        def response(messages):
            calls.append(messages)
            return next(responses)

        monkeypatch.setattr(llm_client, "_chat", response)
        assert tasks.process_analysis_run(run_id) == "succeeded"
        assert tasks.process_analysis_run(run_id) == "succeeded"
        assert len(calls) == 3 and len(errors) == 2
        assert errors[0].reason_code == "invalid_edit_count"
        assert errors[0].repair_hint in calls[1][1]["content"]
        length_error = errors[1]
        assert length_error.reason_code == "invalid_replacement_length"
        assert length_error.unit_id == valid["source_edits"][0]["unit_id"]
        assert length_error.repair_hint in calls[2][1]["content"]
        assert length_error.unit_id in calls[2][1]["content"]
        assert (
            f"Original final length is {len(SOURCE_EXPERIENCE)} characters; "
            f"proposed final length is {len(TAILORED_EXPERIENCE) * 2}."
        ) in calls[2][1]["content"]
        with factory() as db:
            run = db.get(AnalysisRun, run_id)
            resume = db.get(Resume, 10)
            version = db.query(ResumeVersion).one()
            assert run.status == "succeeded" and run.attempt_count == 1
            assert run.usage_state == "committed" and run.committed_units == 10
            assert run.error_code is None
            assert run.result_payload["resume_version_id"] == version.id
            assert version.structured_content["source_edits"] == valid["source_edits"]
            assert resume.source_document == source
            artifact = render_resume_version(version, resume, "docx")
            assert DocxDocument(BytesIO(artifact.content)).paragraphs[-1].text == TAILORED_EXPERIENCE
            assert db.get(User, 1).ai_credits == 40
            events = db.query(UsageEvent).order_by(UsageEvent.created_at).all()
            assert [event.event_type for event in events] == ["reserve", "commit"]
            assert [event.amount for event in events] == [-10, 0]
            assert db.query(ModelCallEvent).one().status == "succeeded"
    finally:
        engine.dispose()


def test_changed_tailoring_constraint_leaves_one_native_width_repair_and_charges_once(monkeypatch):
    engine, factory = _database()
    monkeypatch.setattr(tasks, "SessionLocal", factory)
    calls = []
    try:
        source, units = _configure_pdf_tailoring_source(factory)
        safe = _pdf_tailoring_edit(units[SOURCE_EXPERIENCE], TAILORED_EXPERIENCE)
        overlength = _pdf_tailoring_edit(units[SOURCE_EXPERIENCE], TAILORED_EXPERIENCE * 2)
        overflow = _pdf_tailoring_edit(units[PDF_OVERFLOW_ORIGINAL], PDF_OVERFLOW_REPLACEMENT)
        assert overlength["unit_id"] != overflow["unit_id"]
        with pytest.raises(ResumeLayoutError, match="does not fit its original PDF text slot") as native_error:
            apply_source_edits(source, "pdf", [overflow])
        width_hint = native_error.value.repair_hint
        assert width_hint and "width" in width_hint.lower()
        assert native_error.value.unit_id == overflow["unit_id"]
        responses = iter([
            json.dumps({"source_edits": [overlength], "evidence_needed": []}),
            json.dumps({"source_edits": [overflow], "evidence_needed": []}),
            json.dumps({"source_edits": [safe], "evidence_needed": []}),
        ])

        def response(messages):
            calls.append(messages)
            return next(responses)

        monkeypatch.setattr(llm_client, "_chat", response)
        run_id = _create_tailoring_run(factory)
        assert tasks.process_analysis_run(run_id) == "succeeded"
        assert tasks.process_analysis_run(run_id) == "succeeded"
        assert len(calls) == 3
        assert (
            f"Original final length is {len(SOURCE_EXPERIENCE)} characters; "
            f"proposed final length is {len(TAILORED_EXPERIENCE) * 2}."
        ) in calls[1][1]["content"]
        assert overflow["unit_id"] in calls[2][1]["content"]
        assert width_hint in calls[2][1]["content"]
        with factory() as db:
            run = db.get(AnalysisRun, run_id)
            resume = db.get(Resume, 10)
            version = db.query(ResumeVersion).one()
            assert run.status == "succeeded" and run.attempt_count == 1
            assert run.usage_state == "committed" and run.committed_units == 10
            assert run.result_payload["content"]["source_edits"] == [safe]
            assert version.structured_content["source_edits"] == [safe]
            assert resume.source_document == source
            artifact = render_resume_version(version, resume, "pdf")
            exported_text = PdfReader(BytesIO(artifact.content)).pages[0].extract_text()
            assert TAILORED_EXPERIENCE in exported_text
            assert PDF_OVERFLOW_ORIGINAL in exported_text
            assert PDF_OVERFLOW_REPLACEMENT not in exported_text
            assert db.get(User, 1).ai_credits == 40
            events = db.query(UsageEvent).order_by(UsageEvent.created_at).all()
            assert [event.event_type for event in events] == ["reserve", "commit"]
            assert [event.amount for event in events] == [-10, 0]
    finally:
        engine.dispose()


@pytest.mark.parametrize("failure", ["malformed", "unsafe-layout"])
def test_invalid_native_tailoring_fails_without_version_and_refunds_once(monkeypatch, failure):
    engine, factory = _database()
    monkeypatch.setattr(tasks, "SessionLocal", factory)
    repair_notes = []

    def invalid_tailor(**kwargs):
        repair_notes.append(kwargs["repair_note"])
        if failure == "malformed":
            return {"source_edits": "not an edit list"}
        return _source_edit_response(
            kwargs["source_units"], replacement=TAILORED_EXPERIENCE * 5
        )

    monkeypatch.setattr(llm_client, "tailor_resume_from_evidence", invalid_tailor)
    try:
        run_id = _create_tailoring_run(factory)
        assert tasks.process_analysis_run(run_id) == "failed"
        assert tasks.process_analysis_run(run_id) == "failed"
        assert len(repair_notes) == 2
        assert repair_notes[0] == ""
        assert repair_notes[1]
        with factory() as db:
            run = db.get(AnalysisRun, run_id)
            assert run.status == "failed"
            assert run.usage_state == "released"
            assert run.committed_units == 0
            assert run.result_payload is None
            assert run.error_code == "TailoringOutputError"
            assert run.attempt_count == 1
            assert db.query(ResumeVersion).count() == 0
            assert db.get(User, 1).ai_credits == 50
            events = db.query(UsageEvent).order_by(UsageEvent.created_at).all()
            assert [event.event_type for event in events] == ["reserve", "release"]
            assert [event.amount for event in events] == [-10, 10]
    finally:
        engine.dispose()


def test_native_tailoring_repairs_an_unsafe_edit_and_commits_one_charge(monkeypatch):
    engine, factory = _database()
    monkeypatch.setattr(tasks, "SessionLocal", factory)
    repair_notes = []

    def repaired_tailor(**kwargs):
        repair_notes.append(kwargs["repair_note"])
        replacement = TAILORED_EXPERIENCE * 5 if len(repair_notes) == 1 else TAILORED_EXPERIENCE
        return _source_edit_response(kwargs["source_units"], replacement=replacement)

    monkeypatch.setattr(llm_client, "tailor_resume_from_evidence", repaired_tailor)
    try:
        run_id = _create_tailoring_run(factory)
        assert tasks.process_analysis_run(run_id) == "succeeded"
        assert tasks.process_analysis_run(run_id) == "succeeded"
        assert len(repair_notes) == 2
        assert repair_notes[0] == ""
        assert repair_notes[1]
        with factory() as db:
            run = db.get(AnalysisRun, run_id)
            assert run.usage_state == "committed"
            assert run.committed_units == 10
            assert run.attempt_count == 1
            assert db.get(User, 1).ai_credits == 40
            version = db.query(ResumeVersion).one()
            assert version.structured_content["source_edits"][0]["replacement_text"] == TAILORED_EXPERIENCE
            events = db.query(UsageEvent).order_by(UsageEvent.created_at).all()
            assert [event.event_type for event in events] == ["reserve", "commit"]
            assert [event.amount for event in events] == [-10, 0]
    finally:
        engine.dispose()


def test_failed_tailoring_keeps_layout_diagnostics_and_cause_after_rollback(monkeypatch, caplog):
    engine, factory = _database()
    monkeypatch.setattr(tasks, "SessionLocal", factory)
    repair_notes = []
    unit_ids = []
    errors = []
    execute = tasks.execute_operation

    def invalid_tailor(**kwargs):
        repair_notes.append(kwargs["repair_note"])
        response = _source_edit_response(
            kwargs["source_units"], replacement=TAILORED_EXPERIENCE * 5
        )
        unit_ids.append(response["source_edits"][0]["unit_id"])
        return response

    def capture_failure(db, run):
        try:
            return execute(db, run)
        except llm_client.TailoringOutputError as exc:
            errors.append(exc)
            raise

    monkeypatch.setattr(llm_client, "tailor_resume_from_evidence", invalid_tailor)
    monkeypatch.setattr(tasks, "execute_operation", capture_failure)
    try:
        run_id = _create_tailoring_run(factory)
        with caplog.at_level(logging.WARNING, logger="hirewiz.analysis.operations"):
            assert tasks.process_analysis_run(run_id) == "failed"
            assert tasks.process_analysis_run(run_id) == "failed"
        assert len(errors) == 1
        cause = errors[0].__cause__
        assert isinstance(cause, ResumeLayoutError)
        assert cause.unit_id == unit_ids[0]
        assert repair_notes[0] == ""
        assert repair_notes[1].startswith(f"Source unit {unit_ids[0]}:")
        diagnostic_messages = [
            record.getMessage() for record in caplog.records
            if record.name == "hirewiz.analysis.operations"
        ]
        assert len(diagnostic_messages) == 2
        assert all(f"run_id={run_id}" in message for message in diagnostic_messages)
        assert all("type=ResumeLayoutError" in message for message in diagnostic_messages)
        assert all("reason=docx_paragraph_overflow" in message for message in diagnostic_messages)
        assert all(f"unit_id={unit_ids[0]}" in message for message in diagnostic_messages)
        assert all("source_units=1 edits=1" in message for message in diagnostic_messages)
        assert SOURCE_EXPERIENCE not in caplog.text
        assert TAILORED_EXPERIENCE not in caplog.text
        with factory() as db:
            run = db.get(AnalysisRun, run_id)
            assert run.usage_state == "released" and run.committed_units == 0
            assert run.error_code == "TailoringOutputError"
            assert run.error_message == "Could not apply useful changes while preserving the resume format"
            assert run.result_payload is None
            assert db.query(ResumeVersion).count() == 0
            assert db.query(ModelCallEvent).count() == 0
            assert db.get(User, 1).ai_credits == 50
            events = db.query(UsageEvent).order_by(UsageEvent.created_at).all()
            assert [event.event_type for event in events] == ["reserve", "release"]
            assert [event.amount for event in events] == [-10, 10]
    finally:
        engine.dispose()


def test_zero_edit_provider_output_keeps_private_diagnostics_and_refunds(monkeypatch, caplog):
    engine, factory = _database()
    monkeypatch.setattr(tasks, "SessionLocal", factory)
    calls = []
    errors = []
    execute = tasks.execute_operation

    def zero_edits(messages):
        calls.append(messages)
        return json.dumps({"source_edits": [], "evidence_needed": []})

    def capture_failure(db, run):
        try:
            return execute(db, run)
        except llm_client.TailoringOutputError as exc:
            errors.append(exc)
            raise

    monkeypatch.setattr(llm_client, "_chat", zero_edits)
    monkeypatch.setattr(tasks, "execute_operation", capture_failure)
    try:
        run_id = _create_tailoring_run(factory)
        with caplog.at_level(logging.WARNING, logger="hirewiz.analysis.operations"):
            assert tasks.process_analysis_run(run_id) == "failed"
        assert len(calls) == 2 and len(errors) == 1
        assert isinstance(errors[0].__cause__, llm_client.TailoringOutputError)
        assert str(errors[0].__cause__) == "Return between one and twelve useful source replacements"
        diagnostics = [record.getMessage() for record in caplog.records
                       if record.name == "hirewiz.analysis.operations"]
        assert len(diagnostics) == 2
        assert all("reason=invalid_edit_count" in message for message in diagnostics)
        assert all("unit_id=none" in message for message in diagnostics)
        assert SOURCE_EXPERIENCE not in caplog.text
        assert TAILORED_EXPERIENCE not in caplog.text
        with factory() as db:
            run = db.get(AnalysisRun, run_id)
            assert run.usage_state == "released" and run.committed_units == 0
            assert db.query(ResumeVersion).count() == 0
            assert db.query(ModelCallEvent).count() == 0
            assert db.get(User, 1).ai_credits == 50
    finally:
        engine.dispose()


def test_tailoring_diagnostics_and_repair_exclude_unknown_unit_identity(monkeypatch, caplog):
    engine, factory = _database()
    monkeypatch.setattr(tasks, "SessionLocal", factory)
    unknown_unit = "untrusted-private-unit@example.com"
    repair_notes = []

    def invalid_tailor(**kwargs):
        repair_notes.append(kwargs["repair_note"])
        error = llm_client.TailoringOutputError("Return a JSON object containing source_edits")
        error.unit_id = unknown_unit
        raise error

    monkeypatch.setattr(llm_client, "tailor_resume_from_evidence", invalid_tailor)
    try:
        run_id = _create_tailoring_run(factory)
        with caplog.at_level(logging.WARNING, logger="hirewiz.analysis.operations"):
            assert tasks.process_analysis_run(run_id) == "failed"
        assert len(repair_notes) == 2
        assert unknown_unit not in " ".join(repair_notes)
        assert unknown_unit not in caplog.text
        assert "reason=invalid_output_format unit_id=none" in caplog.text
    finally:
        engine.dispose()


def test_evidence_tailoring_validates_approval_before_reservation():
    engine, factory = _database()
    try:
        with factory() as db:
            db.add(
                Opportunity(
                    id="opp_no_evidence",
                    user_id=1,
                    resume_id=10,
                    title="Platform Engineer",
                    company="Example Co",
                    job_description="Build and operate reliable Python services for global customers.",
                    job_snapshot={},
                )
            )
            db.commit()
        with factory() as db:
            try:
                analysis_service.create_run(
                    db,
                    user_id=1,
                    payload=schemas.AnalysisRunCreate(
                        operation="resume_tailor",
                        opportunity_id="opp_no_evidence",
                        input={},
                    ),
                    header_idempotency_key="test-request",
                )
                raise AssertionError("Expected evidence validation failure")
            except HTTPException as exc:
                assert exc.status_code == 422
        with factory() as db:
            assert db.get(User, 1).ai_credits == 50
            assert db.query(AnalysisRun).count() == 0
            assert db.query(UsageEvent).count() == 0
    finally:
        engine.dispose()
