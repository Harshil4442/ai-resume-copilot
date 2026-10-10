"""Structural/parser and transport tests; fake ClamAV is never engine proof."""
from __future__ import annotations

import base64
import hashlib
import http.client
import importlib.util
import io
import json
import socket
import subprocess
import sys
import tarfile
import threading
import time
from datetime import UTC, datetime, timedelta
from pathlib import Path
from zipfile import ZipFile, ZipInfo

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.errors import PdfReadError
from pypdf.generic import DictionaryObject, NameObject, TextStringObject

ROOT = Path(__file__).resolve().parents[2]


def load(name):
    spec = importlib.util.spec_from_file_location("document_worker_" + name,
        ROOT / "infra" / "document-processing" / (name + ".py"))
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@pytest.fixture
def processor():
    return load("processor")


@pytest.fixture
def service(monkeypatch):
    module = load("service")
    monkeypatch.setenv("DOCUMENT_WORKER_IMAGE_DIGEST", "registry.invalid/document@sha256:" + "a" * 64)
    monkeypatch.setenv("DOCUMENT_WORKER_POLICY_SHA256", "b" * 64)
    return module


def pdf(*, action=None, attachment=False):
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    if action:
        writer._root_object[NameObject("/OpenAction")] = DictionaryObject({
            NameObject("/S"): NameObject(action), NameObject("/JS"): TextStringObject("private-marker")})
    if attachment:
        writer.add_attachment("private-file", b"private-marker")
    stream = io.BytesIO()
    writer.write(stream)
    return stream.getvalue()


def docx():
    document = Document()
    document.add_paragraph("Harshil Candidate")
    document.add_paragraph("SKILLS")
    document.add_paragraph("Python, FastAPI, PostgreSQL")
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def payload(content, kind="pdf", **changes):
    value = {"version": 1, "source_format": kind, "content_base64": base64.b64encode(content).decode(),
             "sha256": hashlib.sha256(content).hexdigest()}
    value.update(changes)
    return json.dumps(value).encode()


@pytest.mark.parametrize("changes", [{"version": True}, {"version": 2}, {"source_format": "tex"},
    {"sha256": "f" * 64}, {"content_base64": "!!"}, {"filename": "private-marker"}])
def test_envelope_refuses_without_any_scan(processor, monkeypatch, changes):
    monkeypatch.setattr(processor, "scan_original", lambda _: pytest.fail("scanner reached"))
    with pytest.raises(processor.InspectionUnavailable, match="^document_inspection_unavailable$"):
        processor.inspect(payload(b"original", **changes))


@pytest.mark.parametrize("raw", [b"{}", b'{"version":1,"version":1}', b'NaN', b'[]', b'\xff'])
def test_strict_request_json(processor, raw):
    with pytest.raises(processor.InspectionUnavailable):
        processor.decode_request(raw)


def test_empty_and_oversize_original_refused(processor):
    for content in (b"", b"x" * (processor.MAX_FILE + 1)):
        with pytest.raises(processor.InspectionDenied):
            processor.decode_request(payload(content))


def test_scan_failure_precedes_all_document_parsers(processor, monkeypatch):
    def deny(_):
        raise processor.InspectionDenied()
    monkeypatch.setattr(processor, "scan_original", deny)
    monkeypatch.setattr(processor, "validate_pdf", lambda _: pytest.fail("parser reached"))
    with pytest.raises(processor.InspectionDenied):
        processor.inspect(payload(b"untrusted"))


@pytest.mark.parametrize("kind", ["pdf", "docx"])
def test_real_structure_and_deterministic_parser_after_fake_clean_scan(processor, monkeypatch, kind):
    # This exercises genuine parsers; the scan is explicitly synthetic.
    calls = []
    def fake_scan(path):
        calls.append(path.read_bytes())
        return {"engine": "ClamAV", "version": "1.4.3", "definitions": "1/2026-10-10T00:00:00+00:00"}
    monkeypatch.setattr(processor, "scan_original", fake_scan)
    content = pdf() if kind == "pdf" else docx()
    result = processor.inspect(payload(content, kind))
    assert calls == [content]
    assert result["sha256"] == hashlib.sha256(content).hexdigest()
    assert result["size_bytes"] == len(content)
    assert result["parsed"]["experience_years"] == 0
    if kind == "docx":
        assert {"Python", "FastAPI", "PostgreSQL"} <= set(result["parsed"]["skills"])
        assert result["parsed"]["contact_info"]["name"] == "Harshil Candidate"


@pytest.mark.parametrize("content", [b"%PDF-malformed", pdf(action="/JavaScript"), pdf(attachment=True)])
def test_pdf_active_and_malformed_content_refused(processor, content):
    with pytest.raises((processor.InspectionDenied, PdfReadError)):
        processor.validate_pdf(content)


def test_pdf_page_and_encryption_bounds(processor):
    for encrypted in (False, True):
        writer = PdfWriter()
        for _ in range(1 if encrypted else 21):
            writer.add_blank_page(200, 200)
        if encrypted:
            writer.encrypt("private-marker")
        stream = io.BytesIO()
        writer.write(stream)
        with pytest.raises(processor.InspectionDenied):
            processor.validate_pdf(stream.getvalue())


def replace_zip(original, name, value):
    output = io.BytesIO()
    with ZipFile(io.BytesIO(original)) as old, ZipFile(output, "w") as new:
        for entry in old.infolist():
            new.writestr(entry, value if entry.filename == name else old.read(entry))
        if name not in old.namelist():
            new.writestr(name, value)
    return output.getvalue()


@pytest.mark.parametrize("name,value", [
    ("../escape", b"private-marker"), ("word/vbaProject.bin", b"private-marker"),
    ("word/embeddings/payload", b"private-marker"),
    ("word/document.xml", b'<!DOCTYPE document [<!ENTITY secret SYSTEM "file:///private-marker">]><document/>'),
    ("word/document.xml", b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:altChunk/></w:document>'),
    ("word/_rels/document.xml.rels", b'<Relationships><Relationship TargetMode="External" Target="https://private.invalid" Type="template"/></Relationships>')])
def test_docx_unsafe_structures(processor, name, value):
    with pytest.raises(processor.InspectionDenied):
        processor.validate_docx(replace_zip(docx(), name, value))


def test_docx_link_duplicate_and_utf16_dtd_refused(processor):
    for variant in ("link", "duplicate", "utf16"):
        output = io.BytesIO()
        with ZipFile(io.BytesIO(docx())) as old, ZipFile(output, "w") as new:
            for entry in old.infolist():
                new.writestr(entry, old.read(entry))
            if variant == "link":
                entry = ZipInfo("word/linked.xml")
                entry.external_attr = 0o120777 << 16
                new.writestr(entry, b"/private-marker")
            elif variant == "duplicate":
                new.writestr("word/document.xml", b"private-marker")
            else:
                new.writestr("word/unsafe.xml", ('<?xml version="1.0" encoding="UTF-16"?>'
                    '<!DOCTYPE x [<!ENTITY secret SYSTEM "file:///private-marker">]><x/>').encode("utf-16"))
        with pytest.raises(processor.InspectionDenied):
            processor.validate_docx(output.getvalue())


def test_changed_scan_target_never_parsed(processor, monkeypatch):
    def change(path):
        path.write_bytes(b"changed-private-marker")
        return {"engine": "ClamAV", "version": "1.4.3", "definitions": "synthetic"}
    monkeypatch.setattr(processor, "scan_original", change)
    monkeypatch.setattr(processor, "validate_pdf", lambda _: pytest.fail("parser reached"))
    with pytest.raises(processor.InspectionUnavailable):
        processor.inspect(payload(pdf()))


@pytest.mark.parametrize("large_field", ["raw", "sections", "contact"])
def test_multilingual_byte_budget_refuses_before_result(processor, monkeypatch, large_field):
    from app.services import parsing
    monkeypatch.setattr(processor, "scan_original", lambda _: {"engine": "ClamAV", "version": "synthetic", "definitions": "synthetic"})
    monkeypatch.setattr(processor, "validate_pdf", lambda _: None)
    raw = "界" * 90000 if large_field == "raw" else ""
    sections = {"summary": "界" * 90000} if large_field == "sections" else {}
    contact = {"name": "界" * 334 if large_field == "contact" else None,
               "email": None, "phone": None, "linkedin": None, "github": None}
    monkeypatch.setattr(parsing, "parse_resume_file", lambda *_a, **_k: (raw, sections, [], 0, contact))
    with pytest.raises(processor.InspectionDenied):
        processor.inspect(payload(pdf()))


def scanner_output(path, **changes):
    fields = {"Known viruses": "100", "Engine version": "1.4.3", "Scanned directories": "0",
              "Scanned files": "1", "Infected files": "0"}
    fields.update(changes)
    return str(path) + ": OK\n" + "\n".join(k + ": " + v for k, v in fields.items())


def test_scan_requires_exact_clean_one_file_and_bounded_flags(processor, monkeypatch, tmp_path):
    path = tmp_path / "original.pdf"
    calls = []
    def command(args, timeout, **options):
        calls.append((args, timeout))
        return 0, "ClamAV 1.4.3/28000/Sat Oct 10 00:00:00 2026\n" if "--version" in args else scanner_output(path)
    monkeypatch.setattr(processor, "_command", command)
    result = processor.scan_original(path, now=datetime(2026, 10, 10, tzinfo=UTC))
    assert result["engine"] == "ClamAV"
    scan = calls[1][0]
    assert scan[-2:] == ["--", str(path)]
    assert {"--alert-exceeds-max=yes", "--alert-encrypted=yes", "--fail-if-cvd-older-than=3",
            "--max-filesize=6M", "--max-scansize=20M", "--max-files=500", "--max-recursion=16"} <= set(scan)
    assert len(calls) == 2


@pytest.mark.parametrize("change", [{"Scanned files": "0"}, {"Scanned files": "2"},
    {"Known viruses": "0"}, {"Infected files": "1"}, {"Engine version": "different"}, {"Total errors": "1"}])
def test_skipped_or_incomplete_scan_not_clean(processor, monkeypatch, tmp_path, change):
    path = tmp_path / "original.pdf"
    monkeypatch.setattr(processor, "_command", lambda args, _, **options: (0, "ClamAV 1.4.3/28000/Sat Oct 10 00:00:00 2026\n"
        if "--version" in args else scanner_output(path, **change)))
    with pytest.raises(processor.InspectionUnavailable):
        processor.scan_original(path, now=datetime(2026, 10, 10, tzinfo=UTC))


@pytest.mark.parametrize("age", [timedelta(days=3, seconds=1), timedelta(seconds=-1)])
def test_stale_and_future_definitions_refused_before_scan(processor, monkeypatch, tmp_path, age):
    calls = []
    def command(args, _, **options):
        calls.append(args)
        return 0, "ClamAV 1.4.3/28000/Sat Oct 10 00:00:00 2026\n"
    monkeypatch.setattr(processor, "_command", command)
    with pytest.raises(processor.InspectionUnavailable):
        processor.scan_original(tmp_path / "original.pdf", now=datetime(2026, 10, 10, tzinfo=UTC) + age)
    assert len(calls) == 1


@pytest.mark.parametrize("failure", [FileNotFoundError("private-marker"), subprocess.TimeoutExpired("private-marker", 30)])
def test_scanner_missing_timeout_fixed_private_errors(processor, monkeypatch, failure):
    monkeypatch.setattr(processor.subprocess, "run", lambda *_args, **_kwargs: (_ for _ in ()).throw(failure))
    with pytest.raises(processor.InspectionUnavailable, match="^document_inspection_unavailable$") as caught:
        processor._command(["/usr/bin/clamscan"], 1)
    assert caught.value.__cause__ is None
    assert "private-marker" not in str(caught.value)


def result_for(content):
    return {"version": 1, "sha256": hashlib.sha256(content).hexdigest(), "size_bytes": len(content),
        "source_format": "pdf", "scan": {"engine": "ClamAV", "version": "1.4.3", "definitions": "28000/current"},
        "parsed": {"raw_text": "", "sections": {}, "skills": [], "experience_years": 0,
                   "contact_info": {"name": None, "email": None, "phone": None, "linkedin": None, "github": None}}}


def test_sandbox_exact_import_lifecycle_and_cleanup_before_result(service, monkeypatch):
    content = pdf()
    original = payload(content)
    calls = []
    tar_paths = []
    monkeypatch.setattr(service.os.path, "isfile", lambda _: True)
    monkeypatch.setattr(service.os, "access", lambda *_: True)
    def run(args, **options):
        calls.append(args)
        if args[1] == "do":
            path = Path(next(a.split("=", 1)[1] for a in args if a.startswith("--import-tar=")))
            tar_paths.append(path)
            assert path.stat().st_mode & 0o777 == 0o600
            with tarfile.open(path) as archive:
                assert len(archive.getmembers()) == 1
                entry = archive.getmembers()[0]
                assert entry.name == "tmp/request.json" and entry.isreg() and entry.mode == 0o444
                assert archive.extractfile(entry).read() == original
            return 0, json.dumps(result_for(content)).encode()
        return 0, b""
    monkeypatch.setattr(service, "_run", run)
    result = service.inspect_in_sandbox(original)
    assert [args[1] for args in calls] == ["do", "delete"]
    assert calls[0][2] == "--sandbox-name=" + calls[1][2]
    assert all(args[0] == service.SANDBOX for args in calls)
    assert not any(a.startswith(("--allow-egress", "--env", "--mount", "--export", "--sync")) for args in calls for a in args)
    assert all(not path.exists() for path in tar_paths)
    assert result["worker_image"].endswith("@sha256:" + "a" * 64)
    assert result["policy_sha256"] == "b" * 64


@pytest.mark.parametrize("failure_step", ["do", "delete"])
def test_failed_do_or_delete_never_releases_result(service, monkeypatch, failure_step):
    content = pdf()
    calls = []
    monkeypatch.setattr(service.os.path, "isfile", lambda _: True)
    monkeypatch.setattr(service.os, "access", lambda *_: True)
    def run(args, **_):
        calls.append(args[1])
        if args[1] == failure_step:
            raise service.InspectionDenied()
        return 0, json.dumps(result_for(content)).encode() if args[1] == "do" else b""
    monkeypatch.setattr(service, "_run", run)
    with pytest.raises(service.InspectionDenied):
        service.inspect_in_sandbox(payload(content))
    assert calls[-1] == "delete" and calls.count("delete") == 1


def test_missing_launcher_has_no_host_fallback(service, monkeypatch):
    monkeypatch.setattr(service.os.path, "isfile", lambda _: False)
    monkeypatch.setattr(service, "_run", lambda *_a, **_k: pytest.fail("process reached"))
    with pytest.raises(service.InspectionDenied):
        service.inspect_in_sandbox(payload(pdf()))


@pytest.mark.parametrize("field", ["DOCUMENT_WORKER_IMAGE_DIGEST", "DOCUMENT_WORKER_POLICY_SHA256"])
def test_unpinned_worker_configuration_refused(service, monkeypatch, field):
    monkeypatch.setenv(field, "unreviewed")
    with pytest.raises(service.InspectionDenied):
        service.configuration()


def test_bounded_process_output_and_wall_clock(service):
    # Actual harmless local children test the supervisor, not a Cloud sandbox.
    with pytest.raises(service.InspectionDenied):
        service._run([sys.executable, "-I", "-c", "print('x'*10000)"], timeout=2, limit=100)
    with pytest.raises(service.InspectionDenied):
        service._run([sys.executable, "-I", "-c", "import time; time.sleep(5)"], timeout=1, limit=100)


def test_parser_only_image_has_no_config_clients_or_source_tree():
    text = (ROOT / "infra/document-processing/Dockerfile").read_text()
    copied = "\n".join(line for line in text.splitlines() if line.startswith("COPY"))
    assert "COPY backend/app " not in text and "COPY . " not in text
    assert not any(word in copied for word in ("database", "config", "llm_client", "models", ".env", "auth"))
    assert "freshclam" in text and "USER 65534:65534" in text


@pytest.mark.parametrize("kind", ["good", "wrong-policy", "duplicate-policy", "duplicate-length", "wrong-path", "compressed"])
def test_actual_http_adapter_fixed_boundary(service, monkeypatch, kind, capsys):
    calls = []
    def inspect(raw):
        calls.append(raw)
        return {"version": 1, "fixed": "synthetic-result"}
    monkeypatch.setattr(service, "inspect_in_sandbox", inspect)
    server = service.InspectionHTTPServer(("127.0.0.1", 0), service.Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    original = payload(pdf())
    connection = http.client.HTTPConnection("127.0.0.1", server.server_address[1], timeout=2)
    try:
        connection.putrequest("POST", "/other" if kind == "wrong-path" else "/inspect")
        connection.putheader("Content-Type", "application/json")
        connection.putheader("Content-Length", str(len(original)))
        connection.putheader("X-HireWiz-Document-Policy", "private-marker" if kind == "wrong-policy" else "b" * 64)
        if kind == "duplicate-policy":
            connection.putheader("X-HireWiz-Document-Policy", "b" * 64)
        if kind == "duplicate-length":
            connection.putheader("Content-Length", str(len(original)))
        if kind == "compressed":
            connection.putheader("Content-Encoding", "gzip")
        connection.endheaders(original)
        response = connection.getresponse()
        result = json.loads(response.read())
        assert response.getheader("Cache-Control") == "no-store"
        if kind == "good":
            assert response.status == 200 and calls == [original]
        else:
            assert response.status == 503 and calls == []
            assert result == {"error": "document_inspection_refused"}
    finally:
        connection.close()
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)
    captured = capsys.readouterr()
    assert captured.out == captured.err == ""


@pytest.mark.parametrize("instruction", ['INCLUDETEXT "https://fixture.invalid/candidate.docx"',
    'INCLUDEPICTURE "https://fixture.invalid/image"', 'DDE private-marker',
    'DDEAUTO private-marker', 'DATABASE private-marker', 'LINK private-marker'])
@pytest.mark.parametrize("complex_field", [False, True])
def test_docx_active_field_instructions_refused(processor, instruction, complex_field):
    from html import escape
    if complex_field:
        # Instruction tokens may be split across Word runs.
        first, second = instruction[:3], instruction[3:]
        field = ('<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
                 '<w:r><w:instrText>' + escape(first) + '</w:instrText></w:r>'
                 '<w:r><w:instrText>' + escape(second) + '</w:instrText></w:r>'
                 '<w:r><w:fldChar w:fldCharType="end"/></w:r>')
    else:
        field = '<w:fldSimple w:instr="' + escape(instruction, quote=True) + '"/>'
    xml = ('<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
           '<w:body><w:p>' + field + '</w:p></w:body></w:document>').encode()
    with pytest.raises(processor.InspectionDenied):
        processor.validate_docx(replace_zip(docx(), "word/document.xml", xml))


@pytest.mark.parametrize("complex_field", [False, True])
def test_docx_passive_page_field_remains_supported(processor, complex_field):
    if complex_field:
        field = ('<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
                 '<w:r><w:instrText>NUM</w:instrText></w:r>'
                 '<w:r><w:instrText>PAGES</w:instrText></w:r>'
                 '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
                 '<w:r><w:t>1</w:t></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r>')
    else:
        field = r'<w:fldSimple w:instr=" PAGE \* MERGEFORMAT "/>'
    xml = ('<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
           '<w:body><w:p>' + field + '</w:p></w:body></w:document>').encode()
    processor.validate_docx(replace_zip(docx(), "word/document.xml", xml))


@pytest.mark.parametrize("stage", ["headers", "body"])
def test_http_absolute_read_deadline_closes_dripping_request(service, monkeypatch, stage, capsys):
    monkeypatch.setattr(service, "READ_SECONDS", 0.12, raising=False)
    monkeypatch.setattr(service, "inspect_in_sandbox", lambda _: pytest.fail("inspection reached"))
    server = service.InspectionHTTPServer(("127.0.0.1", 0), service.Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    connection = socket.create_connection(server.server_address, timeout=1)
    try:
        prefix = b"POST /inspect HTTP/1.1\r\n"
        if stage == "body":
            prefix += (b"Content-Type: application/json\r\nContent-Length: 100\r\n"
                       b"X-HireWiz-Document-Policy: " + b"b" * 64 + b"\r\n\r\n")
        else:
            prefix += b"Host: "
        connection.sendall(prefix)
        for _ in range(6):
            time.sleep(0.03)  # Each gap is less than an inactivity timeout.
            try:
                connection.sendall(b"x")
            except OSError:
                break
        connection.settimeout(0.2)
        try:
            response = connection.recv(1024)
        except ConnectionResetError:
            response = b""  # Native TCP reset also proves the read was terminated.
        except TimeoutError:
            pytest.fail("dripping request outlived the absolute read deadline")
        assert response == b"" or b"503" in response
    finally:
        connection.close()
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)
    assert capsys.readouterr() == ("", "")


def test_http_read_timer_cancelled_before_processing_and_on_finish(service, monkeypatch, capsys):
    monkeypatch.setattr(service, "READ_SECONDS", 0.1, raising=False)
    timers = []
    original_timer = threading.Timer
    def timer(*args, **kwargs):
        value = original_timer(*args, **kwargs)
        timers.append(value)
        return value
    monkeypatch.setattr(service.threading, "Timer", timer)
    def inspect(_):
        time.sleep(0.2)  # Inspection may exceed the completed body-read window.
        return {"version": 1, "fixed": "synthetic-result"}
    monkeypatch.setattr(service, "inspect_in_sandbox", inspect)
    server = service.InspectionHTTPServer(("127.0.0.1", 0), service.Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    connection = http.client.HTTPConnection(*server.server_address, timeout=1)
    try:
        connection.request("POST", "/inspect", b"{}", headers={"Content-Type": "application/json",
            "X-HireWiz-Document-Policy": "b" * 64})
        response = connection.getresponse()
        assert response.status == 200
        assert json.loads(response.read()) == {"version": 1, "fixed": "synthetic-result"}
    finally:
        connection.close()
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)
    for value in timers:
        value.join(timeout=1)
    assert len(timers) == 1 and not timers[0].is_alive()
    assert capsys.readouterr() == ("", "")


@pytest.mark.parametrize("expiry", ["timer-fired", "deadline-with-delayed-timer"])
def test_complete_body_cannot_dispatch_after_read_deadline(service, monkeypatch, expiry):
    from email.message import Message
    handler = service.Handler.__new__(service.Handler)
    handler._read_lock = threading.Lock()
    handler._read_active = True
    handler._read_expired = False
    handler._read_deadline = 2
    cancellations = []
    shutdowns = []
    class Timer:
        def cancel(self):
            cancellations.append(True)
    class Connection:
        def shutdown(self, _):
            shutdowns.append(True)
    handler._read_timer = Timer()
    handler.connection = Connection()
    handler.path = "/inspect"
    handler.headers = Message()
    handler.headers["Content-Type"] = "application/json"
    handler.headers["Content-Length"] = "2"
    handler.headers["X-HireWiz-Document-Policy"] = "b" * 64
    handler.rfile = io.BytesIO(b"{}")
    calls = []
    replies = []
    monkeypatch.setattr(service, "inspect_in_sandbox", lambda raw: calls.append(raw) or {})
    monkeypatch.setattr(handler, "_reply", lambda status, result: replies.append((status, result)))
    original_finish = handler._finish_request_read
    def paused_finish(*args, **kwargs):
        # The complete body has been read, but timer/admission have not locked.
        if expiry == "timer-fired":
            handler._expire_read()
        return original_finish(*args, **kwargs)
    monkeypatch.setattr(handler, "_finish_request_read", paused_finish)
    monkeypatch.setattr(service.time, "monotonic", lambda: 2)
    handler.do_POST()
    assert calls == []
    assert replies == [(503, {"error": "document_inspection_refused"})]
    assert cancellations == [True]
    assert len(shutdowns) == (1 if expiry == "timer-fired" else 0)
    assert not service.LOCK.locked()


@pytest.mark.parametrize("now,allowed", [(1.999, True), (2.0, False), (2.001, False)])
def test_read_admission_boundary_and_cleanup_are_idempotent(service, monkeypatch, now, allowed):
    handler = service.Handler.__new__(service.Handler)
    handler._read_lock = threading.Lock()
    handler._read_active = True
    handler._read_expired = False
    handler._read_deadline = 2.0
    cancellations = []
    class Timer:
        def cancel(self):
            cancellations.append(True)
    handler._read_timer = Timer()
    monkeypatch.setattr(service.time, "monotonic", lambda: now)
    if allowed:
        handler._finish_request_read(admit=True)
    else:
        with pytest.raises(service.InspectionDenied):
            handler._finish_request_read(admit=True)
    handler._finish_request_read()
    handler._finish_request_read()
    assert cancellations == [True, True, True] and not handler._read_active


@pytest.mark.parametrize(
    "failure", [FileNotFoundError("private-marker"), subprocess.TimeoutExpired("private-marker", 1)]
)
def test_scanner_infrastructure_has_distinct_private_unavailable_code(
    processor, monkeypatch, failure
):
    monkeypatch.setattr(processor.subprocess, "run", lambda *a, **k: (_ for _ in ()).throw(failure))
    with pytest.raises(ValueError) as caught:
        processor._command(["/usr/bin/clamscan"], 1)
    assert str(caught.value) == "document_inspection_unavailable"
    assert caught.value.__cause__ is None


@pytest.mark.parametrize("age", [timedelta(days=3, seconds=1), timedelta(seconds=-1)])
def test_definition_freshness_is_unavailable_not_document_refusal(
    processor, monkeypatch, tmp_path, age
):
    monkeypatch.setattr(
        processor.subprocess,
        "run",
        lambda *a, **k: subprocess.CompletedProcess(
            a[0], 0, b"ClamAV 1.4.3/28000/Sat Oct 10 00:00:00 2026\n", b""
        ),
    )
    with pytest.raises(ValueError) as caught:
        processor.scan_original(
            tmp_path / "original.pdf", now=datetime(2026, 10, 10, tzinfo=UTC) + age
        )
    assert str(caught.value) == "document_inspection_unavailable"


@pytest.mark.parametrize("permanent,exit_code", [(True, 1), (False, 2)])
def test_processor_main_outcome_is_fixed_versioned_and_exact_byte_bound(
    processor, monkeypatch, capsys, permanent, exit_code
):
    import builtins

    original = pdf()
    monkeypatch.setattr(processor.sys, "argv", ["processor.py", "/tmp/request.json"])
    monkeypatch.setattr(processor.sys, "path", list(processor.sys.path))
    monkeypatch.setattr(processor.resource, "setrlimit", lambda *a: None)
    monkeypatch.setattr(processor.os, "umask", lambda *a: None)
    monkeypatch.setattr(processor.logging, "disable", lambda *a: None)
    monkeypatch.setattr(processor.warnings, "simplefilter", lambda *a: None)
    monkeypatch.setattr(builtins, "open", lambda *a, **k: io.BytesIO(payload(original)))
    failure = (
        processor.InspectionDenied() if permanent else RuntimeError("private-candidate-marker")
    )
    monkeypatch.setattr(processor, "inspect", lambda raw: (_ for _ in ()).throw(failure))
    with pytest.raises(SystemExit) as caught:
        processor.main()
    assert caught.value.code == exit_code
    observed = json.loads(capsys.readouterr().out)
    assert observed == {
        "version": 1,
        "sha256": hashlib.sha256(original).hexdigest(),
        "size_bytes": len(original),
        "source_format": "pdf",
        "error": "document_inspection_refused" if permanent else "document_inspection_unavailable",
    }


@pytest.mark.parametrize(
    "error,status", [("document_inspection_refused", 422), ("document_inspection_unavailable", 503)]
)
def test_http_controlled_outcome_retains_bound_fixed_protocol(service, monkeypatch, error, status):
    source = pdf()
    result = {
        "version": 1,
        "sha256": hashlib.sha256(source).hexdigest(),
        "size_bytes": len(source),
        "source_format": "pdf",
        "worker_image": "registry.invalid/document@sha256:" + "a" * 64,
        "policy_sha256": "b" * 64,
        "error": error,
    }
    monkeypatch.setattr(service, "inspect_in_sandbox", lambda raw: result)
    server = service.InspectionHTTPServer(("127.0.0.1", 0), service.Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    connection = http.client.HTTPConnection("127.0.0.1", server.server_port, timeout=5)
    try:
        connection.request(
            "POST",
            "/inspect",
            body=payload(source),
            headers={"Content-Type": "application/json", "X-HireWiz-Document-Policy": "b" * 64},
        )
        response = connection.getresponse()
        assert response.status == status and json.loads(response.read()) == result
    finally:
        connection.close()
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)


@pytest.mark.parametrize(
    "signature",
    ["Win.Test.EICAR_HDB-1", "Heuristics.Encrypted.PDF", "Heuristics.Limits.Exceeded.MaxScanSize"],
)
def test_coherent_one_file_detection_is_permanent_and_precedes_parsing(
    processor, monkeypatch, tmp_path, signature
):
    path = tmp_path / "original.pdf"
    output = scanner_output(path, **{"Infected files": "1"}).replace(
        str(path) + ": OK", str(path) + ": " + signature + " FOUND"
    )

    def command(args, **options):
        if "--version" in args:
            return subprocess.CompletedProcess(
                args, 0, b"ClamAV 1.4.3/28000/Sat Oct 10 00:00:00 2026\n", b""
            )
        return subprocess.CompletedProcess(args, 1, output.encode(), b"")

    monkeypatch.setattr(processor.subprocess, "run", command)
    with pytest.raises(processor.InspectionDenied, match="^document_inspection_refused$"):
        processor.scan_original(path, now=datetime(2026, 10, 10, tzinfo=UTC))


@pytest.mark.parametrize(
    "code,variant",
    [
        (2, "found"),
        (0, "found"),
        (1, "clean"),
        (1, "count"),
        (1, "foreign"),
        (1, "error"),
        (1, "warning"),
        (1, "missing"),
        (1, "duplicate"),
        (1, "stderr"),
        (1, "oversize"),
        (3, "found"),
    ],
)
def test_ambiguous_scanner_outcome_is_unavailable(processor, monkeypatch, tmp_path, code, variant):
    path = tmp_path / "original.pdf"
    output = scanner_output(path, **{"Infected files": "1"}).replace(
        str(path) + ": OK", str(path) + ": Win.Test.EICAR_HDB-1 FOUND"
    )
    if variant == "clean":
        output = scanner_output(path)
    if variant == "count":
        output = output.replace("Infected files: 1", "Infected files: 0")
    if variant == "foreign":
        output = output.replace(str(path), str(tmp_path / "different.pdf"))
    if variant == "error":
        output += "\nTotal errors: 1"
    if variant == "warning":
        output += "\nWARNING: private-marker"
    if variant == "missing":
        output = output.replace("Scanned files: 1", "")
    if variant == "duplicate":
        output += "\nEngine version: 1.4.3"
    if variant == "oversize":
        output += "x" * 65536

    def command(args, **options):
        if "--version" in args:
            return subprocess.CompletedProcess(
                args, 0, b"ClamAV 1.4.3/28000/Sat Oct 10 00:00:00 2026\n", b""
            )
        return subprocess.CompletedProcess(
            args, code, output.encode(), b"private-marker" if variant == "stderr" else b""
        )

    monkeypatch.setattr(processor.subprocess, "run", command)
    with pytest.raises(processor.InspectionUnavailable) as caught:
        processor.scan_original(path, now=datetime(2026, 10, 10, tzinfo=UTC))
    assert str(caught.value) == "document_inspection_unavailable" and caught.value.__cause__ is None


@pytest.mark.parametrize("kind", ["pdf", "docx"])
def test_unexpected_parser_failure_is_unavailable_not_permanent(processor, monkeypatch, kind):
    from app.services import parsing

    monkeypatch.setattr(
        processor,
        "scan_original",
        lambda _: {"engine": "ClamAV", "version": "synthetic", "definitions": "synthetic"},
    )
    monkeypatch.setattr(
        parsing,
        "parse_resume_file",
        lambda *a, **k: (_ for _ in ()).throw(RuntimeError("private-marker")),
    )
    with pytest.raises(processor.InspectionUnavailable) as caught:
        processor.inspect(payload(pdf() if kind == "pdf" else docx(), kind))
    assert caught.value.__cause__ is None and "private-marker" not in str(caught.value)


def private_outcome(content, code):
    return {
        "version": 1,
        "sha256": hashlib.sha256(content).hexdigest(),
        "size_bytes": len(content),
        "source_format": "pdf",
        "error": "document_inspection_refused" if code == 1 else "document_inspection_unavailable",
    }


@pytest.mark.parametrize("code", [1, 2])
def test_controlled_exec_outcome_requires_successful_cleanup_before_release(
    service, monkeypatch, code
):
    content = pdf()
    calls = []
    monkeypatch.setattr(service.os.path, "isfile", lambda _: True)
    monkeypatch.setattr(service.os, "access", lambda *a: True)

    def run(args, **options):
        calls.append(args[1])
        assert options.get("allow_outcome", False) is (args[1] == "do")
        return (
            (code, json.dumps(private_outcome(content, code)).encode())
            if args[1] == "do"
            else (0, b"")
        )

    monkeypatch.setattr(service, "_run", run)
    result = service.inspect_in_sandbox(payload(content))
    assert calls == ["do", "delete"]
    assert result == {
        **private_outcome(content, code),
        "worker_image": "registry.invalid/document@sha256:" + "a" * 64,
        "policy_sha256": "b" * 64,
    }


@pytest.mark.parametrize(
    "code,field,value",
    [
        (0, None, None),
        (1, "error", "document_inspection_unavailable"),
        (2, "error", "document_inspection_refused"),
        (1, "version", True),
        (1, "sha256", "f" * 64),
        (1, "size_bytes", 1),
        (1, "source_format", "docx"),
        (1, "trace", "private-marker"),
    ],
)
def test_exec_error_protocol_disagreement_is_unavailable_and_always_deleted(
    service, monkeypatch, code, field, value
):
    content = pdf()
    calls = []
    outcome = private_outcome(content, code)
    if field is not None:
        outcome[field] = value
    monkeypatch.setattr(service.os.path, "isfile", lambda _: True)
    monkeypatch.setattr(service.os, "access", lambda *a: True)

    def run(args, **options):
        calls.append(args[1])
        return (code, json.dumps(outcome).encode()) if args[1] == "do" else (0, b"")

    monkeypatch.setattr(service, "_run", run)
    with pytest.raises(service.InspectionDenied) as caught:
        service.inspect_in_sandbox(payload(content))
    assert "private-marker" not in str(caught.value) and calls == ["do", "delete"]


@pytest.mark.parametrize("code", [1, 2])
def test_failed_delete_overrides_controlled_document_outcome(service, monkeypatch, code):
    content = pdf()
    calls = []
    monkeypatch.setattr(service.os.path, "isfile", lambda _: True)
    monkeypatch.setattr(service.os, "access", lambda *a: True)

    def run(args, **options):
        calls.append(args[1])
        if args[1] == "delete":
            raise service.InspectionDenied()
        return (
            (code, json.dumps(private_outcome(content, code)).encode())
            if args[1] == "do"
            else (0, b"")
        )

    monkeypatch.setattr(service, "_run", run)
    with pytest.raises(service.InspectionDenied):
        service.inspect_in_sandbox(payload(content))
    assert calls == ["do", "delete"]


@pytest.mark.parametrize("code,allowed", [(1, False), (2, False), (1, True), (2, True), (3, True)])
def test_only_explicit_exec_mode_consumes_controlled_nonzero_supervisor_result(
    service, code, allowed
):
    # Actual harmless local children exercise pipe/exit handling, not Cloud Run.
    args = [sys.executable, "-I", "-c", "print('fixed');raise SystemExit(" + str(code) + ")"]
    if allowed and code in {1, 2}:
        assert service._run(args, timeout=2, limit=100, allow_outcome=allowed) == (code, b"fixed\n")
    else:
        with pytest.raises(service.InspectionDenied):
            service._run(args, timeout=2, limit=100, allow_outcome=allowed)


@pytest.mark.parametrize("code", [1, 2])
def test_exec_duplicate_keys_never_release_a_controlled_result(service, monkeypatch, code):
    content = pdf()
    calls = []
    raw = json.dumps(private_outcome(content, code)).encode()[:-1] + b',"version":1}'
    monkeypatch.setattr(service.os.path, "isfile", lambda _: True)
    monkeypatch.setattr(service.os, "access", lambda *a: True)

    def run(args, **options):
        calls.append(args[1])
        return (code, raw) if args[1] == "do" else (0, b"")

    monkeypatch.setattr(service, "_run", run)
    with pytest.raises(service.InspectionDenied):
        service.inspect_in_sandbox(payload(content))
    assert calls == ["do", "delete"]
