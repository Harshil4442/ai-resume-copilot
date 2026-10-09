"""Connected authenticated backend journey; no external HTTP, model or portal send.

Only external HTTP transport and the database connection boundary are replaced.
Signup, JWT auth, parsing, feed normalization, ranking, ownership, packages,
review, approval, payments and ledgers execute their actual implementations.
"""

from __future__ import annotations

import hashlib
import hmac
import io
import json
import os
import socket
from datetime import UTC, datetime, timedelta

import httpx
import pytest
from backend.app import models as core
from backend.app.database import Base, get_db
from backend.app.domains.analysis import tasks as analysis_tasks
from backend.app.domains.dispatch import service as dispatch
from backend.app.domains.dispatch.models import DispatchOutbox
from backend.app.domains.employer import models
from backend.app.domains.employer import tasks as employer_tasks
from backend.app.domains.notifications import service as notifications
from backend.app.services import llm_client
from backend.app.services.generation_budget import GenerationBudget
from docx import Document
from fastapi.testclient import TestClient
from google import genai
from sqlalchemy import create_engine, event
from sqlalchemy.orm import sessionmaker

MEDIA = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
CANDIDATE = "candidate.synthetic@example.com"
ADMIN = "operator.synthetic@example.com"
WEBHOOK_SECRET = "synthetic-journey-webhook-secret-not-live"


def _resume(label):
    document = Document()
    document.add_heading("Synthetic Candidate", 0)
    document.add_paragraph(label)
    document.add_heading("Summary", 1)
    document.add_paragraph("Python software engineer building documented services.")
    document.add_heading("Skills", 1)
    document.add_paragraph("Python, PostgreSQL, Docker")
    document.add_heading("Experience", 1)
    document.add_paragraph(
        "Built a Python service using PostgreSQL; improved documented response time by 20%."
    )
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _body(response, expected=200):
    assert response.status_code == expected, response.text
    return response.json()


def _sign(payload):
    raw = json.dumps(payload, sort_keys=True, separators=(",", ":")).encode()
    signature = hmac.new(WEBHOOK_SECRET.encode(), raw, hashlib.sha256).hexdigest()
    return raw, {
        "content-type": "application/json",
        "x-razorpay-event-id": "evt_synthetic_journey",
        "x-razorpay-signature": signature,
    }


@pytest.fixture
def journey(tmp_path, monkeypatch):
    # conftest intentionally supplies synthetic pricing for other tests. This
    # journey removes it and any keys before application construction.
    for name in (
        "LLM_MODEL_COST_POLICY_JSON",
        "LLM_API_KEY",
        "GOOGLE_API_KEY",
        "GEMINI_API_KEY",
        "EMPLOYER_ARTIFACT_BUCKET",
        "REDIS_URL",
        "WORKER_ALLOWED_TOPICS",
    ):
        monkeypatch.delenv(name, raising=False)
    assert all(
        os.getenv(name) is None
        for name in (
            "LLM_MODEL_COST_POLICY_JSON",
            "LLM_API_KEY",
            "GOOGLE_API_KEY",
            "GEMINI_API_KEY",
        )
    )
    monkeypatch.setenv("APP_ENV", "test")
    monkeypatch.setenv("ANALYSIS_TASKS_MODE", "inline")
    monkeypatch.setenv("OPTIONAL_AI_GENERATION_ENABLED", "false")
    monkeypatch.setenv("EMPLOYER_DISCOVERY_ENABLED", "true")
    monkeypatch.setenv("EMPLOYER_AUTO_SUBMIT_ENABLED", "false")
    monkeypatch.setenv("EMPLOYER_SEARCH_CREDITS_PER_JOB", "2")
    monkeypatch.setenv("EMPLOYER_APPLY_CREDITS_PER_JOB", "5")
    monkeypatch.setenv("ADMIN_EMAILS", ADMIN)
    monkeypatch.setenv("LIFECYCLE_EMAILS_ENABLED", "true")
    monkeypatch.setenv("RESEND_API_KEY", "synthetic-only")
    monkeypatch.setenv("EMAIL_FROM", "HireWiz <noreply@hirewiz.example>")
    for name, value in {
        "RAZORPAY_CHECKOUT_ENABLED": "true",
        "RAZORPAY_ACCOUNT_APPROVED": "true",
        "PAYMENTS_GO_LIVE_REVIEW_COMPLETE": "true",
        "RAZORPAY_MODE": "test",
        "RAZORPAY_KEY_ID": "rzp_test_synthetic_journey",
        "RAZORPAY_KEY_SECRET": "synthetic-order-secret",
        "RAZORPAY_WEBHOOK_SECRET": WEBHOOK_SECRET,
    }.items():
        monkeypatch.setenv(name, value)
    reached = []

    def forbidden(*args, **kwargs):
        reached.append("forbidden-generation")
        pytest.fail("A no-AI cold-cache journey reached a model SDK or writing call")

    for name in (
        "_chat",
        "_chat_with_budget",
        "chat_json",
        "rewrite_bullets",
        "generate_interview_questions",
        "tailor_resume_from_evidence",
        "tailor_resume_mega_llm",
        "extract_skills_llm",
        "extract_jd_skills_llm",
        "analyze_job_match_mega_llm",
        "generate_learning_strategy_llm",
    ):
        monkeypatch.setattr(llm_client, name, forbidden)
    monkeypatch.setattr(genai, "Client", forbidden)
    monkeypatch.setattr(GenerationBudget, "admit", forbidden)

    now = datetime.now(UTC)
    rows = []
    for identity, days, title in (
        (101, 1, "Python Software Engineer"),
        (102, 2, "Python Backend Engineer"),
        (103, 60, "Python Software Engineer"),
        (104, 1, "Product Designer"),
    ):
        rows.append(
            {
                "id": identity,
                "internal_job_id": identity + 1000,
                "title": title,
                "location": {"name": "Bengaluru, India"},
                "absolute_url": f"https://job-boards.greenhouse.io/syntheticjourney/jobs/{identity}",
                "content": "<p>Python engineer develops services with PostgreSQL and Docker.</p>"
                if identity != 104
                else "<p>Product design and visual collaboration.</p>",
                "first_published": (now - timedelta(days=days)).isoformat(),
                "updated_at": now.isoformat(),
            }
        )
    http_calls, email_deliveries, supplier_orders = [], [], []

    def transport(request):
        http_calls.append((request.method, request.url.host, request.url.path))
        if request.url.host == "boards-api.greenhouse.io":
            assert request.method == "GET", "No employer submission may be sent"
            assert "authorization" not in request.headers and "cookie" not in request.headers
            base = "/v1/boards/syntheticjourney/jobs"
            if request.url.path == base:
                return httpx.Response(200, json={"jobs": rows, "meta": {"total": len(rows)}})
            selected = next(
                row for row in rows if str(row["id"]) == request.url.path.removeprefix(base + "/")
            )
            return httpx.Response(
                200,
                json={
                    **selected,
                    "questions": [
                        {
                            "label": "First name",
                            "required": True,
                            "fields": [{"name": "first_name", "type": "input_text"}],
                        },
                        {
                            "label": "Last name",
                            "required": True,
                            "fields": [{"name": "last_name", "type": "input_text"}],
                        },
                        {
                            "label": "Email",
                            "required": True,
                            "fields": [{"name": "email", "type": "input_text"}],
                        },
                        {
                            "label": "Resume",
                            "required": True,
                            "fields": [{"name": "resume", "type": "input_file"}],
                        },
                    ],
                },
            )
        if request.url.host == "api.razorpay.com" and request.url.path == "/v1/orders":
            assert request.method == "POST"
            payload = json.loads(request.content)
            assert payload["amount"] == 64_900 and payload["currency"] == "INR"
            assert payload["partial_payment"] is False
            supplier_orders.append(payload)
            return httpx.Response(
                200,
                json={
                    **payload,
                    "entity": "order",
                    "id": "order_synthetic_journey",
                    "status": "created",
                },
            )
        if request.url.host == "api.resend.com" and request.url.path == "/emails":
            assert request.method == "POST"
            payload = json.loads(request.content)
            assert payload["to"] == [CANDIDATE]
            assert "welcome" in payload["subject"].lower()
            email_deliveries.append(payload)
            return httpx.Response(200, json={"id": "email_synthetic_journey"})
        pytest.fail(
            f"Unexpected external HTTP boundary: {request.method} {request.url.host} {request.url.path}"
        )

    original_init = httpx.Client.__init__

    def client_init(client, *args, **kwargs):
        # Preserve TestClient's ASGI transport; all would-be external clients
        # receive the strict synthetic boundary with no actual sockets.
        if kwargs.get("transport") is None:
            kwargs["transport"] = httpx.MockTransport(transport)
        original_init(client, *args, **kwargs)

    monkeypatch.setattr(httpx.Client, "__init__", client_init)

    def socket_forbidden(*args, **kwargs):
        pytest.fail("An isolated backend journey attempted a real network connection")

    monkeypatch.setattr(socket.socket, "connect", socket_forbidden)
    monkeypatch.setattr(socket, "create_connection", socket_forbidden)

    engine = create_engine(
        "sqlite:///" + str(tmp_path / "journey.db"), connect_args={"check_same_thread": False}
    )

    @event.listens_for(engine, "connect")
    def foreign_keys(connection, _):
        connection.execute("PRAGMA foreign_keys=ON")

    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, autoflush=False)
    for module in (dispatch, analysis_tasks, employer_tasks):
        monkeypatch.setattr(module, "SessionLocal", factory)
    from backend.app.main import app

    def database():
        with factory() as db:
            yield db

    previous = dict(app.dependency_overrides)
    app.dependency_overrides[get_db] = database
    with factory() as db:
        assert db.query(core.User).count() == db.query(core.Resume).count() == 0
        assert db.query(core.AnalysisRun).count() == db.query(models.EmployerPosting).count() == 0
        assert (
            db.query(core.ModelCallEvent).count()
            == db.query(models.ServiceCreditEvent).count()
            == 0
        )
        assert db.query(core.SkillCoverage).count() == 0
    try:
        with TestClient(app) as client:
            yield factory, client, reached, http_calls, email_deliveries, supplier_orders
    finally:
        app.dependency_overrides.clear()
        app.dependency_overrides.update(previous)
        engine.dispose()


def _signup_login(client, email):
    user = _body(
        client.post(
            "/api/auth/register",
            json={
                "email": email,
                "password": "synthetic-strong-password-123",
                "accepted_terms": True,
                "confirmed_age_18": True,
            },
        )
    )
    token = _body(
        client.post(
            "/api/auth/login", json={"email": email, "password": "synthetic-strong-password-123"}
        )
    )
    assert token["user_id"] == user["id"]
    return user, {"Authorization": "Bearer " + token["access_token"]}


def test_fresh_authenticated_no_ai_discovery_review_and_manual_handoff(journey):
    factory, client, reached, calls, emails, supplier_orders = journey
    assert client.get("/api/resume/list").status_code == 401
    user, owner = _signup_login(client, CANDIDATE)
    assert user["ai_credits"] == 50
    with factory() as db:
        assert db.get(core.User, user["id"]).job_service_credits == 0
        summary = notifications.deliver_pending_notifications(db)
        assert summary.sent == 1 and summary.failed == summary.retried == 0
        welcome = db.query(core.NotificationOutbox).filter_by(notification_type="welcome").one()
        assert welcome.status == "sent" and welcome.attempt_count == 1
    assert len(emails) == 1
    _, operator = _signup_login(client, ADMIN)

    profile = _body(
        client.put(
            "/api/auth/profile",
            headers=owner,
            json={
                "full_name": "Synthetic Candidate",
                "target_role": "Python Engineer",
                "location": "Bengaluru, India",
                "preferred_location": "Bengaluru",
                "skills": ["Python", "PostgreSQL", "Docker"],
            },
        )
    )
    assert profile["full_name"] == "Synthetic Candidate"
    original, custom = _resume("Original source resume"), _resume("Candidate-written custom resume")
    parsed = _body(
        client.post(
            "/api/resume/parse", headers=owner, files={"file": ("original.docx", original, MEDIA)}
        )
    )
    alternate = _body(
        client.post(
            "/api/resume/parse", headers=owner, files={"file": ("custom.docx", custom, MEDIA)}
        )
    )
    assert parsed["extraction_mode"] == alternate["extraction_mode"] == "deterministic"
    assert parsed["enrichment_state"] == alternate["enrichment_state"] == "not_requested"
    assert "Python" in parsed["skills"]
    assert (
        client.get(f"/api/resume/{parsed['resume_id']}/source", headers=owner).content == original
    )
    assert (
        client.get(f"/api/resume/{parsed['resume_id']}/source", headers=operator).status_code == 404
    )
    evidence = _body(
        client.post(
            "/api/v1/evidence-items",
            headers=owner,
            json={
                "resume_id": parsed["resume_id"],
                "category": "experience",
                "title": "Documented Python service",
                "evidence_text": "Built a Python service using PostgreSQL; improved documented response time by 20%.",
                "skills": ["Python", "PostgreSQL"],
            },
        ),
        201,
    )
    evidence = _body(
        client.patch(
            "/api/v1/evidence-items/" + evidence["id"],
            headers=owner,
            json={"approval_state": "approved"},
        )
    )
    assert evidence["approval_state"] == "approved"
    assert (
        client.patch(
            "/api/v1/evidence-items/" + evidence["id"],
            headers=operator,
            json={"approval_state": "rejected"},
        ).status_code
        == 404
    )

    source = _body(
        client.post(
            "/api/v1/employer-jobs/sources",
            headers=operator,
            json={
                "employer": "Synthetic Origin Employer",
                "employer_key": "synthetic-origin-employer",
                "platform": "greenhouse",
                "board_token": "syntheticjourney",
                "careers_url": "https://employer.example/careers",
                "allowed_hosts": ["employer.example"],
                "verification_url": "https://employer.example/careers",
                "verification_note": "Synthetic official careers fixture attributes this exact employer tenant; no real ownership claim.",
            },
        ),
        201,
    )
    assert source["application_mode"] == "manual"
    _body(
        client.post(f"/api/v1/employer-jobs/sources/{source['id']}/refresh", headers=operator), 202
    )
    dispatch.dispatch_pending(limit=10, topics={"employer.refresh"})
    with factory() as db:
        assert db.query(models.EmployerPosting).count() == 4
        assert db.get(models.EmployerSource, source["id"]).status == "healthy"
    catalog = _body(client.get("/api/v1/employer-jobs/catalog", headers=owner))
    assert catalog["balance"] == 0 and catalog["auto_submit_enabled"] is False
    assert catalog["coverage"]["worldwide_recall_verified"] is False

    checkout = _body(
        client.post(
            "/api/billing/orders",
            headers=owner,
            json={"sku": "starter_bundle", "billing_country": "IN"},
        )
    )
    assert len(supplier_orders) == 1
    assert _body(client.get("/api/v1/employer-jobs/credits", headers=owner))["balance"] == 0
    payload = {
        "event": "payment.captured",
        "payload": {
            "payment": {
                "entity": {
                    "id": "pay_synthetic_journey",
                    "entity": "payment",
                    "amount": checkout["amount_minor"],
                    "currency": "INR",
                    "status": "captured",
                    "captured": True,
                    "order_id": checkout["provider_order_id"],
                    "method": "upi",
                    "international": False,
                    "fee": 1500,
                    "tax": 229,
                    "notes": {
                        "hirewiz_order_id": checkout["order_id"],
                        "sku": "starter_bundle",
                        "billing_country": "IN",
                    },
                }
            }
        },
    }
    raw, signature = _sign(payload)
    _body(client.post("/api/billing/webhooks/razorpay", content=raw, headers=signature))
    _body(client.post("/api/billing/webhooks/razorpay", content=raw, headers=signature))
    assert _body(client.get("/api/v1/employer-jobs/credits", headers=owner))["balance"] == 100
    with factory() as db:
        assert db.get(core.User, user["id"]).ai_credits == 52

    search_input = {
        "resume_id": parsed["resume_id"],
        "role": "Python engineer",
        "location": "Bengaluru",
        "desired_count": 3,
        "published_within_days": 7,
        "idempotency_key": "journey-cold-search",
    }
    search = _body(
        client.post("/api/v1/employer-jobs/searches", headers=owner, json=search_input), 201
    )
    assert search["desired_count"] == 3 and search["delivered_count"] == 2
    assert (
        search["reserved_credits"] == 6
        and search["charged_credits"] == 4
        and search["refunded_credits"] == 2
    )
    assert {item["posting"]["external_id"] for item in search["items"]} == {"101", "102"}
    assert all(
        item["posting"]["apply_url"].startswith(
            "https://job-boards.greenhouse.io/syntheticjourney/jobs/"
        )
        for item in search["items"]
    )
    repeated = _body(
        client.post("/api/v1/employer-jobs/searches", headers=owner, json=search_input), 201
    )
    assert repeated["id"] == search["id"]
    second = _body(
        client.post(
            "/api/v1/employer-jobs/searches",
            headers=owner,
            json={**search_input, "idempotency_key": "journey-repeat-search"},
        ),
        201,
    )
    assert second["charged_credits"] == 0
    assert _body(client.get("/api/v1/employer-jobs/credits", headers=owner))["balance"] == 96

    posting = search["items"][0]["posting"]
    opportunity = _body(
        client.post(
            "/api/v1/opportunities",
            headers=owner,
            json={
                "title": posting["title"],
                "company": posting["employer"],
                "location": posting["location"],
                "source": "verified_employer",
                "source_url": posting["canonical_url"],
                "job_description": posting["description"],
                "resume_id": parsed["resume_id"],
            },
        ),
        201,
    )

    def analysis(operation, mode):
        queued = _body(
            client.post(
                "/api/v1/analysis-runs",
                headers={**owner, "Idempotency-Key": "journey-" + operation},
                json={
                    "operation": operation,
                    "opportunity_id": opportunity["id"],
                    "input": {"mode": mode, "resume_id": parsed["resume_id"]},
                },
            ),
            202,
        )
        state = _body(client.get("/api/v1/analysis-runs/" + queued["id"], headers=owner))
        assert state["status"] == "succeeded" and state["generation_attempt_count"] == 0
        return _body(
            client.get("/api/v1/analysis-runs/" + queued["id"] + "/result", headers=owner)
        )["result"]

    match = analysis("job_match", "basic")
    interview = analysis("interview_questions", "curated")
    assert match["mode"] == "basic" and match["provenance"] == "local_rules"
    assert len(interview["questions"]) == 8
    assert interview["mode"] == interview["provenance"] == "curated"
    answer = _body(
        client.post(
            "/api/rag/ask",
            headers=owner,
            json={
                "job_match_id": match["match_id"],
                "resume_id": parsed["resume_id"],
                "question": "Show my approved evidence.",
                "mode": "basic",
            },
        )
    )
    assert answer["mode"] == "direct" and answer["sources"] == ["evidence:" + evidence["id"]]
    assert "Documented Python service" in answer["answer"]

    application = _body(
        client.post(
            "/api/v1/employer-jobs/applications",
            headers=owner,
            json={
                "posting_id": posting["id"],
                "resume_id": parsed["resume_id"],
                "resume_choice": "original",
                "idempotency_key": "journey-original-application",
            },
        ),
        201,
    )
    application_id = application["id"]
    assert application["application_mode"] == "manual"
    answers = {"first_name": "Synthetic", "last_name": "Candidate", "email": CANDIDATE}
    application = _body(
        client.put(
            f"/api/v1/employer-jobs/applications/{application_id}/package",
            headers=owner,
            json={
                "resume_id": parsed["resume_id"],
                "resume_choice": "original",
                "answers": answers,
                "consents": {},
            },
        )
    )
    original_digest = application["package_digest"]
    preview = client.get(
        f"/api/v1/employer-jobs/applications/{application_id}/artifact", headers=owner
    )
    assert preview.status_code == 200 and preview.content == original
    assert preview.headers["x-artifact-sha256"] == hashlib.sha256(original).hexdigest()
    assert (
        client.get(
            f"/api/v1/employer-jobs/applications/{application_id}/artifact", headers=operator
        ).status_code
        == 404
    )
    assert (
        client.post(
            f"/api/v1/employer-jobs/applications/{application_id}/execute",
            headers=owner,
            json={"package_digest": original_digest},
        ).status_code
        == 409
    )
    _body(
        client.post(
            f"/api/v1/employer-jobs/applications/{application_id}/approve",
            headers=owner,
            json={
                "package_digest": original_digest,
                "allowed_actions": ["fill", "upload"],
            },
        )
    )
    application = _body(
        client.put(
            f"/api/v1/employer-jobs/applications/{application_id}/package",
            headers=owner,
            json={
                "resume_id": alternate["resume_id"],
                "resume_choice": "custom",
                "answers": answers,
                "consents": {},
            },
        )
    )
    custom_digest = application["package_digest"]
    assert custom_digest != original_digest and application["status"] == "ready"
    with factory() as db:
        assert db.get(models.EmployerApplication, application_id).approved_digest is None
    preview = client.get(
        f"/api/v1/employer-jobs/applications/{application_id}/artifact", headers=owner
    )
    assert preview.status_code == 200 and preview.content == custom
    assert (
        client.post(
            f"/api/v1/employer-jobs/applications/{application_id}/execute",
            headers=owner,
            json={"package_digest": original_digest},
        ).status_code
        == 409
    )
    _body(
        client.post(
            f"/api/v1/employer-jobs/applications/{application_id}/approve",
            headers=owner,
            json={
                "package_digest": custom_digest,
                "allowed_actions": ["fill", "upload"],
            },
        )
    )
    handoff = _body(
        client.post(
            f"/api/v1/employer-jobs/applications/{application_id}/execute",
            headers=owner,
            json={"package_digest": custom_digest},
        )
    )
    assert handoff["status"] == "manual_handoff"
    assert handoff["handoff_url"] == posting["apply_url"]
    assert handoff["reserved_credits"] == handoff["charged_credits"] == 0
    assert handoff["receipt"] is None
    assert _body(client.get("/api/v1/employer-jobs/credits", headers=owner))["balance"] == 96
    with factory() as db:
        assert db.query(core.ModelCallEvent).count() == 0
        runs = db.query(core.AnalysisRun).filter_by(user_id=user["id"]).all()
        assert len(runs) == 3
        assert {run.operation for run in runs} == {
            "job_match",
            "interview_questions",
            "match_question",
        }
        assert all(
            run.generation_attempt_count == 0 and run.model_cost_quote is None for run in runs
        )
        assert all(run.usage_state == "committed" and run.committed_units == 1 for run in runs)
        assert db.get(core.User, user["id"]).ai_credits == 49
        assert (
            sum(item.amount for item in db.query(core.UsageEvent).filter_by(
                user_id=user["id"], source_type="analysis_run"
            ))
            == -3
        )
        grants = db.query(core.UsageEvent).filter_by(
            user_id=user["id"], source_type="payment_order", event_type="grant"
        ).all()
        assert len(grants) == 1 and grants[0].amount == 2
        assert grants[0].source_id == checkout["order_id"]
        assert db.query(core.SkillCoverage).count() == 0
        assert (
            db.query(models.ServiceCreditEvent)
            .filter_by(event_type="grant", user_id=user["id"])
            .count()
            == 1
        )
        assert (
            db.query(models.ServiceCreditReservation).filter_by(operation="job_application").count()
            == 0
        )
        assert db.query(models.EmployerApplicationAttempt).count() == 0
        assert db.query(DispatchOutbox).filter_by(topic="employer.apply").count() == 0
        approvals = (
            db.query(models.EmployerApplicationApproval)
            .filter_by(application_id=application_id)
            .order_by(models.EmployerApplicationApproval.approved_at)
            .all()
        )
        assert (
            len(approvals) == 2
            and approvals[0].revoked_at is not None
            and approvals[1].revoked_at is None
        )
        assert approvals[1].review_snapshot["resume_choice"] == "custom"
        assert approvals[1].review_snapshot["artifact_sha256"] == hashlib.sha256(custom).hexdigest()
        assert db.get(core.User, user["id"]).job_service_credits == 96
        assert db.get(core.Resume, parsed["resume_id"]).source_document == original
    assert reached == []
    assert all(method == "GET" for method, host, _ in calls if host == "boards-api.greenhouse.io")
