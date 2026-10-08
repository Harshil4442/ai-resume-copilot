"""Synthetic schema-0009 quiescence checks; no external requests or credentials."""
from __future__ import annotations

import asyncio
import io
from types import SimpleNamespace

import pytest
from backend.app import models, schemas
from backend.app.database import Base
from backend.app.domains.analysis import schemas as analysis_schemas
from backend.app.domains.analysis import service as analysis_service
from backend.app.domains.analysis import tasks
from backend.app.domains.dispatch import service as dispatch
from backend.app.domains.dispatch.models import DispatchOutbox
from backend.app.routers import jobs, resume, worker
from backend.app.services import llm_client, parsing
from backend.app.services.generation_budget import GenerationBudget
from backend.app.services.generation_gate import GenerationPaused, generation_enabled
from backend.app.services.guardrails import billable_operation
from backend.app.services.rag.chat import ask_match_ai
from docx import Document
from fastapi import HTTPException, UploadFile
from google import genai
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool
from starlette.datastructures import Headers
from starlette.requests import Request


def forbidden(*args, **kwargs):
    pytest.fail("Quiesced generation must not construct a client, admit an attempt, or send a request")


@pytest.fixture
def factory(monkeypatch):
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, autoflush=False)
    with factory() as db:
        db.add(models.User(id=1, email="synthetic@example.com", ai_credits=100, tier="free"))
        db.add(models.Resume(id=1, user_id=1, original_filename="resume.docx", raw_text="Python engineer",
                             skills=["Python"], sections={"experience": "Developed Python services"}))
        db.add(models.Opportunity(id="opp_gate", user_id=1, resume_id=1, title="Software Engineer",
                                  company="Example", job_description="Build Python and Docker services for customers.", job_snapshot={}))
        db.commit()
    monkeypatch.setattr(tasks, "SessionLocal", factory)
    monkeypatch.setattr(dispatch, "SessionLocal", factory)
    monkeypatch.setenv("APP_ENV", "test")
    monkeypatch.delenv("ANALYSIS_TASK_TOKEN", raising=False)
    monkeypatch.delenv("WORKER_ALLOWED_TOPICS", raising=False)
    try:
        yield factory
    finally:
        engine.dispose()


def create(factory, operation="job_match", mode=None, key="gate-synthetic-001"):
    data = {"resume_id": 1, "job_description": "Build Python and Docker services for customers."} if operation == "job_match" else {}
    if mode is not None:
        data["mode"] = mode
    with factory() as db:
        run, _ = analysis_service.create_run(db, user_id=1,
            payload=analysis_schemas.AnalysisRunCreate(operation=operation, opportunity_id="opp_gate", input=data),
            header_idempotency_key=key)
        return run.id


@pytest.mark.parametrize("model", ["gemini-1.5-flash", "gpt-4o-mini"])
def test_provider_gate_precedes_client_and_attempt_admission(monkeypatch, model):
    monkeypatch.setenv("OPTIONAL_AI_GENERATION_ENABLED", "false")
    monkeypatch.setenv("LLM_API_KEY", "synthetic-not-a-credential")
    monkeypatch.setattr(llm_client, "LLM_MODEL", model)
    monkeypatch.setattr(genai, "Client", forbidden)
    monkeypatch.setattr(llm_client.httpx, "post", forbidden)
    monkeypatch.setattr(GenerationBudget, "admit", forbidden)
    with pytest.raises(GenerationPaused):
        llm_client._chat([{"role": "user", "content": "synthetic"}])
    with pytest.raises(GenerationPaused):
        llm_client._chat_with_budget([{"role": "user", "content": "synthetic"}], GenerationBudget())


def test_gate_rechecks_before_gemini_fallback(monkeypatch):
    calls = []
    def send(**kwargs):
        calls.append(kwargs["model"])
        monkeypatch.setenv("OPTIONAL_AI_GENERATION_ENABLED", "false")
        raise RuntimeError("503 synthetic provider outage")
    monkeypatch.setenv("OPTIONAL_AI_GENERATION_ENABLED", "true")
    monkeypatch.setenv("LLM_API_KEY", "synthetic-not-a-credential")
    monkeypatch.setattr(llm_client, "LLM_MODEL", "gemini-1.5-flash")
    monkeypatch.setattr(genai, "Client", lambda **kwargs: SimpleNamespace(models=SimpleNamespace(generate_content=send)))
    with pytest.raises(GenerationPaused):
        llm_client._chat([{"role": "user", "content": "synthetic"}])
    assert calls == ["gemini-1.5-flash"]


@pytest.mark.parametrize("value", ["false", "FALSE", "0", "", "invalid"])
def test_gate_rejects_non_enabled_configuration(monkeypatch, value):
    monkeypatch.setenv("OPTIONAL_AI_GENERATION_ENABLED", value)
    assert generation_enabled() is False


@pytest.mark.parametrize("operation,mode", [("job_match", "enhanced"), ("interview_questions", "enhanced")])
def test_async_admission_creates_no_run_outbox_or_unit_hold(factory, monkeypatch, operation, mode):
    monkeypatch.setenv("OPTIONAL_AI_GENERATION_ENABLED", "false")
    with pytest.raises(HTTPException) as caught:
        create(factory, operation, mode)
    assert caught.value.status_code == 503
    with factory() as db:
        assert db.query(models.AnalysisRun).count() == db.query(DispatchOutbox).count() == db.query(models.UsageEvent).count() == 0
        assert db.get(models.User, 1).ai_credits == 100


@pytest.mark.parametrize("operation,payload", [
    ("rewrite_bullets", {}), ("interview_questions_legacy", {}), ("resume_tailor_legacy", {}),
    ("job_match_legacy", {"mode": "enhanced"}), ("match_question", {"mode": "enhanced"}),
    ("learning_strategy", {"mode": "enhanced"}), ("resume_enrichment", {"mode": "enhanced"}),
])
def test_legacy_admission_rejects_before_product_reservation(factory, monkeypatch, operation, payload):
    monkeypatch.setenv("OPTIONAL_AI_GENERATION_ENABLED", "false")
    with factory() as db:
        with pytest.raises(HTTPException) as caught:
            with billable_operation(user_id=1, db=db, operation=operation, amount=1, input_payload=payload):
                forbidden()
        assert caught.value.status_code == 503
        assert db.query(models.AnalysisRun).count() == db.query(models.UsageEvent).count() == 0
        assert db.get(models.User, 1).ai_credits == 100


@pytest.mark.parametrize("through_dispatch", [False, True])
def test_paused_worker_returns_503_before_any_run_or_dispatch_claim(factory, monkeypatch, through_dispatch):
    monkeypatch.setenv("OPTIONAL_AI_GENERATION_ENABLED", "true")
    run_id = create(factory, mode="enhanced")
    with factory() as db:
        event_id = db.query(DispatchOutbox).one().id
    monkeypatch.setenv("OPTIONAL_AI_GENERATION_ENABLED", "false")
    monkeypatch.setattr(llm_client, "_chat_with_budget", forbidden)
    with pytest.raises(HTTPException) as caught:
        if through_dispatch:
            worker.execute_dispatch_task(event_id, x_cloudtasks_taskname="synthetic", x_hirewiz_task_token=None)
        else:
            worker.execute_analysis_task(run_id, x_cloudtasks_taskname="synthetic", x_hirewiz_task_token=None)
    assert caught.value.status_code == 503
    with factory() as db:
        run, event = db.get(models.AnalysisRun, run_id), db.get(DispatchOutbox, event_id)
        assert run.status == "queued" and run.attempt_count == run.generation_attempt_count == 0
        assert event.status == "pending" and event.execution_attempts == 0 and event.lease_token is None
        assert run.usage_state == "reserved" and db.get(models.User, 1).ai_credits == 99
        assert db.query(models.ModelCallEvent).count() == 0


def test_basic_and_curated_work_still_execute_with_gate_closed(factory, monkeypatch):
    monkeypatch.setenv("OPTIONAL_AI_GENERATION_ENABLED", "false")
    monkeypatch.setattr(llm_client, "_chat_with_budget", forbidden)
    match = create(factory)
    interview = create(factory, "interview_questions", key="gate-curated-001")
    assert tasks.process_analysis_run(match) == "succeeded"
    with factory() as db:
        event_id = db.query(DispatchOutbox).filter_by(aggregate_id=interview).one().id
    assert dispatch.process_event(event_id) == "succeeded"
    with factory() as db:
        assert db.get(models.AnalysisRun, match).result_payload["mode"] == "basic"
        assert len(db.get(models.AnalysisRun, interview).result_payload["questions"]) == 8
        assert db.query(models.ModelCallEvent).count() == 0


def test_actual_legacy_job_route_preserves_basic_and_blocks_enhanced(factory, monkeypatch):
    monkeypatch.setenv("OPTIONAL_AI_GENERATION_ENABLED", "false")
    monkeypatch.setattr(llm_client, "_chat_with_budget", forbidden)
    with factory() as db:
        user = db.get(models.User, 1)
        result = jobs.match_job(schemas.JobMatchRequest(resume_id=1, job_title="Engineer",
            job_description="Build Python and Docker services for customers.", mode="basic"), db=db, current_user=user)
        assert result.mode == "basic"
        before = db.query(models.AnalysisRun).count()
        with pytest.raises(HTTPException) as caught:
            jobs.match_job(schemas.JobMatchRequest(resume_id=1, job_title="Engineer",
                job_description="Build Python and Docker services for customers.", mode="enhanced"), db=db, current_user=user)
        assert caught.value.status_code == 503
        assert db.query(models.AnalysisRun).count() == before


def test_direct_match_answer_and_plain_parser_do_not_need_generation(monkeypatch):
    monkeypatch.setenv("OPTIONAL_AI_GENERATION_ENABLED", "false")
    monkeypatch.setattr(llm_client, "_chat_with_budget", forbidden)
    answer = ask_match_ai(resume=SimpleNamespace(id=1, skills=["Python"], sections={}),
        match=SimpleNamespace(id=1, true_gaps=["Docker"], required_skills=["Python", "Docker"], partial_matches=[],
            full_matches=["Python"], match_score=50, fit_summary="", improvement_tips=[]),
        question="Which skills am I missing?", recent_messages=[])
    assert answer.mode == "direct"
    document = Document()
    document.add_paragraph("Synthetic Candidate")
    document.add_heading("Technical Skills", level=2)
    document.add_paragraph("Python and Docker")
    output = io.BytesIO()
    document.save(output)
    assert "Python" in parsing.parse_resume_file(output.getvalue(), "synthetic.docx")[2]


def test_requested_enrichment_keeps_parsed_resume_without_hold(factory, monkeypatch):
    monkeypatch.setenv("OPTIONAL_AI_GENERATION_ENABLED", "false")
    monkeypatch.setattr(llm_client, "_chat_with_budget", forbidden)
    document = Document()
    document.add_paragraph("Synthetic Candidate")
    document.add_heading("Technical Skills", level=2)
    document.add_paragraph("Python and Docker")
    output = io.BytesIO()
    document.save(output)
    upload = UploadFile(io.BytesIO(output.getvalue()), filename="synthetic.docx",
        headers=Headers({"content-type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"}))
    request = Request({"type": "http", "method": "POST", "path": "/api/resume/parse", "headers": [],
                       "client": ("127.0.0.1", 1), "app": SimpleNamespace(state=SimpleNamespace())})
    with factory() as db:
        result = asyncio.run(resume.parse_resume.__wrapped__(request=request, file=upload, enrich_skills=True,
            db=db, current_user=db.get(models.User, 1)))
        assert result.extraction_mode == "deterministic" and result.enrichment_units == 0
        assert result.source_available and result.warnings
        assert db.query(models.AnalysisRun).count() == db.query(models.UsageEvent).count() == 0


def test_paused_admission_replays_exact_and_semantically_equivalent_requests(factory, monkeypatch):
    monkeypatch.setenv("OPTIONAL_AI_GENERATION_ENABLED", "true")
    run_id = create(factory, mode="enhanced")
    monkeypatch.setenv("OPTIONAL_AI_GENERATION_ENABLED", "false")
    assert create(factory, mode="enhanced") == run_id
    assert create(factory, mode="enhanced", key="equivalent-replay-001") == run_id
    with factory() as db:
        changed = analysis_schemas.AnalysisRunCreate(operation="job_match", opportunity_id="opp_gate",
            input={"resume_id": 1, "mode": "enhanced", "job_description": "A different Python position with different requirements."})
        with pytest.raises(HTTPException) as conflict:
            analysis_service.create_run(db, user_id=1, payload=changed, header_idempotency_key="gate-synthetic-001")
        assert conflict.value.status_code == 409
        with pytest.raises(HTTPException) as paused:
            analysis_service.create_run(db, user_id=1, payload=changed, header_idempotency_key="new-generation-001")
        assert paused.value.status_code == 503
        assert db.query(models.AnalysisRun).count() == db.query(DispatchOutbox).count() == 1
        assert db.query(models.UsageEvent).count() == 1
        assert db.get(models.User, 1).ai_credits == 99


def test_employer_dispatch_still_claims_and_completes_with_gate_closed(factory, monkeypatch):
    monkeypatch.setenv("OPTIONAL_AI_GENERATION_ENABLED", "false")
    monkeypatch.setenv("WORKER_ALLOWED_TOPICS", "employer.search")
    calls = []
    monkeypatch.setattr(dispatch, "_execute", lambda topic, aggregate_id: calls.append((topic, aggregate_id)) or "succeeded")
    with factory() as db:
        event = dispatch.enqueue(db, topic="employer.search", aggregate_id="search_synthetic",
            payload={"search_id": "search_synthetic"}, key="synthetic-employer-dispatch-001")
        db.commit()
        event_id = event.id
    assert worker.execute_dispatch_task(event_id, x_cloudtasks_taskname="synthetic", x_hirewiz_task_token=None)["status"] == "succeeded"
    assert calls == [("employer.search", "search_synthetic")]
    with factory() as db:
        event = db.get(DispatchOutbox, event_id)
        assert event.status == "completed" and event.execution_attempts == 1
        assert event.lease_token is None and event.completed_at is not None


def test_employer_worker_scope_denies_analysis_before_paused_gate(factory, monkeypatch):
    monkeypatch.setenv("OPTIONAL_AI_GENERATION_ENABLED", "true")
    run_id = create(factory, mode="enhanced")
    with factory() as db:
        event_id = db.query(DispatchOutbox).one().id
    monkeypatch.setenv("OPTIONAL_AI_GENERATION_ENABLED", "false")
    monkeypatch.setenv("WORKER_ALLOWED_TOPICS", "employer.search")
    for handler, identifier in [(worker.execute_dispatch_task, event_id), (worker.execute_analysis_task, run_id)]:
        with pytest.raises(HTTPException) as denied:
            handler(identifier, x_cloudtasks_taskname="synthetic", x_hirewiz_task_token=None)
        assert denied.value.status_code == 403
    with factory() as db:
        assert db.get(DispatchOutbox, event_id).execution_attempts == 0
        assert db.get(models.AnalysisRun, run_id).attempt_count == 0


def test_terminal_generation_replay_and_duplicate_tasks_remain_idempotent(factory, monkeypatch):
    monkeypatch.setenv("OPTIONAL_AI_GENERATION_ENABLED", "true")
    run_id = create(factory, mode="enhanced")
    monkeypatch.setattr(tasks, "execute_operation", lambda db, run: {"mode": "enhanced", "synthetic": True})
    with factory() as db:
        event_id = db.query(DispatchOutbox).one().id
    assert dispatch.process_event(event_id) == "succeeded"
    monkeypatch.setenv("OPTIONAL_AI_GENERATION_ENABLED", "false")
    monkeypatch.setattr(tasks, "execute_operation", forbidden)
    assert create(factory, mode="enhanced") == run_id
    assert tasks.process_analysis_run(run_id) == "succeeded"
    assert dispatch.process_event(event_id) == "completed"
    with factory() as db:
        run = db.get(models.AnalysisRun, run_id)
        assert run.attempt_count == 1 and run.committed_units == 1 and run.usage_state == "committed"
        assert db.get(DispatchOutbox, event_id).execution_attempts == 1
        assert db.query(models.UsageEvent).count() == 2
        assert db.get(models.User, 1).ai_credits == 99
