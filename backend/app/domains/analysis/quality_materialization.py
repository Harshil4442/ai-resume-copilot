"""Offline native sources for frozen synthetic cases; never quality receipts.

Only this packet's declarative recipes are supported. No source text is executed,
no font substitution is permitted, and no renderer/model/network is invoked.
"""
from __future__ import annotations

import io
import re
import shutil
from datetime import datetime
from importlib.metadata import version
from pathlib import Path
from xml.sax.saxutils import escape
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from . import quality_packet as packet

VERSION = "unified-quality-materialization-v1"
EPOCH = (2000, 1, 1, 0, 0, 0)


class UnsupportedRecipe(ValueError):
    pass


def _check(case: dict) -> None:
    recipe = case["source_recipe"]
    if (case["source_format"] not in packet.FORMATS or recipe["format"] != case["source_format"]
            or recipe["columns"] not in (1, 2)
            or recipe["sections_in_order"] != [s["heading"] for s in case["resume"]["sections"]]
            or not set(recipe["table_sections"]) <= set(recipe["sections_in_order"])
            or not 0.5 <= recipe["margin_inches"] <= 1.2
            or not re.fullmatch(r"synth-[0-9]{3}", case["case_id"])):
        raise UnsupportedRecipe("Recipe is outside the explicitly implemented layout contract")
    actual = "\n\n".join(s["heading"] + "\n" + s["text"] for s in case["resume"]["sections"])
    if actual != case["resume"]["raw_text"]:
        raise UnsupportedRecipe("Section content differs from the frozen complete resume text")
    engine = recipe.get("engine")
    if engine is not None and (
        case["source_format"] not in {"tex", "texzip"}
        or engine != "pdftex-no-shell-escape-v1"
        or recipe["font"] not in {"Latin Modern Roman", "Latin Modern Mono"}
    ):
        raise UnsupportedRecipe("TeX engine or font is outside the supported native policy")


def _zip(members: dict[str, bytes]) -> bytes:
    out = io.BytesIO()
    with ZipFile(out, "w", compression=ZIP_DEFLATED) as archive:
        for name, content in sorted(members.items()):
            info = ZipInfo(name, date_time=EPOCH)
            info.compress_type = ZIP_DEFLATED
            info.external_attr = 0o100644 << 16
            archive.writestr(info, content)
    return out.getvalue()


def _docx(case: dict) -> bytes:
    from docx import Document
    from docx.enum.section import WD_SECTION_START
    from docx.enum.text import WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    recipe, resume = case["source_recipe"], case["resume"]
    doc = Document()
    props = doc.core_properties
    props.author, props.last_modified_by = "HireWiz synthetic fixture", "HireWiz synthetic fixture"
    props.created = props.modified = datetime(*EPOCH)
    props.title = resume["name"]
    for name, size in (("Normal", 10), ("Title", 16), ("Heading 1", 12)):
        style = doc.styles[name]
        style.font.name, style.font.size = recipe["font"], Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        # Theme attributes override an otherwise explicit font/color in some
        # Word/LibreOffice templates. Keep the recipe authoritative.
        rpr = style.element.get_or_add_rPr()
        for element in (rpr.rFonts, rpr.color):
            for key in list(element.attrib):
                if "theme" in key.lower():
                    del element.attrib[key]
        ppr = style.element.get_or_add_pPr()
        border = ppr.find(qn("w:pBdr"))
        if border is not None:
            ppr.remove(border)
        style.paragraph_format.space_after = Pt(6)
    for section in doc.sections:
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.top_margin = section.bottom_margin = Inches(recipe["margin_inches"])
        section.left_margin = section.right_margin = Inches(recipe["margin_inches"])
    doc.add_paragraph(resume["name"], "Title")
    doc.add_paragraph(resume["email"])
    for url in recipe["required_links"]:
        paragraph = doc.add_paragraph()
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), paragraph.part.relate_to(url,
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True))
        run, text = OxmlElement("w:r"), OxmlElement("w:t")
        text.text = url
        run.append(text)
        link.append(run)
        paragraph._p.append(link)
    if recipe["columns"] == 2:
        section = doc.add_section(WD_SECTION_START.CONTINUOUS)
        cols = section._sectPr.find(qn("w:cols"))
        cols.set(qn("w:num"), "2")
        cols.set(qn("w:space"), "360")
    for section in resume["sections"]:
        if recipe["columns"] == 2 and section["heading"] == "Projects":
            doc.add_paragraph().add_run().add_break(WD_BREAK.COLUMN)
        doc.add_paragraph(section["heading"], "Heading 1")
        lines = section["text"].split("\n")
        if section["heading"] in recipe["table_sections"]:
            table = doc.add_table(rows=0, cols=1)
            table.style = "Table Grid"
            table.autofit = False
            table.columns[0].width = Inches((8.5 - 2 * recipe["margin_inches"] - (0.25 if recipe["columns"] == 2 else 0)) / recipe["columns"])
            for line in lines:
                cell = table.add_row().cells[0]
                cell.text = line
        else:
            for line in lines:
                doc.add_paragraph(line)
    out = io.BytesIO()
    doc.save(out)
    with ZipFile(io.BytesIO(out.getvalue())) as archive:
        return _zip({name: archive.read(name) for name in archive.namelist()})


def _tex_escape(text: str) -> str:
    chars = {"\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "$": r"\$", "#": r"\#",
             "_": r"\_", "{": r"\{", "}": r"\}", "~": r"\textasciitilde{}", "^": r"\textasciicircum{}"}
    return "".join(chars.get(char, char) for char in text)


def _tex(case: dict) -> bytes:
    recipe, resume = case["source_recipe"], case["resume"]
    engine = recipe.get("engine")
    if engine == "pdftex-no-shell-escape-v1":
        if recipe["font"] not in {"Latin Modern Roman", "Latin Modern Mono"}:
            raise UnsupportedRecipe("The selected pdfLaTeX font is unsupported")
        font_setup = [r"\usepackage[T1]{fontenc}", r"\usepackage{lmodern}"]
        if recipe["font"] == "Latin Modern Mono":
            font_setup.append(r"\renewcommand{\familydefault}{\ttdefault}")
    elif engine is None:
        font_setup = [r"\usepackage{fontspec}", r"\setmainfont{" + _tex_escape(recipe["font"]) + "}"]
    else:
        raise UnsupportedRecipe("The selected TeX engine is unsupported")
    lines = [r"\documentclass[10pt,letterpaper]{article}", *font_setup,
             r"\usepackage[margin=" + str(recipe["margin_inches"]) + r"in]{geometry}",
             r"\usepackage{multicol,tabularx,hyperref}", r"\hypersetup{hidelinks}",
             r"\setlength{\parindent}{0pt}", r"\setlength{\parskip}{6pt}", r"\setlength{\columnsep}{0.25in}",
             r"\pagestyle{empty}", r"\begin{document}",
             r"{\Large " + _tex_escape(resume["name"]) + r"}\par", _tex_escape(resume["email"]) + r"\par"]
    for url in recipe["required_links"]:
        lines.append(r"\href{" + _tex_escape(url) + "}{" + _tex_escape(url) + r"}\par")
    if recipe["columns"] == 2:
        lines.append(r"\begin{multicols}{2}")
    for section in resume["sections"]:
        if recipe["columns"] == 2 and section["heading"] == "Projects":
            lines.append(r"\columnbreak")
        lines.append(r"\section*{" + _tex_escape(section["heading"]) + "}")
        if section["heading"] in recipe["table_sections"]:
            # These frozen tables have at most three short rows. Never replace
            # an unsupported overflowing table with a different layout.
            lines.append(r"\begin{tabularx}{\linewidth}{|X|}\hline")
            lines.extend(_tex_escape(line) + r" \\ \hline" for line in section["text"].split("\n"))
            lines.append(r"\end{tabularx}\par")
        else:
            lines.extend(_tex_escape(line) + r"\par" for line in section["text"].split("\n"))
    if recipe["columns"] == 2:
        lines.append(r"\end{multicols}")
    lines.append(r"\end{document}")
    return ("\n".join(lines) + "\n").encode()


def _font(recipe: dict, fonts: dict[str, Path]):
    from reportlab.pdfbase.ttfonts import TTFError, TTFont

    supplied = fonts.get(recipe["font"])
    if supplied is None or not supplied.is_file():
        raise UnsupportedRecipe("font_unavailable: exact " + recipe["font"] + " font file is required")
    name = "fixture-" + packet.file_digest(supplied)[:20]
    try:
        font = TTFont(name, str(supplied))
    except TTFError as exc:
        raise UnsupportedRecipe("font_unreadable: a supported exact-family TTF file is required") from exc
    actual = font.face.familyName
    actual = actual.decode("utf-8") if isinstance(actual, bytes) else str(actual)
    if actual.casefold() != recipe["font"].casefold():
        raise UnsupportedRecipe("font_mismatch: the supplied font family differs from the frozen recipe")
    return name, font, supplied


def _pdf(case: dict, fonts: dict[str, Path]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.platypus import (
        BaseDocTemplate,
        Frame,
        FrameBreak,
        LayoutError,
        PageTemplate,
        Paragraph,
        Table,
        TableStyle,
    )

    recipe, resume = case["source_recipe"], case["resume"]
    name, font, _ = _font(recipe, fonts)
    text = resume["name"] + resume["email"] + resume["raw_text"] + "".join(recipe["required_links"])
    if any(ord(char) not in font.face.charToGlyph for char in text if char not in "\n\r\t"):
        raise UnsupportedRecipe("font_glyph_unavailable: source text cannot be represented without substitution")
    pdfmetrics.registerFont(font)
    normal = ParagraphStyle("body", fontName=name, fontSize=10, leading=13, spaceAfter=6)
    title = ParagraphStyle("title", parent=normal, fontSize=16, leading=20, spaceAfter=10)
    heading = ParagraphStyle("heading", parent=normal, fontSize=12, leading=15, spaceBefore=10, keepWithNext=True)
    out = io.BytesIO()
    margin, gap = recipe["margin_inches"] * 72, 18
    width, height = 612 - margin * 2, 792 - margin * 2
    column = (width - gap * (recipe["columns"] - 1)) / recipe["columns"]
    # Header and body share the same columns; source order follows frame flow.
    frames = [Frame(margin + index * (column + gap), margin, column, height,
                    leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0,
                    id=f"column-{index}") for index in range(recipe["columns"])]
    doc = BaseDocTemplate(out, pagesize=(612, 792), invariant=True, title=resume["name"],
                          author="HireWiz synthetic fixture", pageCompression=1)
    doc.addPageTemplates(PageTemplate(id="resume", frames=frames))
    story = [Paragraph(escape(resume["name"]), title), Paragraph(escape(resume["email"]), normal)]
    for url in recipe["required_links"]:
        story.append(Paragraph('<link href="' + escape(url, {'"': '&quot;'}) + '">' + escape(url) + '</link>', normal))
    for section in resume["sections"]:
        if recipe["columns"] == 2 and section["heading"] == "Projects":
            story.append(FrameBreak())
        story.append(Paragraph(escape(section["heading"]), heading))
        paragraphs = [Paragraph(escape(line), normal) for line in section["text"].split("\n")]
        if section["heading"] in recipe["table_sections"]:
            table = Table([[paragraph] for paragraph in paragraphs], colWidths=[column])
            table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D9D9D9")),
                                      ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                      ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
            story.append(table)
        else:
            story.extend(paragraphs)
    try:
        doc.build(story)
    except LayoutError as exc:
        raise UnsupportedRecipe("layout_unavailable: source content cannot fit the requested layout without review") from exc
    return out.getvalue()


def materialize(root: Path, destination: Path, *, fonts: dict[str, Path] | None = None,
                case_ids: list[str] | None = None) -> dict:
    validation = packet.validate_packet(root)
    if not validation.get("draft_inputs_valid"):
        raise ValueError("The frozen packet must validate before materialization")
    manifest, _, cases = packet.load_packet(root)
    supported_v2 = manifest["schema_version"] == packet.SUPPORTED_VERSION
    if supported_v2:
        from .quality_supported_v2 import verify_lineage

        if not verify_lineage(root.parent / "unified_credit_v1", root)["source_lineage_valid"]:
            raise ValueError("The supported packet must match its frozen v1 lineage")
    if case_ids is not None:
        if not case_ids or len(set(case_ids)) != len(case_ids) or set(case_ids) - {c["case_id"] for c in cases}:
            raise ValueError("Select distinct existing frozen case IDs")
        cases = [case for case in cases if case["case_id"] in case_ids]
    if destination.resolve() == root.resolve() or root.resolve() in destination.resolve().parents:
        raise ValueError("Output must remain separate from the frozen packet")
    destination.mkdir(parents=True, exist_ok=False)
    fonts = fonts or {}
    rows = []
    for case in cases:
        _check(case)
        folder = destination / case["case_id"]
        folder.mkdir()
        (folder / "case.json").write_bytes(packet.canonical(case))
        (folder / "resume.txt").write_text(case["resume"]["raw_text"], encoding="utf-8")
        source_format = case["source_format"]
        # The v1 sidecar remains unchanged; v2 has only its selected source.
        if not supported_v2 or source_format in {"tex", "texzip"}:
            (folder / "source.tex").write_bytes(_tex(case))
        row = {"case_id": case["case_id"], "input_sha256": packet.digest(packet.canonical(case)),
               "resume_content_sha256": packet.digest(packet.canonical(case["resume"])),
               "source_format": source_format, "source_recipe": case["source_recipe"],
               "input_artifact": None, "materialization_state": "source_created",
               "rendered_proof": "pending", "visual_review": "pending", "qualification": "not_evaluated",
               "compiler": ("not_applicable" if supported_v2 and source_format in {"pdf", "docx"}
                            else "available_not_executed" if shutil.which("pdflatex" if supported_v2 else "xelatex")
                            else "unavailable"),
               "font": {"family": case["source_recipe"]["font"], "state": "unavailable"},
               "issues": []}
        if supported_v2 and source_format in {"tex", "texzip"}:
            row["font"] = {"family": case["source_recipe"]["font"],
                           "state": "policy_declared_native_unverified"}
        else:
            try:
                _, _, selected = _font(case["source_recipe"], fonts)
                row["font"] = {"family": case["source_recipe"]["font"], "state": "exact_file_verified",
                               "sha256": packet.file_digest(selected)}
            except (UnsupportedRecipe, OSError, ValueError) as exc:
                row["issues"].append(str(exc))
        try:
            if source_format == "pdf":
                (folder / "source.pdf").write_bytes(_pdf(case, fonts))
            elif source_format == "docx":
                (folder / "source.docx").write_bytes(_docx(case))
            elif source_format == "texzip":
                (folder / "source.texzip").write_bytes(_zip({"main.tex": _tex(case)}))
            source = folder / ("source." + source_format)
            row["input_artifact"] = {"path": str(source.relative_to(destination)), "sha256": packet.file_digest(source)}
        except (UnsupportedRecipe, OSError, ValueError) as exc:
            row["materialization_state"] = "native_source_pending"
            row["issues"].append(str(exc))
        row["files"] = [{"path": str(path.relative_to(destination)), "sha256": packet.file_digest(path)}
                        for path in sorted(folder.iterdir())]
        rows.append(row)
    result = {"schema_version": "unified-quality-materialization-v2" if supported_v2 else VERSION,
              "state": "synthetic_native_sources_not_reviewed",
              "packet_manifest_sha256": packet.file_digest(root / "manifest.json"),
              "corpus_sha256": packet.file_digest(root / "cases.jsonl"),
              "tool_sha256": packet.file_digest(Path(__file__)),
              "libraries": {name: version(name) for name in ("python-docx", "reportlab")},
              "case_count": len(rows), "materialized_count": sum(r["input_artifact"] is not None for r in rows),
              "cases": rows, "authorizes_release": False, "quality_status": "not_evaluated"}
    (destination / "materialization-manifest.json").write_bytes(packet.canonical(result))
    return result
