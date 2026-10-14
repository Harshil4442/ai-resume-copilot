"""Synthetic document-routing diagnostic; offline unless --execute is supplied.

Run only in the fixed Cloud Run Job with the replacement runtime identity and
Direct VPC routing. This never deploys, grants IAM, reads customer input, retries
a processing POST, or writes a token/response/document to an artifact or log.
"""
from __future__ import annotations

import argparse
import base64
import hashlib
import http.client
import importlib.util
import io
import json
import os
import signal
import socket
import ssl
import sys
import time
from pathlib import Path
from typing import Any
from urllib.parse import quote, urlsplit

PROJECT = "ai-resume-parser-482412"
PROJECT_NUMBER = "157225724590"
REGION = "us-central1"
JOB = "hirewiz-document-connectivity-20261014"
IDENTITY = f"hirewiz-runtime-20261014@{PROJECT}.iam.gserviceaccount.com"
ORIGIN = "https://hirewiz-document-inspector-mqff62d75a-uc.a.run.app"
IMAGE = (f"{REGION}-docker.pkg.dev/{PROJECT}/cloud-run-source-deploy/hirewiz-document-inspector"
         "@sha256:624597726ce2fefa2b3574fbf0a8ec4ccdc02a826380d76c5edc8ebbd59aa107")
POLICY = "dbd15ce3120bb6211bb6c1d7b166aaddbd8b80e00a72d564e6274ef10db3aaf1"
RENDERER = "LibreOffice 25.2.3.2 520(Build:2)"
FONTS = "abaf429caeeeb5fbcc000c878a315c89e23ecd3dc912bb74df10761c144f7d60"
PRIVATE_IPS = frozenset({"199.36.153.8", "199.36.153.9", "199.36.153.10", "199.36.153.11"})
EXTERNAL_ORIGIN = "https://api.github.com"
_ERROR = "document_native_connectivity_unverified"


class ConnectivityDenied(ValueError):
    """A finite reason; external response bodies/exceptions never escape."""


def _require(value: Any) -> None:
    if not value:
        raise ConnectivityDenied(_ERROR)


def _canonical(value: Any) -> bytes:
    return json.dumps(value, sort_keys=True, separators=(",", ":"),
                      ensure_ascii=True, allow_nan=False).encode()


def _sha(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def _pairs(items: list[tuple[str, Any]]) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for key, value in items:
        _require(key not in result)
        result[key] = value
    return result


def _decode(value: bytes) -> dict[str, Any]:
    _require(type(value) is bytes and 0 < len(value) <= 1024 * 1024)
    result = json.loads(value, object_pairs_hook=_pairs,
                        parse_constant=lambda _: (_ for _ in ()).throw(ConnectivityDenied(_ERROR)))
    _require(type(result) is dict)
    return result


def _expire(*_: Any) -> None:
    raise ConnectivityDenied(_ERROR)


def _request(origin: str, path: str, *, method: str = "GET", payload: bytes | None = None,
             token: str | None = None, seconds: float = 5, limit: int = 32768,
             metadata: bool = False) -> tuple[int, dict[str, str], bytes]:
    """Fixed destinations, no proxy/redirect/default-credential/TLS-env behavior."""
    _require((metadata and origin == "http://metadata.google.internal" and token is None and method == "GET")
             or (not metadata and origin in {ORIGIN, EXTERNAL_ORIGIN}
                 and (token is None or origin == ORIGIN)))
    _require(0 < seconds <= 75 and 0 < limit <= 1024 * 1024)
    parsed = urlsplit(origin)
    headers = {"Accept-Encoding": "identity", "User-Agent": "HireWiz-routing-diagnostic/1"}
    if metadata:
        connection = http.client.HTTPConnection(parsed.hostname, timeout=min(seconds, 5))
        headers["Metadata-Flavor"] = "Google"
    else:
        import certifi
        context = ssl.SSLContext(ssl.PROTOCOL_TLS_CLIENT)
        context.load_verify_locations(cafile=certifi.where())
        connection = http.client.HTTPSConnection(parsed.hostname, timeout=min(seconds, 5), context=context)
    if token is not None:
        _require(type(token) is str and 0 < len(token) <= 16384
                 and not any(ord(character) < 33 or ord(character) > 126 for character in token))
        headers["Authorization"] = "Bearer " + token
    if payload is not None:
        headers.update({"Content-Type": "application/json", "X-HireWiz-Document-Policy": POLICY})
    previous = signal.signal(signal.SIGALRM, _expire)
    signal.setitimer(signal.ITIMER_REAL, seconds)
    try:
        connection.connect()
        _require(connection.sock is not None)
        connection.sock.settimeout(seconds)
        connection.request(method, path, body=payload, headers=headers)
        response = connection.getresponse()
        response_headers: dict[str, str] = {}
        for key, value in response.getheaders():
            lowered = key.lower()
            _require(lowered not in response_headers or lowered not in {
                "content-type", "content-encoding", "content-length", "metadata-flavor"})
            response_headers[lowered] = value
        content = response.read(limit + 1)
        _require(len(content) <= limit and response_headers.get("content-encoding", "identity").lower() == "identity")
        return response.status, response_headers, content
    finally:
        connection.close()
        signal.setitimer(signal.ITIMER_REAL, 0)
        signal.signal(signal.SIGALRM, previous)


def _metadata(path: str, limit: int) -> str:
    code, headers, raw = _request("http://metadata.google.internal", "/computeMetadata/v1/" + path,
                                  metadata=True, limit=limit)
    _require(code == 200 and headers.get("metadata-flavor") == "Google")
    return raw.decode("ascii").strip()


def _identity_token() -> str:
    _require(os.environ.get("CLOUD_RUN_JOB") == JOB and os.environ.get("CLOUD_RUN_TASK_ATTEMPT") == "0")
    _require(_metadata("project/project-id", 128) == PROJECT
             and _metadata("project/numeric-project-id", 128) == PROJECT_NUMBER
             and _metadata("instance/service-accounts/default/email", 256) == IDENTITY)
    return _metadata("instance/service-accounts/default/identity?audience=" + quote(ORIGIN, safe="")
                     + "&format=full", 16384)


def _private_dns(*, seconds: float = 5) -> None:
    _require(0 < seconds <= 5)
    previous = signal.signal(signal.SIGALRM, _expire)
    signal.setitimer(signal.ITIMER_REAL, seconds)
    try:
        rows = socket.getaddrinfo(urlsplit(ORIGIN).hostname, 443, socket.AF_UNSPEC, socket.SOCK_STREAM)
        _require(rows and all(row[0] == socket.AF_INET and row[4][0] in PRIVATE_IPS for row in rows))
    finally:
        signal.setitimer(signal.ITIMER_REAL, 0)
        signal.signal(signal.SIGALRM, previous)


def _startup(token: str, *, seconds: int = 75) -> int:
    """Retry read-only readiness only. Processing requests are always sent once."""
    _require(type(seconds) is int and 1 <= seconds <= 75)
    deadline = time.monotonic() + seconds
    attempts = 0
    while time.monotonic() < deadline:
        attempts += 1
        try:
            _private_dns(seconds=min(5, deadline - time.monotonic()))
            remaining = deadline - time.monotonic()
            if remaining <= 0:
                break
            status, headers, body = _request(ORIGIN, "/", token=token, seconds=min(5, remaining))
            _require(status == 404 and headers.get("content-type", "").split(";", 1)[0] == "application/json"
                     and _decode(body) == {"error": "document_inspection_refused"})
            return attempts
        except Exception:
            remaining = deadline - time.monotonic()
            if remaining > 0:
                time.sleep(min(2, remaining))
    raise ConnectivityDenied(_ERROR) from None


def _software(name: str):
    """Load one ordinary module, never the app package or release runner."""
    _require(name in {"resume_layout", "docx_layout_validation", "document_ingestion"})
    bases = (Path("/app/backend/app/services"), Path("/app/app/services"),
             Path(__file__).resolve().parents[2] / "backend/app/services") if "__file__" in globals() else (
                 Path("/app/backend/app/services"), Path("/app/app/services"))
    path = next((base / (name + ".py") for base in bases if (base / (name + ".py")).is_file()), None)
    _require(path is not None and not path.is_symlink())
    module_name = "hirewiz_document_diagnostic_" + name
    spec = importlib.util.spec_from_file_location(module_name, path)
    _require(spec is not None and spec.loader is not None)
    module = importlib.util.module_from_spec(spec)
    sys.modules[module_name] = module
    spec.loader.exec_module(module)
    return module


def _fixtures() -> tuple[bytes, bytes, list[dict[str, str]]]:
    from docx import Document
    from docx.shared import Pt
    patch = _software("resume_layout")
    original = "Built scalable systems with Python and SQL."
    document = Document()
    document.styles["Normal"].font.name = "Cambria"
    document.styles["Normal"].font.size = Pt(11)
    document.add_paragraph("Synthetic Routing Check")
    document.add_paragraph(original)
    document.add_paragraph("Original education and protected running content.")
    output = io.BytesIO()
    document.save(output)
    source = output.getvalue()
    unit = next(unit for unit in patch.extract_source_units(source, "docx") if unit["text"] == original)
    edits = [{"unit_id": unit["unit_id"], "original_text": original,
              "replacement_text": "Built systems with Python and SQL."}]
    return source, patch._apply_docx(source, edits), edits


def _inspect_payload(source: bytes) -> bytes:
    return _canonical({"version": 1, "source_format": "docx", "sha256": _sha(source),
                       "content_base64": base64.b64encode(source).decode("ascii")})


def _layout_payload(source: bytes, candidate: bytes, edits: list[dict[str, str]]) -> tuple[bytes, dict[str, Any]]:
    binding = {"version": 1, "operation": "docx_layout", "sha256": _sha(source),
               "size_bytes": len(source), "source_format": "docx", "candidate_sha256": _sha(candidate),
               "candidate_size_bytes": len(candidate), "edits_sha256": _sha(_canonical(edits))}
    return _canonical({"version": 1, "source_format": "docx", "sha256": binding["sha256"],
                       "content_base64": base64.b64encode(source).decode("ascii"), "edits": edits,
                       "candidate_sha256": binding["candidate_sha256"], "candidate_size_bytes": len(candidate),
                       "worker_image": IMAGE, "policy_sha256": POLICY}), binding


def _scan(value: Any) -> None:
    _require(type(value) is dict and set(value) == {"engine", "version", "definitions"}
             and value["engine"] == "ClamAV" and all(type(value[key]) is str and 0 < len(value[key]) <= 128
             and all(32 <= ord(c) <= 126 for c in value[key]) for key in ("version", "definitions")))


def _layout_receipt(raw: bytes, binding: dict[str, Any]) -> None:
    receipt = _decode(raw)
    _require(set(receipt) == set(binding) | {"worker_image", "policy_sha256", "scan", "candidate_scan", "layout"}
             and all(type(receipt[key]) is type(value) and receipt[key] == value for key, value in binding.items())
             and receipt["worker_image"] == IMAGE and receipt["policy_sha256"] == POLICY)
    _scan(receipt["scan"])
    _scan(receipt["candidate_scan"])
    _require(receipt["scan"] == receipt["candidate_scan"])
    validator = _software("docx_layout_validation")
    validator.check_layout_receipt(receipt["layout"], source_sha256=binding["sha256"],
        source_size_bytes=binding["size_bytes"], candidate_sha256=binding["candidate_sha256"],
        candidate_size_bytes=binding["candidate_size_bytes"], edits_sha256=binding["edits_sha256"],
        image=IMAGE, policy=POLICY, renderer_version=RENDERER, font_manifest_sha256=FONTS)
    _require(receipt["layout"].get("pages") == 1 and receipt["layout"].get("changed_paragraphs") == 1
             and type(receipt["layout"].get("protected_glyphs")) is int
             and receipt["layout"]["protected_glyphs"] > 0)


def _post(path: str, payload: bytes, token: str) -> bytes:
    _require(path in {"/inspect", "/validate-docx"})
    status, headers, raw = _request(ORIGIN, path, method="POST", payload=payload, token=token,
                                   seconds=75, limit=1024 * 1024 if path == "/inspect" else 32768)
    _require(status == 200 and headers.get("content-type", "").split(";", 1)[0] == "application/json")
    return raw


def run() -> dict[str, Any]:
    outcome = {"version": 1, "diagnostic": "document_private_routing", "native_http_acceptance": False,
               "dns_private": False, "runtime_identity": False, "readiness": False,
               "external_https": False, "inspect": False, "validate_docx": False,
               "processing_posts_attempted": 0, "readiness_attempts": 0,
               "stage": "identity", "release_acceptance": False}
    try:
        token = _identity_token()
        outcome["runtime_identity"] = True
        outcome["stage"] = "startup"
        outcome["readiness_attempts"] = _startup(token)
        outcome.update(dns_private=True, readiness=True)
        outcome["stage"] = "external_https"
        code, _, _ = _request(EXTERNAL_ORIGIN, "/", seconds=10)
        _require(code == 200)
        outcome["external_https"] = True
        outcome["stage"] = "synthetic_fixture"
        source, candidate, edits = _fixtures()
        ingestion = _software("document_ingestion")
        config = ingestion.WorkerConfig(ORIGIN + "/inspect", ORIGIN, IMAGE, POLICY)
        outcome["stage"] = "inspect"
        outcome["processing_posts_attempted"] += 1
        inspected = ingestion._decode_receipt(_post("/inspect", _inspect_payload(source), token), source, "docx", config)
        _require("Synthetic Routing Check" in inspected.parsed[0]
                 and "Built scalable systems with Python and SQL." in inspected.parsed[0])
        outcome["inspect"] = True
        payload, binding = _layout_payload(source, candidate, edits)
        outcome["stage"] = "validate_docx"
        outcome["processing_posts_attempted"] += 1
        _layout_receipt(_post("/validate-docx", payload, token), binding)
        outcome.update(validate_docx=True, native_http_acceptance=True, stage="complete")
    except Exception:
        # No raw exception, response, parsed text, token, fixture or provider name.
        outcome["error"] = _ERROR
    return outcome


def plan() -> dict[str, Any]:
    return {"job": JOB, "project": PROJECT, "region": REGION, "identity": IDENTITY,
            "origin": ORIGIN, "image": IMAGE, "policy_sha256": POLICY,
            "renderer_version": RENDERER, "font_manifest_sha256": FONTS,
            "network": "hirewiz-private", "subnet": "hirewiz-callers-us-central1",
            "vpc_egress": "private-ranges-only", "native_calls": 0,
            "native_http_acceptance": False, "release_acceptance": False,
            "processing_post_retries": 0, "job_task_retries": 0}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--execute", action="store_true")
    args = parser.parse_args()
    outcome = run() if args.execute else plan()
    print(json.dumps(outcome, sort_keys=True, separators=(",", ":"), allow_nan=False))
    if args.execute and not outcome["native_http_acceptance"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
